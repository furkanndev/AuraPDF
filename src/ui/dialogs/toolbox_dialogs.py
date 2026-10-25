from typing import List, Dict, Optional, Any
import copy
import os
import sys
import shutil
import subprocess
import zipfile
import uuid
import re
import base64
import io
import math
import secrets
import hashlib
import time
import pymupdf as fitz

HAS_PPTX = True
HAS_OPENPYXL = True

from PyQt6.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QLabel, QFrame,
                             QPushButton, QListWidget, QListWidgetItem, 
                             QLineEdit, QProgressBar, QMessageBox, QFileDialog, 
                             QComboBox, QCheckBox, QGroupBox, QRadioButton, QSpinBox, 
                             QAbstractSpinBox, QAbstractItemView,
                             QTabWidget, QWidget, QSlider, QColorDialog, QButtonGroup, QApplication,
                             QScrollArea, QGridLayout, QFormLayout)
from PyQt6.QtGui import (QDragEnterEvent, QDropEvent, QImage, QPixmap, QIcon, QColor, QPainter, QPen, QBrush, QFont, QFontMetrics, QShortcut, QKeySequence, QTransform)
from PyQt6.QtCore import (Qt, QThread, pyqtSignal, QSize, QByteArray, QBuffer, QIODevice, QRectF, QPointF, QTimer)
from src.core.i18n import LanguageManager
from src.ui.dialog_theme import apply_modern_dialog_style, ModernInfoCard, get_theme_palette
from src.ui.modern_stepper import ModernSpinBox, ModernDoubleSpinBox
from src.ui.icons import ModernIconFactory


def get_spinbox_stylesheet(parent=None) -> str:
    base_dir = getattr(sys, '_MEIPASS', os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    up_candidates = [
        os.path.join(base_dir, "assets", "spin_up.png"),
        os.path.join(base_dir, "src", "assets", "spin_up.png"),
    ]
    down_candidates = [
        os.path.join(base_dir, "assets", "spin_down.png"),
        os.path.join(base_dir, "src", "assets", "spin_down.png"),
    ]
    up_icon = next((p for p in up_candidates if os.path.isfile(p)), up_candidates[0]).replace("\\", "/")
    down_icon = next((p for p in down_candidates if os.path.isfile(p)), down_candidates[0]).replace("\\", "/")

    pal = get_theme_palette(parent)
    bg = pal.get("card", "#242426")
    text = pal.get("text", "#f1f5f9")
    border = pal.get("border", "rgba(255, 255, 255, 0.12)")
    accent = pal.get("accent", "#38bdf8")

    return f"""
        QSpinBox {{
            background-color: {bg};
            color: {text};
            border: 1px solid {border};
            border-radius: 4px;
            padding: 3px 2px 3px 6px;
            padding-right: 42px;
            font-size: 12px;
            font-weight: bold;
        }}
        QSpinBox#PreviewPageSpin {{
            background-color: {bg};
            color: {text};
            border: 1px solid {border};
            border-radius: 4px;
            padding: 2px 2px;
            font-size: 12px;
            font-weight: bold;
            text-align: center;
        }}
        QSpinBox#PreviewPageSpin:focus {{
            border: 1px solid #38bdf8;
        }}
        QSpinBox:focus {{
            border: 1px solid #38bdf8;
        }}
        QSpinBox::down-button {{
            subcontrol-origin: border;
            subcontrol-position: center right;
            right: 21px;
            width: 18px;
            height: 18px;
            border: 1px solid rgba(255, 255, 255, 0.1);
            border-radius: 3px;
            background-color: rgba(255, 255, 255, 0.06);
            image: url("{down_icon}");
        }}
        QSpinBox::down-button:hover {{
            background-color: rgba(255, 255, 255, 0.12);
            border-color: rgba(255, 255, 255, 0.25);
        }}
        QSpinBox::down-button:pressed {{
            background-color: rgba(255, 255, 255, 0.18);
        }}
        QSpinBox::up-button {{
            subcontrol-origin: border;
            subcontrol-position: center right;
            right: 2px;
            width: 18px;
            height: 18px;
            border: 1px solid rgba(255, 255, 255, 0.1);
            border-radius: 3px;
            background-color: rgba(255, 255, 255, 0.06);
            image: url("{up_icon}");
        }}
        QSpinBox::up-button:hover {{
            background-color: rgba(255, 255, 255, 0.12);
            border-color: rgba(255, 255, 255, 0.25);
        }}
        QSpinBox::up-button:pressed {{
            background-color: rgba(255, 255, 255, 0.18);
        }}
    """


# --- 9. PDF ARAÇLARI (TOOLBOX) & DÖNÜŞTÜRÜCÜ DİYALOGLARI ---
def get_ebook_convert_path():
    p = shutil.which("ebook-convert")
    if p and os.path.isfile(p):
        return p
    candidates = [
        r"C:\Program Files\Calibre2\ebook-convert.exe",
        r"C:\Program Files (x86)\Calibre2\ebook-convert.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Calibre2\ebook-convert.exe"),
        os.path.expandvars(r"%PROGRAMFILES%\Calibre2\ebook-convert.exe"),
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    return None


class EbookConversionThread(QThread):
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, input_pdf: str, output_path: str, format_type: str):
        super().__init__()
        self.input_pdf = input_pdf
        self.output_path = output_path
        self.format_type = format_type.lower()

    def run(self):
        try:
            calibre_bin = get_ebook_convert_path()
            if calibre_bin:
                proc = subprocess.run([calibre_bin, self.input_pdf, self.output_path], capture_output=True, text=True)
                if proc.returncode == 0:
                    self.finished_signal.emit(True, f"Başarıyla dönüştürüldü:\n{self.output_path}")
                else:
                    self.finished_signal.emit(False, f"Dönüştürme hatası:\n{proc.stderr}")
            elif self.format_type == 'epub':
                self.convert_builtin_epub()
                self.finished_signal.emit(True, f"Başarıyla dönüştürüldü (Yerleşik Motor):\n{self.output_path}")
            else:
                self.finished_signal.emit(False, "MOBI dönüştürmesi için sistemde Calibre bulunamadı.")
        except Exception as e:
            self.finished_signal.emit(False, f"Dönüştürme sırasında hata oluştu:\n{e}")

    def convert_builtin_epub(self):
        doc = fitz.open(self.input_pdf)
        book_id = str(uuid.uuid4())
        title = os.path.splitext(os.path.basename(self.input_pdf))[0]
        with zipfile.ZipFile(self.output_path, 'w') as zf:
            zf.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
            container_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>'''
            zf.writestr('META-INF/container.xml', container_xml)
            manifest_items, spine_items = [], []
            for idx in range(len(doc)):
                page = doc[idx]
                text = page.get_text("html")
                if not text.strip():
                    plain = page.get_text("text")
                    text = f"<div><pre>{plain}</pre></div>"
                ch_filename = f"chapter_{idx + 1}.xhtml"
                ch_content = f'''<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml">
<head><title>Sayfa {idx + 1}</title>
<style>body {{ font-family: sans-serif; line-height: 1.5; padding: 1em;   }}</style>
</head><body>{text}</body></html>'''
                zf.writestr(f"OEBPS/{ch_filename}", ch_content)
                manifest_items.append(f'<item id="ch_{idx + 1}" href="{ch_filename}" media-type="application/xhtml+xml"/>')
                spine_items.append(f'<itemref idref="ch_{idx + 1}"/>')

            opf_content = f'''<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="BookId" version="2.0">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title>{title}</dc:title><dc:identifier id="BookId">{book_id}</dc:identifier><dc:language>tr</dc:language>
  </metadata>
  <manifest><item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>{''.join(manifest_items)}</manifest>
  <spine toc="ncx">{''.join(spine_items)}</spine>
</package>'''
            zf.writestr('OEBPS/content.opf', opf_content)
            ncx_content = f'''<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <head><meta name="dtb:uid" content="{book_id}"/></head><docTitle><text>{title}</text></docTitle>
  <navMap><navPoint id="navpoint-1" playOrder="1"><navLabel><text>Başlangıç</text></navLabel><content src="chapter_1.xhtml"/></navPoint></navMap>
</ncx>'''
            zf.writestr('OEBPS/toc.ncx', ncx_content)
        doc.close()


class PDFMergeListWidget(QListWidget):
    def __init__(self, parent_dialog):
        super().__init__()
        self.parent_dialog = parent_dialog
        self.setAcceptDrops(True)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragEnterEvent(event)

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragMoveEvent(event)

    def dropEvent(self, event):
        if event.mimeData().hasUrls():
            for u in event.mimeData().urls():
                path = u.toLocalFile()
                if path.lower().endswith(".pdf") and os.path.isfile(path):
                    self.parent_dialog.add_file(path)
            event.acceptProposedAction()
        else:
            super().dropEvent(event)


class PDFMergeDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.setWindowTitle(LanguageManager.tr("merge_title"))
        self.resize(640, 500)
        self.setMinimumSize(520, 380)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(LanguageManager.tr("merge_label")))

        h_center = QHBoxLayout()
        self.list_widget = PDFMergeListWidget(self)
        h_center.addWidget(self.list_widget)

        v_side = QVBoxLayout()
        btn_add = QPushButton(LanguageManager.tr("merge_btn_add"))
        btn_add_active = QPushButton(LanguageManager.tr("merge_btn_add_active"))
        btn_up = QPushButton(LanguageManager.tr("merge_btn_up"))
        btn_down = QPushButton(LanguageManager.tr("merge_btn_down"))
        btn_remove = QPushButton(LanguageManager.tr("merge_btn_remove"))
        btn_clear = QPushButton(LanguageManager.tr("merge_btn_clear"))

        btn_add.clicked.connect(self.choose_files)
        btn_add_active.clicked.connect(self.add_active_doc)
        btn_up.clicked.connect(self.move_up)
        btn_down.clicked.connect(self.move_down)
        btn_remove.clicked.connect(self.remove_selected)
        btn_clear.clicked.connect(self.list_widget.clear)

        v_side.addWidget(btn_add)
        v_side.addWidget(btn_add_active)
        v_side.addWidget(btn_up)
        v_side.addWidget(btn_down)
        v_side.addWidget(btn_remove)
        v_side.addWidget(btn_clear)
        v_side.addStretch()
        h_center.addLayout(v_side)
        layout.addLayout(h_center)

        self.chk_open_tab = QCheckBox(LanguageManager.tr("merge_chk_open"))
        self.chk_open_tab.setChecked(True)
        layout.addWidget(self.chk_open_tab)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.clicked.connect(self.reject)
        btn_merge = QPushButton(LanguageManager.tr("merge_btn_run"))
        btn_merge.setObjectName("PrimaryBtn")
        btn_merge.clicked.connect(self.execute_merge)
        h_btn.addWidget(btn_cancel)
        h_btn.addWidget(btn_merge)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.add_file(tab_doc.file_path)

    def add_file(self, path):
        for i in range(self.list_widget.count()):
            if self.list_widget.item(i).data(Qt.ItemDataRole.UserRole) == path:
                return
        item = QListWidgetItem(os.path.basename(path))
        item.setData(Qt.ItemDataRole.UserRole, path)
        item.setToolTip(path)
        self.list_widget.addItem(item)

    def add_active_doc(self):
        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.add_file(tab_doc.file_path)
        else:
            QMessageBox.information(self, LanguageManager.tr("info"), LanguageManager.tr("no_active_tab"))

    def choose_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, LanguageManager.tr("merge_choose_title"), "", "PDF (*.pdf)")
        for p in paths:
            self.add_file(p)

    def move_up(self):
        row = self.list_widget.currentRow()
        if row > 0:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(row - 1, item)
            self.list_widget.setCurrentRow(row - 1)

    def move_down(self):
        row = self.list_widget.currentRow()
        if row < self.list_widget.count() - 1:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(row + 1, item)
            self.list_widget.setCurrentRow(row + 1)

    def remove_selected(self):
        row = self.list_widget.currentRow()
        if row >= 0:
            self.list_widget.takeItem(row)

    def execute_merge(self):
        if self.list_widget.count() < 2:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("merge_warn_min_2"))
            return

        save_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("merge_save_title"), "birlestirilmis.pdf", "PDF (*.pdf)")
        if not save_path:
            return

        try:
            out_doc = fitz.open()
            for i in range(self.list_widget.count()):
                fpath = self.list_widget.item(i).data(Qt.ItemDataRole.UserRole)
                sub_doc = fitz.open(fpath)
                out_doc.insert_pdf(sub_doc)
                sub_doc.close()

            out_doc.save(save_path, garbage=4, deflate=True, clean=True)
            out_doc.close()

            if self.chk_open_tab.isChecked():
                self.main_window.add_pdf_tab(save_path)

            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("merge_success", path=save_path))
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("merge_error", error=e))


class PDFSplitDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.setWindowTitle(LanguageManager.tr("split_title"))
        self.resize(600, 460)
        self.setMinimumSize(480, 360)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)

        layout.addWidget(QLabel(LanguageManager.tr("split_src_label")))
        h_src = QHBoxLayout()
        self.txt_source = QLineEdit()
        btn_browse_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_src.clicked.connect(self.browse_source)
        h_src.addWidget(self.txt_source)
        h_src.addWidget(btn_browse_src)
        layout.addLayout(h_src)

        grp_opts = QGroupBox(LanguageManager.tr("split_grp_opts"))
        v_opts = QVBoxLayout(grp_opts)

        self.rb_all = QRadioButton(LanguageManager.tr("split_opt_all"))
        self.rb_all.setChecked(True)
        v_opts.addWidget(self.rb_all)

        h_range = QHBoxLayout()
        self.rb_range = QRadioButton(LanguageManager.tr("split_opt_range"))
        self.txt_range = QLineEdit()
        self.txt_range.setPlaceholderText(LanguageManager.tr("split_opt_range_ph"))
        self.txt_range.setEnabled(False)
        self.rb_range.toggled.connect(self.txt_range.setEnabled)
        h_range.addWidget(self.rb_range)
        h_range.addWidget(self.txt_range)
        v_opts.addLayout(h_range)

        h_n = QHBoxLayout()
        self.rb_n = QRadioButton(LanguageManager.tr("split_opt_n"))
        self.spin_n = ModernSpinBox(min_val=1, max_val=9999, default_val=1, parent=self)
        self.spin_n.setFixedWidth(115)
        self.spin_n.setEnabled(False)
        self.rb_n.toggled.connect(self.spin_n.setEnabled)
        h_n.addWidget(self.rb_n)
        h_n.addWidget(self.spin_n)
        v_opts.addLayout(h_n)
        layout.addWidget(grp_opts)

        layout.addWidget(QLabel(LanguageManager.tr("split_out_label")))
        h_out = QHBoxLayout()
        self.txt_dest = QLineEdit()
        btn_browse_dest = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_dest.clicked.connect(self.browse_dest)
        h_out.addWidget(self.txt_dest)
        h_out.addWidget(btn_browse_dest)
        layout.addLayout(h_out)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.clicked.connect(self.reject)
        btn_split = QPushButton(LanguageManager.tr("split_btn_run"))
        btn_split.setObjectName("PrimaryBtn")
        btn_split.clicked.connect(self.execute_split)
        h_btn.addWidget(btn_cancel)
        h_btn.addWidget(btn_split)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.txt_source.setText(tab_doc.file_path)
            self.txt_dest.setText(os.path.dirname(tab_doc.file_path))

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("split_choose_pdf"), "", "PDF (*.pdf)")
        if path:
            self.txt_source.setText(path)
            if not self.txt_dest.text():
                self.txt_dest.setText(os.path.dirname(path))

    def browse_dest(self):
        folder = QFileDialog.getExistingDirectory(self, LanguageManager.tr("split_choose_dir"))
        if folder:
            self.txt_dest.setText(folder)

    def execute_split(self):
        src_path = self.txt_source.text().strip()
        out_dir = self.txt_dest.text().strip()

        if not os.path.isfile(src_path):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_src"))
            return
        if not os.path.isdir(out_dir):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_dest"))
            return

        try:
            doc = fitz.open(src_path)
            total = len(doc)
            base_name = os.path.splitext(os.path.basename(src_path))[0]
            created_count = 0

            if self.rb_all.isChecked():
                for i in range(total):
                    new_doc = fitz.open()
                    new_doc.insert_pdf(doc, from_page=i, to_page=i)
                    out_path = os.path.join(out_dir, f"{base_name}_sayfa_{i + 1}.pdf")
                    new_doc.save(out_path, garbage=4, deflate=True, clean=True)
                    new_doc.close()
                    created_count += 1

            elif self.rb_range.isChecked():
                range_text = self.txt_range.text().strip()
                if not range_text:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_range"))
                    return

                for part in range_text.split(','):
                    part = part.strip()
                    if not part: continue
                    pages = []
                    if '-' in part:
                        try:
                            s, e = map(int, part.split('-'))
                            pages = [p - 1 for p in range(s, e + 1) if 1 <= p <= total]
                        except ValueError:
                            continue
                    else:
                        try:
                            p = int(part)
                            if 1 <= p <= total:
                                pages = [p - 1]
                        except ValueError:
                            continue

                    if pages:
                        new_doc = fitz.open()
                        for p_idx in pages:
                            new_doc.insert_pdf(doc, from_page=p_idx, to_page=p_idx)
                        out_path = os.path.join(out_dir, f"{base_name}_aralik_{part}.pdf")
                        new_doc.save(out_path, garbage=4, deflate=True, clean=True)
                        new_doc.close()
                        created_count += 1

            elif self.rb_n.isChecked():
                n = self.spin_n.value()
                chunk_num = 1
                for i in range(0, total, n):
                    end_idx = min(i + n - 1, total - 1)
                    new_doc = fitz.open()
                    new_doc.insert_pdf(doc, from_page=i, to_page=end_idx)
                    out_path = os.path.join(out_dir, f"{base_name}_parca_{chunk_num}_s{i+1}-{end_idx+1}.pdf")
                    new_doc.save(out_path, garbage=4, deflate=True, clean=True)
                    new_doc.close()
                    chunk_num += 1
                    created_count += 1

            doc.close()
            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("split_success", count=created_count, folder=out_dir))
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("split_error", error=e))


class PDFToEbookDialog(QDialog):
    def __init__(self, main_window, target_format="epub"):
        super().__init__(main_window)
        self.main_window = main_window
        self.target_format = target_format.lower()
        self.thread = None

        self.setWindowTitle(LanguageManager.tr("ebook_title", fmt=self.target_format.upper()))
        self.resize(560, 340)
        self.setMinimumSize(460, 280)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)

        layout.addWidget(QLabel(LanguageManager.tr("ebook_src_label")))
        h_src = QHBoxLayout()
        self.txt_source = QLineEdit()
        btn_browse_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_src.clicked.connect(self.browse_source)
        h_src.addWidget(self.txt_source)
        h_src.addWidget(btn_browse_src)
        layout.addLayout(h_src)

        layout.addWidget(QLabel(LanguageManager.tr("ebook_dest_label", fmt=self.target_format.upper())))
        h_dest = QHBoxLayout()
        self.txt_dest = QLineEdit()
        btn_browse_dest = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_dest.clicked.connect(self.browse_dest)
        h_dest.addWidget(self.txt_dest)
        h_dest.addWidget(btn_browse_dest)
        layout.addLayout(h_dest)

        calibre_path = get_ebook_convert_path()
        if calibre_path:
            engine_text = LanguageManager.tr("ebook_engine_calibre")
            engine_color = "#4caf50"
        elif self.target_format == 'epub':
            engine_text = LanguageManager.tr("ebook_engine_pure")
            engine_color = "#ff9800"
        else:
            engine_text = LanguageManager.tr("ebook_engine_mobi_err")
            engine_color = "#f44336"

        lbl_engine = QLabel(engine_text)
        lbl_engine.setStyleSheet(f"color: {engine_color}; font-weight: bold;")
        layout.addWidget(lbl_engine)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        self.btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        self.btn_cancel.clicked.connect(self.reject)
        self.btn_convert = QPushButton(LanguageManager.tr("ebook_btn_run", fmt=self.target_format.upper()))
        self.btn_convert.setObjectName("PrimaryBtn")
        self.btn_convert.clicked.connect(self.execute_conversion)

        if not calibre_path and self.target_format == 'mobi':
            self.btn_convert.setEnabled(False)

        h_btn.addWidget(self.btn_cancel)
        h_btn.addWidget(self.btn_convert)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.txt_source.setText(tab_doc.file_path)
            base = os.path.splitext(tab_doc.file_path)[0]
            self.txt_dest.setText(f"{base}.{self.target_format}")

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_source.setText(path)
            base = os.path.splitext(path)[0]
            self.txt_dest.setText(f"{base}.{self.target_format}")

    def browse_dest(self):
        ext = self.target_format
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), self.txt_dest.text(), f"{ext.upper()} (*.{ext})")
        if path:
            self.txt_dest.setText(path)

    def execute_conversion(self):
        src = self.txt_source.text().strip()
        dest = self.txt_dest.text().strip()

        if not os.path.isfile(src):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("ebook_warn_src"))
            return
        if not dest:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("ebook_warn_dest", fmt=self.target_format.upper()))
            return

        self.btn_convert.setEnabled(False)
        self.btn_cancel.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)

        self.thread = EbookConversionThread(src, dest, self.target_format)
        self.thread.finished_signal.connect(self.on_conversion_finished)
        self.thread.start()

    def on_conversion_finished(self, success, message):
        self.progress_bar.setVisible(False)
        self.btn_convert.setEnabled(True)
        self.btn_cancel.setEnabled(True)

        if success:
            QMessageBox.information(self, LanguageManager.tr("success"), message)
            self.accept()
        else:
            QMessageBox.critical(self, LanguageManager.tr("error"), message)


class WordConversionThread(QThread):
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, input_pdf: str, output_path: str, page_range: Optional[tuple] = None, temp_file_to_cleanup: Optional[str] = None):
        super().__init__()
        self.input_pdf = input_pdf
        self.output_path = output_path
        self.page_range = page_range
        self.temp_file_to_cleanup = temp_file_to_cleanup

    def run(self):
        temp_norm_pdf = None
        try:
            # Okuyucudaki görsel görünüm ile birebir eşleşmesi için sayfa yönlerini ve döndürmelerini kontrol et
            doc_in = fitz.open(self.input_pdf)
            start_p = 0
            end_p = len(doc_in)
            if self.page_range:
                start_p = max(0, self.page_range[0])
                end_p = min(len(doc_in), self.page_range[1])

            page_rotations = [doc_in[i].rotation for i in range(start_p, end_p)]
            has_rotation = any(rot != 0 for rot in page_rotations)

            conversion_input = self.input_pdf
            conversion_range = self.page_range

            if has_rotation:
                # Döndürülmüş sayfaları PDF okuyucusundaki düzgün görsel yöne (rotation=0) dönüştür/pişir
                import tempfile
                with tempfile.NamedTemporaryFile(suffix="_norm.pdf", delete=False) as tf:
                    temp_norm_pdf = tf.name

                doc_norm = fitz.open()
                for pno in range(start_p, end_p):
                    p = doc_in[pno]
                    new_p = doc_norm.new_page(width=p.rect.width, height=p.rect.height)
                    new_p.show_pdf_page(new_p.rect, doc_in, pno)
                doc_norm.save(temp_norm_pdf)
                doc_norm.close()

                conversion_input = temp_norm_pdf
                conversion_range = None  # Normalizasyon dosyasında sayfalar zaten start_p..end_p olarak ayarlandı

            doc_in.close()

            # 1. Temel Dönüştürücü: pdf2docx (yüksek kaliteli metin, tablo, düzen)
            conv_success = False
            try:
                from pdf2docx import Converter
                cv = Converter(conversion_input)
                if conversion_range:
                    s_p, e_p = conversion_range
                    cv.convert(self.output_path, start=s_p, end=e_p)
                else:
                    cv.convert(self.output_path)
                cv.close()
                conv_success = True
            except Exception:
                conv_success = False

            # 2. İkinci Düzey Dönüştürücü: python-docx + pymupdf
            if not conv_success:
                self.convert_builtin_docx(conversion_input, conversion_range)

            # 3. Word Dosyasını İyileştir: Yatay/Dikey yönlendirmeyi ve resim açılarını düzelt
            self.post_process_docx(page_rotations if has_rotation else None)

            self.finished_signal.emit(True, LanguageManager.tr("word_success", path=self.output_path))
        except Exception as e:
            self.finished_signal.emit(False, LanguageManager.tr("word_error", error=str(e)))
        finally:
            if temp_norm_pdf and os.path.exists(temp_norm_pdf):
                try:
                    os.remove(temp_norm_pdf)
                except Exception:
                    pass
            if self.temp_file_to_cleanup and os.path.exists(self.temp_file_to_cleanup):
                try:
                    os.remove(self.temp_file_to_cleanup)
                except Exception:
                    pass

    def post_process_docx(self, page_rotations: Optional[list] = None):
        import docx
        from docx.enum.section import WD_ORIENT
        import zipfile, io, re
        from PIL import Image

        if not os.path.exists(self.output_path):
            return

        # 1. Word sayfa boyutuna göre Yatay (Landscape) ve Dikey (Portrait) yönünü eşitle
        try:
            doc_word = docx.Document(self.output_path)
            changed = False
            for s in doc_word.sections:
                if s.page_width > s.page_height:
                    if s.orientation != WD_ORIENT.LANDSCAPE:
                        s.orientation = WD_ORIENT.LANDSCAPE
                        changed = True
                else:
                    if s.orientation != WD_ORIENT.PORTRAIT:
                        s.orientation = WD_ORIENT.PORTRAIT
                        changed = True
            if changed:
                doc_word.save(self.output_path)
        except Exception:
            pass

        # 2. Döndürülmüş sayfalardaki resimlerin Word içinde de PDF okuyucusundaki gibi doğru açıda olmasını sağla
        if page_rotations and any(r != 0 for r in page_rotations):
            temp_docx = self.output_path + ".tmp.docx"
            try:
                # İlgili resimlerin hangi sayfalara/bölümlere ait olduğunu tespit et
                image_rot_map = {}
                with zipfile.ZipFile(self.output_path, "r") as zin:
                    if "word/_rels/document.xml.rels" in zin.namelist() and "word/document.xml" in zin.namelist():
                        rels_text = zin.read("word/_rels/document.xml.rels").decode("utf-8")
                        rid_map = {}
                        for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels_text):
                            target = m.group(2)
                            if not target.startswith("word/"):
                                target = "word/" + target.lstrip("/")
                            rid_map[m.group(1)] = target

                        doc_xml = zin.read("word/document.xml").decode("utf-8")
                        parts = re.split(r'<w:sectPr[^>]*>.*?</w:sectPr>', doc_xml)
                        for i, part in enumerate(parts):
                            p_rot = page_rotations[i] if i < len(page_rotations) else page_rotations[-1]
                            for m in re.finditer(r'r:embed="([^"]+)"', part):
                                rid = m.group(1)
                                if rid in rid_map:
                                    image_rot_map[rid_map[rid]] = p_rot

                    with zipfile.ZipFile(temp_docx, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                        for item in zin.infolist():
                            buffer = zin.read(item.filename)
                            if item.filename.startswith("word/media/"):
                                rot = image_rot_map.get(item.filename, 0)
                                if rot != 0:
                                    try:
                                        im = Image.open(io.BytesIO(buffer))
                                        pil_angle = (360 - rot) % 360
                                        im_rot = im.rotate(pil_angle, expand=True)
                                        out_buf = io.BytesIO()
                                        im_rot.save(out_buf, format=im.format or "PNG")
                                        buffer = out_buf.getvalue()
                                    except Exception:
                                        pass
                            elif item.filename == "word/document.xml":
                                try:
                                    xml = buffer.decode("utf-8")
                                    # 90 veya 270 derece döndürülmüş resimlerin genişlik/yükseklik oranını düzenle
                                    for target_file, rot in image_rot_map.items():
                                        if rot in (90, 270):
                                            def swap_ext(m):
                                                return f'<wp:extent cx="{m.group(2)}" cy="{m.group(1)}"/>'
                                            xml = re.sub(r'<wp:extent cx="(\d+)" cy="(\d+)"/>', swap_ext, xml)
                                            def swap_aext(m):
                                                return f'<a:ext cx="{m.group(2)}" cy="{m.group(1)}"/>'
                                            xml = re.sub(r'<a:ext cx="(\d+)" cy="(\d+)"/>', swap_aext, xml)
                                    buffer = xml.encode("utf-8")
                                except Exception:
                                    pass
                            zout.writestr(item, buffer)
                os.replace(temp_docx, self.output_path)
            except Exception:
                if os.path.exists(temp_docx):
                    try:
                        os.remove(temp_docx)
                    except Exception:
                        pass

    def convert_builtin_docx(self, input_pdf: Optional[str] = None, page_range: Optional[tuple] = None):
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
        import tempfile

        pdf_path = input_pdf or self.input_pdf
        doc_in = fitz.open(pdf_path)
        doc_out = docx.Document()

        start_p = 0
        end_p = len(doc_in)
        p_range = page_range if page_range is not None else self.page_range
        if p_range:
            start_p = max(0, p_range[0])
            end_p = min(len(doc_in), p_range[1])

        for page_idx in range(start_p, end_p):
            page = doc_in[page_idx]
            if page_idx == start_p:
                sec = doc_out.sections[0]
            else:
                sec = doc_out.add_section()

            if page.rect.width > page.rect.height:
                sec.orientation = WD_ORIENT.LANDSCAPE
                sec.page_width = Pt(page.rect.width)
                sec.page_height = Pt(page.rect.height)
            else:
                sec.orientation = WD_ORIENT.PORTRAIT
                sec.page_width = Pt(page.rect.width)
                sec.page_height = Pt(page.rect.height)

            text = page.get_text("text")
            if text.strip():
                for para in text.split("\n\n"):
                    clean_para = para.strip()
                    if clean_para:
                        doc_out.add_paragraph(clean_para)
            else:
                try:
                    pix = page.get_pixmap(dpi=150)
                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tf:
                        temp_png = tf.name
                    pix.save(temp_png)
                    doc_out.add_picture(temp_png, width=Pt(page.rect.width - 72), height=Pt(page.rect.height - 72))
                    try:
                        os.remove(temp_png)
                    except Exception:
                        pass
                except Exception:
                    pass

        doc_out.save(self.output_path)
        doc_in.close()


class PDFToWordDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.thread = None

        self.setWindowTitle(LanguageManager.tr("word_title"))
        self.resize(600, 400)
        self.setMinimumSize(480, 320)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        layout.addWidget(QLabel(LanguageManager.tr("word_src_label")))
        h_src = QHBoxLayout()
        self.txt_source = QLineEdit()
        btn_browse_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_src.clicked.connect(self.browse_source)
        h_src.addWidget(self.txt_source)
        h_src.addWidget(btn_browse_src)
        layout.addLayout(h_src)

        layout.addWidget(QLabel(LanguageManager.tr("word_dest_label")))
        h_dest = QHBoxLayout()
        self.txt_dest = QLineEdit()
        btn_browse_dest = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_dest.clicked.connect(self.browse_dest)
        h_dest.addWidget(self.txt_dest)
        h_dest.addWidget(btn_browse_dest)
        layout.addLayout(h_dest)

        grp_pages = QGroupBox(LanguageManager.tr("print_grp_pages"))
        v_pages = QVBoxLayout(grp_pages)
        self.rb_all = QRadioButton(LanguageManager.tr("word_pages_all"))
        self.rb_all.setChecked(True)
        self.rb_range = QRadioButton(LanguageManager.tr("word_pages_range"))
        
        h_range = QHBoxLayout()
        h_range.addWidget(self.rb_range)
        self.txt_range = QLineEdit()
        self.txt_range.setPlaceholderText("Örn: 1-5")
        self.txt_range.setEnabled(False)
        self.rb_range.toggled.connect(self.txt_range.setEnabled)
        h_range.addWidget(self.txt_range)
        
        v_pages.addWidget(self.rb_all)
        v_pages.addLayout(h_range)
        layout.addWidget(grp_pages)

        try:
            import pdf2docx
            engine_text = LanguageManager.tr("word_engine_badge")
            engine_color = "#4caf50"
        except ImportError:
            engine_text = LanguageManager.tr("word_engine_pure")
            engine_color = "#ff9800"

        lbl_engine = QLabel(engine_text)
        lbl_engine.setStyleSheet(f"color: {engine_color}; font-weight: bold;")
        layout.addWidget(lbl_engine)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        self.btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        self.btn_cancel.clicked.connect(self.reject)
        self.btn_convert = QPushButton(LanguageManager.tr("word_btn_run"))
        self.btn_convert.setObjectName("PrimaryBtn")
        self.btn_convert.clicked.connect(self.execute_conversion)

        h_btn.addWidget(self.btn_cancel)
        h_btn.addWidget(self.btn_convert)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.txt_source.setText(tab_doc.file_path)
            base = os.path.splitext(tab_doc.file_path)[0]
            self.txt_dest.setText(f"{base}.docx")

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_source.setText(path)
            base = os.path.splitext(path)[0]
            self.txt_dest.setText(f"{base}.docx")

    def browse_dest(self):
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), self.txt_dest.text(), "Word Document (*.docx)")
        if path:
            if not path.lower().endswith(".docx"):
                path += ".docx"
            self.txt_dest.setText(path)

    def execute_conversion(self):
        src = self.txt_source.text().strip()
        dest = self.txt_dest.text().strip()

        if not os.path.isfile(src):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("ebook_warn_src"))
            return
        if not dest:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("ebook_warn_dest", fmt="DOCX"))
            return

        page_range = None
        if self.rb_range.isChecked():
            txt = self.txt_range.text().strip()
            if "-" in txt:
                try:
                    s_str, e_str = txt.split("-", 1)
                    s_val = max(1, int(s_str.strip())) - 1
                    e_val = max(s_val + 1, int(e_str.strip()))
                    page_range = (s_val, e_val)
                except ValueError:
                    QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_range"))
                    return
            elif txt.isdigit():
                p = int(txt) - 1
                page_range = (p, p + 1)
            else:
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_range"))
                return

        self.btn_convert.setEnabled(False)
        self.btn_cancel.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, 0)

        # Kullanıcı PDF okuyucusunda sayfaları döndürdüyse veya düzenlediyse,
        # aktif belgenin hafızadaki güncel durumunu Word dönüştürücüsüne aktar
        temp_input_pdf = None
        if hasattr(self, 'main_window') and self.main_window:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc and getattr(tab_doc, 'file_path', None):
                try:
                    if os.path.abspath(tab_doc.file_path) == os.path.abspath(src):
                        import tempfile
                        with tempfile.NamedTemporaryFile(suffix="_active.pdf", delete=False) as tf:
                            temp_input_pdf = tf.name
                        temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                        tab_doc.export_drawings_to_pdf(temp_doc)
                        temp_doc.save(temp_input_pdf)
                        temp_doc.close()
                        src = temp_input_pdf
                except Exception:
                    temp_input_pdf = None

        self.thread = WordConversionThread(src, dest, page_range, temp_file_to_cleanup=temp_input_pdf)
        self.thread.finished_signal.connect(self.on_conversion_finished)
        self.thread.start()

    def on_conversion_finished(self, success, message):
        self.progress_bar.setVisible(False)
        self.btn_convert.setEnabled(True)
        self.btn_cancel.setEnabled(True)

        if success:
            dest = self.txt_dest.text().strip()
            reply = QMessageBox.question(
                self,
                LanguageManager.tr("success"),
                f"{message}\n\n{LanguageManager.tr('word_open_now')}",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes
            )
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    os.startfile(dest)
                except Exception:
                    pass
            self.accept()
        else:
            QMessageBox.critical(self, LanguageManager.tr("error"), message)


# --- SAYFA ORGANİZE DİYALOĞU (SIRALAMA, SİLME & NUMARALANDIRMA) ---
def parse_page_range(text: str, total_pages: int) -> List[int]:
    """Kullanıcının girdiği '1-5, 8, 11' gibi sayfa metnini 0-indeksli liste olarak döner."""
    pages = set()
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            tokens = part.split("-")
            try:
                start = max(1, int(tokens[0].strip()))
                end = min(total_pages, int(tokens[1].strip()))
                for p in range(start, end + 1):
                    pages.add(p - 1)
            except ValueError:
                continue
        else:
            try:
                p = int(part)
                if 1 <= p <= total_pages:
                    pages.add(p - 1)
            except ValueError:
                continue
    return sorted(list(pages))


def to_roman(n: int) -> str:
    if n <= 0: return str(n)
    val = [1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1]
    syb = ["M", "CM", "D", "CD", "C", "XC", "L", "XL", "X", "IX", "V", "IV", "I"]
    roman = ""
    i = 0
    while n > 0:
        for _ in range(n // val[i]):
            roman += syb[i]
            n -= val[i]
        i += 1
    return roman


def to_alpha(n: int) -> str:
    if n <= 0: return str(n)
    result = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        result = chr(65 + remainder) + result
    return result


def format_number_by_type(n: int, num_type: str) -> str:
    if num_type == "zero2":
        return f"{n:02d}"
    elif num_type == "zero3":
        return f"{n:03d}"
    elif num_type == "roman_upper":
        return to_roman(n)
    elif num_type == "roman_lower":
        return to_roman(n).lower()
    elif num_type == "alpha_upper":
        return to_alpha(n)
    elif num_type == "alpha_lower":
        return to_alpha(n).lower()
    return str(n)


def get_badge_layout(page_w: float, page_h: float, text_w: float, font_size: float,
                     pos: str, mx: float, my: float, shape_style: str,
                     page_rect: Optional[fitz.Rect] = None,
                     custom_norm_pos: Optional[QPointF] = None):
    """
    Returns (badge_rect, text_rect, align)
    """
    if page_rect is not None:
        x0, y0, x1, y1 = page_rect.x0, page_rect.y0, page_rect.x1, page_rect.y1
    else:
        x0, y0, x1, y1 = 0.0, 0.0, page_w, page_h

    eff_w = x1 - x0
    eff_h = y1 - y0

    px = max(10.0, font_size * 0.9)
    py = max(4.0, font_size * 0.35)
    lh = font_size * 1.35
    bw = max(text_w + px * 2.0, font_size * 2.6)
    bh = max(lh + py * 2.0, font_size * 1.8)

    if shape_style == "circle":
        diam = max(bw, bh, font_size * 2.4)
        bw = diam
        bh = diam
    elif shape_style in ("modern_bar", "accent_dot"):
        bw = max(text_w + px * 2.2 + 8.0, font_size * 2.8)
    elif shape_style == "bracket_modern":
        bw = max(text_w + px * 2.4 + 10.0, font_size * 2.8)

    if pos == "custom" and custom_norm_pos is not None:
        cx = x0 + custom_norm_pos.x() * eff_w
        cy = y0 + custom_norm_pos.y() * eff_h
        bx0 = max(x0, min(x1 - bw, cx - bw / 2.0))
        bx1 = bx0 + bw
        by0 = max(y0, min(y1 - bh, cy - bh / 2.0))
        by1 = by0 + bh
        badge_rect = fitz.Rect(bx0, by0, bx1, by1)
        text_rect = badge_rect
        return badge_rect, text_rect, fitz.TEXT_ALIGN_CENTER

    if pos.startswith("bottom"):
        by1 = y1 - my
        by0 = by1 - bh
    else:
        by0 = y0 + my
        by1 = by0 + bh

    if pos.endswith("center"):
        bx0 = x0 + (eff_w - bw) / 2.0
        bx1 = bx0 + bw
        align = fitz.TEXT_ALIGN_CENTER
    elif pos.endswith("right"):
        bx1 = x1 - mx
        bx0 = bx1 - bw
        align = fitz.TEXT_ALIGN_CENTER if shape_style != "none" else fitz.TEXT_ALIGN_RIGHT
    else:
        bx0 = x0 + mx
        bx1 = bx0 + bw
        align = fitz.TEXT_ALIGN_CENTER if shape_style != "none" else fitz.TEXT_ALIGN_LEFT

    badge_rect = fitz.Rect(bx0, by0, bx1, by1)

    if shape_style == "none":
        box_h = max(26.0, font_size * 2.8)
        if pos.startswith("bottom"):
            text_rect = fitz.Rect(x0 + mx, y1 - my - box_h, x1 - mx, y1 - my)
        else:
            text_rect = fitz.Rect(x0 + mx, y0 + my, x1 - mx, y0 + my + box_h)
        if pos.endswith("center"):
            t_align = fitz.TEXT_ALIGN_CENTER
        elif pos.endswith("right"):
            t_align = fitz.TEXT_ALIGN_RIGHT
        else:
            t_align = fitz.TEXT_ALIGN_LEFT
        return badge_rect, text_rect, t_align

    return badge_rect, badge_rect, fitz.TEXT_ALIGN_CENTER


def embed_page_numbering_info(doc, settings: dict, rects_by_page: dict):
    if not hasattr(doc, 'embfile_names'):
        return
    try:
        import json
        info = {
            "version": 1,
            "settings": settings,
            "rects": {str(k): [list(r) for r in v] for k, v in rects_by_page.items()}
        }
        if "aurapdf_pagenum.json" in doc.embfile_names():
            try:
                doc.embfile_del("aurapdf_pagenum.json")
            except Exception:
                pass
        doc.embfile_add("aurapdf_pagenum.json", json.dumps(info).encode("utf-8"))
    except Exception as e:
        print("Error embedding page numbering info:", e)


def remove_page_numbering_from_doc(doc) -> bool:
    removed = False
    if hasattr(doc, 'embfile_names') and "aurapdf_pagenum.json" in doc.embfile_names():
        try:
            import json
            info = json.loads(doc.embfile_get("aurapdf_pagenum.json").decode("utf-8"))
            rects_map = info.get("rects", {})
            for p_idx_str, r_list in rects_map.items():
                p_idx = int(p_idx_str)
                if 0 <= p_idx < len(doc):
                    page = doc[p_idx]
                    for r_coords in r_list:
                        rect = fitz.Rect(r_coords)
                        exp_rect = fitz.Rect(rect.x0 - 2.0, rect.y0 - 2.0, rect.x1 + 2.0, rect.y1 + 2.0)
                        page.add_redact_annot(exp_rect)
                    page.apply_redactions()
            doc.embfile_del("aurapdf_pagenum.json")
            removed = True
        except Exception as e:
            print("Error removing page numbers:", e)
    if hasattr(doc, 'set_page_labels'):
        try:
            doc.set_page_labels([{'startpage': 0, 'style': 'D', 'firstpagenum': 1}])
        except Exception:
            pass
    return removed


def apply_numbering_to_doc(doc, format_func, get_font_func, font_size: float, opacity: float,
                           num_color: QColor, position: str, mx: float, my: float, shape_style: str,
                           custom_norm_pos: Optional[QPointF] = None, settings_dict: Optional[dict] = None):
    # Eğer belgede önceden eklenmiş sayfa numarası varsa önce temizle
    remove_page_numbering_from_doc(doc)

    total_pages = len(doc)
    font_code = get_font_func()
    base_col = (num_color.redF(), num_color.greenF(), num_color.blueF())

    import os
    from src.ui.direct_text_editor import get_unicode_font_path
    num_font_file = get_unicode_font_path(font_code, bold=False, italic=False)
    if not num_font_file or not os.path.isfile(num_font_file):
        num_font_file = get_unicode_font_path("Arial", bold=False, italic=False)

    try:
        if num_font_file and os.path.isfile(num_font_file):
            fitz_font = fitz.Font(fontfile=num_font_file)
        else:
            fitz_font = fitz.Font(font_code)
    except Exception:
        fitz_font = None

    rects_by_page = {}

    for i in range(total_pages):
        text = format_func(i, total_pages)
        if not text:
            continue
        page = doc[i]
        r = page.rect

        if fitz_font is not None:
            try:
                tw = fitz_font.text_length(text, fontsize=font_size)
            except Exception:
                tw = len(text) * font_size * 0.58
        else:
            tw = len(text) * font_size * 0.58

        rot = page.rotation % 360
        derot = page.derotation_matrix

        badge_rect, text_rect, align = get_badge_layout(
            r.width, r.height, tw, font_size, position, mx, my, shape_style, page_rect=r,
            custom_norm_pos=custom_norm_pos
        )

        badge_rect_page = (badge_rect * derot).normalize()
        text_rect_page = (text_rect * derot).normalize()

        # Kayıt altına al (ileride temizleyebilmek için)
        drawn_union = badge_rect_page | text_rect_page
        rects_by_page[i] = [drawn_union]

        cur_text_col = base_col

        if shape_style == "dark_pill":
            page.draw_rect(badge_rect_page, color=(0.28, 0.33, 0.41), fill=(0.12, 0.16, 0.23),
                           fill_opacity=0.90 * opacity, width=1.0, radius=0.45, overlay=True)
            cur_text_col = (0.97, 0.98, 0.99)
        elif shape_style == "accent_pill":
            page.draw_rect(badge_rect_page, color=(0.38, 0.65, 0.98), fill=(0.10, 0.45, 0.91),
                           fill_opacity=0.92 * opacity, width=1.0, radius=0.45, overlay=True)
            cur_text_col = (1.0, 1.0, 1.0)
        elif shape_style == "pill":
            page.draw_rect(badge_rect_page, color=(0.80, 0.84, 0.88), fill=(1.0, 1.0, 1.0),
                           fill_opacity=0.92 * opacity, width=1.0, radius=0.45, overlay=True)
        elif shape_style == "glassmorphism":
            page.draw_rect(badge_rect_page, color=(0.82, 0.88, 0.95), fill=(0.95, 0.97, 1.0),
                           fill_opacity=0.62 * opacity, width=1.1, radius=0.45, overlay=True)
        elif shape_style == "modern_bar":
            bar_x = badge_rect.x0 + 2.0
            p1 = fitz.Point(bar_x, badge_rect.y0 + 2.5) * derot
            p2 = fitz.Point(bar_x, badge_rect.y1 - 2.5) * derot
            page.draw_line(p1, p2, color=(0.15, 0.48, 0.92), width=2.4, stroke_opacity=opacity, overlay=True)
        elif shape_style == "bracket_modern":
            bx0, bx1 = badge_rect.x0, badge_rect.x1
            by0, by1 = badge_rect.y0 + 2.0, badge_rect.y1 - 2.0
            arm = min(4.0, (bx1 - bx0) * 0.12)
            page.draw_line(fitz.Point(bx0 + arm, by0) * derot, fitz.Point(bx0, by0) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
            page.draw_line(fitz.Point(bx0, by0) * derot, fitz.Point(bx0, by1) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
            page.draw_line(fitz.Point(bx0, by1) * derot, fitz.Point(bx0 + arm, by1) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
            page.draw_line(fitz.Point(bx1 - arm, by0) * derot, fitz.Point(bx1, by0) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
            page.draw_line(fitz.Point(bx1, by0) * derot, fitz.Point(bx1, by1) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
            page.draw_line(fitz.Point(bx1, by1) * derot, fitz.Point(bx1 - arm, by1) * derot, color=(0.35, 0.45, 0.60), width=1.3, stroke_opacity=opacity, overlay=True)
        elif shape_style == "gradient_pill":
            page.draw_rect(badge_rect_page, color=(0.55, 0.36, 0.96), fill=(0.43, 0.23, 0.88),
                           fill_opacity=0.92 * opacity, width=1.0, radius=0.45, overlay=True)
            cur_text_col = (1.0, 1.0, 1.0)
        elif shape_style == "accent_dot":
            dot_rad = max(2.5, font_size * 0.22)
            dot_center = fitz.Point(badge_rect.x0 + dot_rad + 3.0, (badge_rect.y0 + badge_rect.y1) / 2.0) * derot
            page.draw_circle(dot_center, dot_rad, color=(0.15, 0.48, 0.92), fill=(0.15, 0.48, 0.92),
                             fill_opacity=0.95 * opacity, width=0.5, overlay=True)
        elif shape_style == "outline_pill":
            page.draw_rect(badge_rect_page, color=(0.45, 0.55, 0.70), fill=(1.0, 1.0, 1.0),
                           fill_opacity=0.0, width=1.1, radius=0.45, overlay=True)
        elif shape_style == "box":
            page.draw_rect(badge_rect_page, color=(0.70, 0.74, 0.80), fill=(1.0, 1.0, 1.0),
                           fill_opacity=0.92 * opacity, width=1.0, overlay=True)
        elif shape_style == "circle":
            center = fitz.Point((badge_rect.x0 + badge_rect.x1) / 2.0, (badge_rect.y0 + badge_rect.y1) / 2.0) * derot
            rad = min(badge_rect_page.width, badge_rect_page.height) / 2.0
            page.draw_circle(center, rad, color=(0.70, 0.74, 0.80), fill=(1.0, 1.0, 1.0),
                             fill_opacity=0.92 * opacity, width=1.0, overlay=True)
        elif shape_style == "underline":
            p1 = fitz.Point(badge_rect.x0, badge_rect.y1) * derot
            p2 = fitz.Point(badge_rect.x1, badge_rect.y1) * derot
            page.draw_line(p1, p2, color=(0.15, 0.45, 0.90), width=1.8, stroke_opacity=opacity, overlay=True)

        actual_fontname = font_code
        if num_font_file and os.path.isfile(num_font_file):
            try:
                page.insert_font(fontname="F_AuraPDF_NUM", fontfile=num_font_file)
                actual_fontname = "F_AuraPDF_NUM"
            except Exception:
                actual_fontname = font_code

        rc = page.insert_textbox(text_rect_page, text, fontsize=font_size, fontname=actual_fontname,
                                 fontfile=num_font_file if actual_fontname == "F_AuraPDF_NUM" else None,
                                 color=cur_text_col, fill_opacity=opacity, align=align, rotate=rot)
        if rc < 0:
            cx = (badge_rect.x0 + badge_rect.x1) / 2.0
            cy = (badge_rect.y0 + badge_rect.y1) / 2.0 + font_size * 0.35
            page_pt = fitz.Point(cx - tw / 2.0, cy) * derot
            page.insert_text(page_pt, text, fontsize=font_size, fontname=actual_fontname,
                             fontfile=num_font_file if actual_fontname == "F_AuraPDF_NUM" else None,
                             color=cur_text_col, fill_opacity=opacity, rotate=rot)

    try:
        if hasattr(doc, 'subset_fonts'):
            doc.subset_fonts()
    except Exception:
        pass

    if settings_dict is not None:
        embed_page_numbering_info(doc, settings_dict, rects_by_page)
        if hasattr(doc, 'set_page_labels'):
            try:
                is_thesis = (settings_dict.get("page_range_mode") == "thesis")
                if is_thesis:
                    labels = []
                    f_end = int(settings_dict.get("thesis_front_end", 3))
                    f_end_idx = min(total_pages - 1, max(0, f_end - 1))
                    f_type = settings_dict.get("thesis_front_type", "roman_lower")
                    f_skip = bool(settings_dict.get("thesis_skip_cover", True))
                    f_style = 'r' if f_type == "roman_lower" else 'R'
                    f_start = int(settings_dict.get("thesis_front_start", 1))
                    b_start = int(settings_dict.get("thesis_body_start_num", 1))
                    b_type = settings_dict.get("thesis_body_type", "arabic")
                    b_style = 'D'

                    if f_skip:
                        labels.append({'startpage': 0, 'prefix': ''})
                        if f_end_idx >= 1:
                            labels.append({'startpage': 1, 'style': f_style, 'firstpagenum': f_start})
                    else:
                        labels.append({'startpage': 0, 'style': f_style, 'firstpagenum': f_start})

                    if f_end_idx + 1 < total_pages:
                        labels.append({'startpage': f_end_idx + 1, 'style': b_style, 'firstpagenum': b_start})
                    doc.set_page_labels(labels)
                elif settings_dict.get("page_range_mode") != "custom":
                    nt = settings_dict.get("num_type", "arabic")
                    skip_cov = bool(settings_dict.get("skip_cover", False))
                    start_num = int(settings_dict.get("start_num", 1))
                    style_map = {"arabic": "D", "zero2": "D", "zero3": "D", "roman_lower": "r", "roman_upper": "R", "alpha_lower": "a", "alpha_upper": "A"}
                    st = style_map.get(nt, "D")
                    labels = []
                    if skip_cov:
                        labels.append({'startpage': 0, 'prefix': ''})
                        if total_pages > 1:
                            labels.append({'startpage': 1, 'style': st, 'firstpagenum': start_num})
                    else:
                        labels.append({'startpage': 0, 'style': st, 'firstpagenum': start_num})
                    doc.set_page_labels(labels)
            except Exception as e:
                print("Error setting page labels:", e)


class InteractivePageNumberPreviewLabel(QLabel):
    positionChanged = pyqtSignal(QPointF)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.norm_pos = None  # None for presets, QPointF(nx, ny) for custom
        self.is_dragging = False
        self.setMouseTracking(True)
        self.setCursor(Qt.CursorShape.OpenHandCursor)
        self.setToolTip(LanguageManager.tr("page_number_drag_hint"))

    def _event_to_norm_pos(self, event) -> Optional[QPointF]:
        pix = self.pixmap()
        if not pix or pix.isNull():
            return None
        pw = pix.width()
        ph = pix.height()
        ox = (self.width() - pw) / 2.0
        oy = (self.height() - ph) / 2.0
        pos = event.position() if hasattr(event, "position") else event.pos()
        px = pos.x()
        py = pos.y()
        nx = max(0.02, min(0.98, (px - ox) / max(1.0, pw)))
        ny = max(0.02, min(0.98, (py - oy) / max(1.0, ph)))
        return QPointF(nx, ny)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            new_pos = self._event_to_norm_pos(event)
            if new_pos is not None:
                self.is_dragging = True
                self.setCursor(Qt.CursorShape.ClosedHandCursor)
                self.norm_pos = new_pos
                self.positionChanged.emit(self.norm_pos)
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self.is_dragging and (event.buttons() & Qt.MouseButton.LeftButton):
            new_pos = self._event_to_norm_pos(event)
            if new_pos is not None:
                self.norm_pos = new_pos
                self.positionChanged.emit(self.norm_pos)
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.is_dragging = False
            self.setCursor(Qt.CursorShape.OpenHandCursor)
            if self.norm_pos is not None:
                self.positionChanged.emit(self.norm_pos)
        super().mouseReleaseEvent(event)


class PDFOrganizePagesDialog(QDialog):
    def __init__(self, main_window, initial_tab: int = 0):
        super().__init__(main_window)
        self.main_window = main_window
        self.setWindowTitle(LanguageManager.tr("organize_title"))
        self.resize(880, 680)
        self.setMinimumSize(740, 520)
        self.setSizeGripEnabled(True)

        # State
        self.loaded_file_path = None
        self.is_active_tab_doc = False
        self.number_color = QColor("#333333")
        self.custom_preview_pos: Optional[QPointF] = None
        self.preview_current_page = 0
        self.preview_total_pages = 1

        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 14, 14, 14)
        layout.setSpacing(10)

        # 1. Üst Kaynak Dosya Seçici
        h_file = QHBoxLayout()
        h_file.addWidget(QLabel(LanguageManager.tr("reorder_src_label")))
        self.txt_file = QLineEdit()
        self.txt_file.setReadOnly(True)
        btn_browse = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse.clicked.connect(self.choose_file)
        h_file.addWidget(self.txt_file)
        h_file.addWidget(btn_browse)
        layout.addLayout(h_file)

        # 2. Ana Sekmeler
        self.tabs = QTabWidget(self)
        self.tabs.currentChanged.connect(self._on_tab_changed)

        # --- SEKME 0: Sayfa Sıralama & Silme ---
        tab_reorder = QWidget()
        v_reorder = QVBoxLayout(tab_reorder)
        v_reorder.setContentsMargins(12, 12, 12, 12)
        v_reorder.setSpacing(10)

        lbl_hint = QLabel(LanguageManager.tr("reorder_label"))
        lbl_hint.setStyleSheet("")
        v_reorder.addWidget(lbl_hint)

        h_center = QHBoxLayout()
        self.list_widget = QListWidget()
        self.list_widget.setIconSize(QSize(60, 80))
        self.list_widget.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.list_widget.setDefaultDropAction(Qt.DropAction.MoveAction)
        self.list_widget.model().rowsMoved.connect(self._on_rows_moved)
        h_center.addWidget(self.list_widget, stretch=1)

        self.shortcut_del = QShortcut(QKeySequence(Qt.Key.Key_Delete), self.list_widget)
        self.shortcut_del.activated.connect(self.delete_selected_page)

        v_side = QVBoxLayout()
        btn_up = QPushButton(LanguageManager.tr("reorder_btn_up"))
        btn_down = QPushButton(LanguageManager.tr("reorder_btn_down"))
        btn_top = QPushButton(LanguageManager.tr("reorder_btn_top"))
        btn_bottom = QPushButton(LanguageManager.tr("reorder_btn_bottom"))
        btn_reverse = QPushButton(LanguageManager.tr("reorder_btn_reverse"))
        btn_reset = QPushButton(LanguageManager.tr("reorder_btn_reset"))
        btn_rotate_cw = QPushButton(LanguageManager.tr("organize_btn_rotate_cw"))
        btn_rotate_ccw = QPushButton(LanguageManager.tr("organize_btn_rotate_ccw"))
        btn_delete = QPushButton(LanguageManager.tr("reorder_btn_delete"))
        btn_delete.setObjectName("DeleteBtn")

        btn_up.clicked.connect(self.move_up)
        btn_down.clicked.connect(self.move_down)
        btn_top.clicked.connect(self.move_top)
        btn_bottom.clicked.connect(self.move_bottom)
        btn_reverse.clicked.connect(self.reverse_order)
        btn_reset.clicked.connect(self.reset_order)
        btn_rotate_cw.clicked.connect(lambda: self.rotate_selected_page(90))
        btn_rotate_ccw.clicked.connect(lambda: self.rotate_selected_page(-90))
        btn_delete.clicked.connect(self.delete_selected_page)

        v_side.addWidget(btn_up)
        v_side.addWidget(btn_down)
        v_side.addWidget(btn_top)
        v_side.addWidget(btn_bottom)
        v_side.addWidget(btn_reverse)
        v_side.addWidget(btn_reset)
        v_side.addSpacing(6)
        v_side.addWidget(btn_rotate_cw)
        v_side.addWidget(btn_rotate_ccw)
        v_side.addSpacing(6)
        v_side.addWidget(btn_delete)
        v_side.addSpacing(6)
        btn_add_blank = QPushButton(LanguageManager.tr("insert_blank_page"))
        btn_add_blank.setIcon(ModernIconFactory.icon_add_page())
        btn_add_blank.clicked.connect(self.add_blank_page)
        v_side.addWidget(btn_add_blank)
        v_side.addStretch()
        h_center.addLayout(v_side)
        v_reorder.addLayout(h_center)

        # Tab 0 Alt Butonlar
        h_bot_reorder = QHBoxLayout()
        btn_cancel_reorder = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel_reorder.clicked.connect(self.reject)
        self.btn_save_as_reorder = QPushButton(LanguageManager.tr("reorder_btn_save_as"))
        self.btn_save_as_reorder.setObjectName("SaveAsBtn")
        self.btn_save_as_reorder.clicked.connect(self.execute_save_as_reorder)
        self.btn_apply_reorder = QPushButton(LanguageManager.tr("reorder_btn_apply"))
        self.btn_apply_reorder.setObjectName("PrimaryBtn")
        self.btn_apply_reorder.clicked.connect(self.execute_apply_reorder)

        h_bot_reorder.addWidget(btn_cancel_reorder)
        h_bot_reorder.addStretch()
        h_bot_reorder.addWidget(self.btn_save_as_reorder)
        h_bot_reorder.addWidget(self.btn_apply_reorder)
        v_reorder.addLayout(h_bot_reorder)

        self.tabs.addTab(tab_reorder, LanguageManager.tr("organize_tab_reorder"))

        # --- SEKME 1: Sayfa Numaralandırma ---
        tab_number = QWidget()
        h_number = QHBoxLayout(tab_number)
        h_number.setContentsMargins(12, 12, 12, 12)
        h_number.setSpacing(14)

        # Sol Ayarlar Paneli
        v_num_left = QVBoxLayout()
        v_num_left.setSpacing(8)

        # Ayarların dikeyde sıkışmasını ve yazıların kırpılmasını önleyen kaydırma alanı
        scroll_num = QScrollArea()
        scroll_num.setWidgetResizable(True)
        scroll_num.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll_num.setFrameShape(QScrollArea.Shape.NoFrame)
        scroll_num.setStyleSheet("QScrollArea { background: transparent; border: none; }")

        scroll_num_content = QWidget()
        scroll_num_content.setStyleSheet("background: transparent;")
        v_scroll_num = QVBoxLayout(scroll_num_content)
        v_scroll_num.setContentsMargins(0, 0, 6, 0)
        v_scroll_num.setSpacing(10)

        # Grup 1: Format & Şablon
        grp_fmt = QGroupBox(LanguageManager.tr("page_number_fmt_group"))
        v_fmt = QVBoxLayout(grp_fmt)
        v_fmt.setContentsMargins(10, 10, 10, 10)
        v_fmt.setSpacing(6)

        self.widget_type_row = QWidget()
        h_type_row = QHBoxLayout(self.widget_type_row)
        h_type_row.setContentsMargins(0, 0, 0, 0)
        h_type_row.addWidget(QLabel(LanguageManager.tr("page_number_type_lbl")))
        self.cmb_num_type = QComboBox()
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_arabic"), "arabic")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_zero2"), "zero2")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_zero3"), "zero3")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_roman_upper"), "roman_upper")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_roman_lower"), "roman_lower")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_alpha_upper"), "alpha_upper")
        self.cmb_num_type.addItem(LanguageManager.tr("page_number_type_alpha_lower"), "alpha_lower")
        self.cmb_num_type.setCurrentIndex(0)
        h_type_row.addWidget(self.cmb_num_type, stretch=1)
        v_fmt.addWidget(self.widget_type_row)

        is_en = (LanguageManager.current_lang == "en")
        p_word = "Page" if is_en else "Sayfa"
        p_total_word = f"Page {{n}} / Total {{total}} Pages" if is_en else "Sayfa {n} / Toplam {total} Sayfa"

        h_fmt_row = QHBoxLayout()
        h_fmt_row.addWidget(QLabel(LanguageManager.tr("page_number_fmt_preset")))
        self.cmb_preset = QComboBox()
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_only_num"), "{n}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_page_num"), f"{p_word} {{n}}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_page_of"), f"{p_word} {{n}} / {{total}}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_num_of"), "{n} / {total}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_zero_slash"), "{zero_n} / {zero_total}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_modern_capsule"), f"{p_word.upper()} {{zero_n}} • {{zero_total}}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_slash"), "// {n}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_emdash"), "— {n} —")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_pipes"), "| {n} |")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_compact"), f"{'p.' if is_en else 's.'} {{n}} / {{total}}")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_hyphen"), "- {n} -")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_bracket"), "[ {n} ]")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_dots"), "• {n} •")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_tilde"), "~ {n} ~")
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_page_total_words"), p_total_word)
        self.cmb_preset.addItem(LanguageManager.tr("page_number_fmt_custom_opt"), "custom")
        self.cmb_preset.setCurrentIndex(2)
        h_fmt_row.addWidget(self.cmb_preset, stretch=1)
        v_fmt.addLayout(h_fmt_row)

        h_custom_row = QHBoxLayout()
        h_custom_row.addWidget(QLabel(LanguageManager.tr("page_number_fmt_custom")))
        self.txt_custom_format = QLineEdit(f"{p_word} {{n}} / {{total}}")
        self.txt_custom_format.setEnabled(False)
        h_custom_row.addWidget(self.txt_custom_format, stretch=1)
        v_fmt.addLayout(h_custom_row)
        v_scroll_num.addWidget(grp_fmt)

        # Rozet / Şekil & Arka Plan Tasarımı
        grp_shape = QGroupBox(LanguageManager.tr("page_number_shape_group"))
        v_shape = QVBoxLayout(grp_shape)
        v_shape.setContentsMargins(10, 10, 10, 10)
        v_shape.setSpacing(6)

        h_shape_row = QHBoxLayout()
        h_shape_row.addWidget(QLabel(LanguageManager.tr("page_number_shape_lbl")))
        self.cmb_shape_style = QComboBox()
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_none"), "none")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_pill"), "pill")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_glass"), "glassmorphism")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_modern_bar"), "modern_bar")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_bracket"), "bracket_modern")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_gradient_pill"), "gradient_pill")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_accent_dot"), "accent_dot")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_outline_pill"), "outline_pill")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_dark_pill"), "dark_pill")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_accent_pill"), "accent_pill")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_box"), "box")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_circle"), "circle")
        self.cmb_shape_style.addItem(LanguageManager.tr("page_number_shape_underline"), "underline")
        self.cmb_shape_style.setCurrentIndex(0)  # Sade Metin by default!
        h_shape_row.addWidget(self.cmb_shape_style, stretch=1)
        v_shape.addLayout(h_shape_row)
        v_scroll_num.addWidget(grp_shape)

        # Grup 2: Konum & Kenar Boşlukları
        grp_pos = QGroupBox(LanguageManager.tr("page_number_pos_group"))
        v_pos = QVBoxLayout(grp_pos)
        v_pos.setContentsMargins(10, 10, 10, 10)
        v_pos.setSpacing(6)

        h_pos_row = QHBoxLayout()
        h_pos_row.addWidget(QLabel(LanguageManager.tr("page_number_pos_label")))
        self.cmb_position = QComboBox()
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_bottom_center"), "bottom_center")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_bottom_right"), "bottom_right")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_bottom_left"), "bottom_left")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_top_center"), "top_center")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_top_right"), "top_right")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_top_left"), "top_left")
        self.cmb_position.addItem(LanguageManager.tr("page_number_pos_custom"), "custom")
        self.cmb_position.setCurrentIndex(0)
        h_pos_row.addWidget(self.cmb_position, stretch=1)
        v_pos.addLayout(h_pos_row)

        h_pos_row2 = QHBoxLayout()
        h_pos_row2.addWidget(QLabel(LanguageManager.tr("page_number_margin_x")))
        self.spn_margin_x = ModernSpinBox(min_val=5, max_val=200, default_val=36, suffix=" pt", parent=self)
        self.spn_margin_x.setFixedWidth(115)
        h_pos_row2.addWidget(self.spn_margin_x)

        h_pos_row2.addSpacing(8)
        h_pos_row2.addWidget(QLabel(LanguageManager.tr("page_number_margin_y")))
        self.spn_margin_y = ModernSpinBox(min_val=5, max_val=200, default_val=36, suffix=" pt", parent=self)
        self.spn_margin_y.setFixedWidth(115)
        h_pos_row2.addWidget(self.spn_margin_y)
        h_pos_row2.addStretch()
        v_pos.addLayout(h_pos_row2)
        v_scroll_num.addWidget(grp_pos)

        # Grup 3: Sayfa Aralığı & Başlangıç
        grp_range = QGroupBox(LanguageManager.tr("page_number_range_group"))
        v_range = QVBoxLayout(grp_range)
        v_range.setContentsMargins(10, 10, 10, 10)
        v_range.setSpacing(8)

        # Mod Seçimi (Radyo Butonları)
        h_range_mode = QHBoxLayout()
        self.rad_all_pages = QRadioButton(LanguageManager.tr("page_number_range_all"))
        self.rad_all_pages.setChecked(True)
        self.rad_custom_pages = QRadioButton(LanguageManager.tr("page_number_range_custom"))
        self.txt_custom_pages = QLineEdit()
        self.txt_custom_pages.setPlaceholderText(LanguageManager.tr("page_number_range_placeholder"))
        self.txt_custom_pages.setEnabled(False)

        h_range_mode.addWidget(self.rad_all_pages)
        h_range_mode.addWidget(self.rad_custom_pages)
        h_range_mode.addWidget(self.txt_custom_pages, stretch=1)
        v_range.addLayout(h_range_mode)

        self.rad_thesis_mode = QRadioButton(LanguageManager.tr("page_number_range_thesis"))
        v_range.addWidget(self.rad_thesis_mode)

        # 1. Standart Mod Ayarları (Tüm Sayfalar / Özel Sayfalar için)
        self.widget_std_range = QWidget()
        h_std = QHBoxLayout(self.widget_std_range)
        h_std.setContentsMargins(0, 4, 0, 0)
        h_std.setSpacing(10)
        h_std.addWidget(QLabel(LanguageManager.tr("page_number_start_from")))
        self.spn_start_num = ModernSpinBox(min_val=1, max_val=9999, default_val=1, parent=self)
        self.spn_start_num.setFixedWidth(90)
        h_std.addWidget(self.spn_start_num)

        h_std.addSpacing(12)
        self.chk_skip_cover = QCheckBox(LanguageManager.tr("page_number_skip_cover"))
        h_std.addWidget(self.chk_skip_cover)
        h_std.addStretch()
        v_range.addWidget(self.widget_std_range)

        # 2. Bölümlü / Tez Modu Ayarları
        self.widget_thesis = QWidget()
        v_th = QVBoxLayout(self.widget_thesis)
        v_th.setContentsMargins(10, 10, 10, 10)
        v_th.setSpacing(8)
        self.widget_thesis.setStyleSheet("""
            QWidget#ThesisBox {
                background-color: rgba(255, 255, 255, 0.03);
                border: 1px solid rgba(255, 255, 255, 0.1);
                border-radius: 6px;
            }
        """)
        self.widget_thesis.setObjectName("ThesisBox")

        # Bölüm 1: Ön Sayfalar
        lbl_sec1_title = QLabel(f"<b>{LanguageManager.tr('page_number_thesis_sec1_title')}</b>")
        lbl_sec1_title.setStyleSheet(" font-size: 12px;")
        v_th.addWidget(lbl_sec1_title)

        h_th_row1 = QHBoxLayout()
        h_th_row1.setSpacing(8)
        h_th_row1.addWidget(QLabel(LanguageManager.tr("page_number_thesis_front_range")))
        self.spn_thesis_front_end = ModernSpinBox(min_val=1, max_val=9999, default_val=3, parent=self)
        self.spn_thesis_front_end.setFixedWidth(78)
        h_th_row1.addWidget(self.spn_thesis_front_end)

        self.cmb_thesis_front_type = QComboBox()
        self.cmb_thesis_front_type.addItem(LanguageManager.tr("page_number_type_roman_lower"), "roman_lower")
        self.cmb_thesis_front_type.addItem(LanguageManager.tr("page_number_type_roman_upper"), "roman_upper")
        h_th_row1.addWidget(self.cmb_thesis_front_type)

        self.chk_thesis_skip_cover = QCheckBox(LanguageManager.tr("page_number_thesis_skip_cover"))
        self.chk_thesis_skip_cover.setChecked(True)
        h_th_row1.addWidget(self.chk_thesis_skip_cover)

        h_th_row1.addSpacing(6)
        h_th_row1.addWidget(QLabel(LanguageManager.tr("page_number_thesis_front_start")))
        self.spn_thesis_front_start = ModernSpinBox(min_val=1, max_val=100, default_val=1, parent=self)
        self.spn_thesis_front_start.setFixedWidth(72)
        h_th_row1.addWidget(self.spn_thesis_front_start)
        h_th_row1.addStretch()
        v_th.addLayout(h_th_row1)

        # Bölüm 2: Ana Gövde
        lbl_sec2_title = QLabel(f"<b>{LanguageManager.tr('page_number_thesis_sec2_title')}</b>")
        lbl_sec2_title.setStyleSheet(" font-size: 12px;")
        v_th.addWidget(lbl_sec2_title)

        h_th_row2 = QHBoxLayout()
        h_th_row2.setSpacing(8)
        self.lbl_thesis_body_range = QLabel(LanguageManager.tr("page_number_thesis_body_range", start=4, total=1))
        h_th_row2.addWidget(self.lbl_thesis_body_range)

        self.cmb_thesis_body_type = QComboBox()
        self.cmb_thesis_body_type.addItem(LanguageManager.tr("page_number_type_arabic"), "arabic")
        self.cmb_thesis_body_type.addItem(LanguageManager.tr("page_number_type_zero2"), "zero2")
        h_th_row2.addWidget(self.cmb_thesis_body_type)

        h_th_row2.addWidget(QLabel(LanguageManager.tr("page_number_thesis_body_start")))
        self.spn_thesis_body_start_num = ModernSpinBox(min_val=1, max_val=9999, default_val=1, parent=self)
        self.spn_thesis_body_start_num.setFixedWidth(78)
        h_th_row2.addWidget(self.spn_thesis_body_start_num)
        h_th_row2.addStretch()
        v_th.addLayout(h_th_row2)

        self.widget_thesis.setVisible(False)
        v_range.addWidget(self.widget_thesis)

        v_scroll_num.addWidget(grp_range)

        # Grup 4: Yazı Tipi & Görünüm
        grp_font = QGroupBox(LanguageManager.tr("page_number_font_group"))
        v_font = QVBoxLayout(grp_font)
        v_font.setContentsMargins(10, 10, 10, 10)
        v_font.setSpacing(6)

        h_font_row1 = QHBoxLayout()
        h_font_row1.addWidget(QLabel(LanguageManager.tr("page_number_font_family")))
        self.cmb_font_family = QComboBox()
        self.cmb_font_family.addItem("Helvetica (Sans-Serif)", "Helvetica")
        self.cmb_font_family.addItem("Times-Roman (Serif)", "Times-Roman")
        self.cmb_font_family.addItem("Courier (Monospace)", "Courier")
        h_font_row1.addWidget(self.cmb_font_family, stretch=1)
        v_font.addLayout(h_font_row1)

        h_font_row1_1 = QHBoxLayout()
        h_font_row1_1.addWidget(QLabel(LanguageManager.tr("page_number_font_size")))
        self.spn_font_size = ModernSpinBox(min_val=6, max_val=72, default_val=10, suffix=" pt", parent=self)
        self.spn_font_size.setFixedWidth(115)
        h_font_row1_1.addWidget(self.spn_font_size)

        h_font_row1_1.addSpacing(10)
        h_font_row1_1.addWidget(QLabel(LanguageManager.tr("page_number_style_label")))
        self.cmb_font_style = QComboBox()
        self.cmb_font_style.addItem(LanguageManager.tr("font_style_regular"), "regular")
        self.cmb_font_style.addItem(LanguageManager.tr("font_style_bold"), "bold")
        self.cmb_font_style.addItem(LanguageManager.tr("font_style_italic"), "italic")
        h_font_row1_1.addWidget(self.cmb_font_style)
        h_font_row1_1.addStretch()
        v_font.addLayout(h_font_row1_1)

        h_font_row2 = QHBoxLayout()
        h_font_row2.addWidget(QLabel(LanguageManager.tr("page_number_font_color")))
        self.btn_color = QPushButton("      ")
        self.btn_color.setFixedWidth(50)
        self._update_color_btn()
        self.btn_color.clicked.connect(self._choose_color)
        h_font_row2.addWidget(self.btn_color)

        h_font_row2.addSpacing(14)
        h_font_row2.addWidget(QLabel(LanguageManager.tr("page_number_opacity")))
        self.sld_opacity = QSlider(Qt.Orientation.Horizontal)
        self.sld_opacity.setRange(10, 100)
        self.sld_opacity.setValue(100)
        self.lbl_opacity_val = QLabel("100%")
        self.lbl_opacity_val.setFixedWidth(40)
        self.sld_opacity.valueChanged.connect(lambda v: self.lbl_opacity_val.setText(f"{v}%"))
        h_font_row2.addWidget(self.sld_opacity)
        h_font_row2.addWidget(self.lbl_opacity_val)
        v_font.addLayout(h_font_row2)
        v_scroll_num.addWidget(grp_font)
        v_scroll_num.addStretch()

        scroll_num.setWidget(scroll_num_content)
        v_num_left.addWidget(scroll_num, stretch=1)

        # Tab 1 Alt Butonlar - 2 Seviyeli Ferah Düzen (Sıkışmayı ve Taşmayı Önler)
        v_bot_number = QVBoxLayout()
        v_bot_number.setSpacing(8)

        h_bot_num_top = QHBoxLayout()
        self.chk_open_tab = QCheckBox(LanguageManager.tr("encrypt_chk_open"))
        self.chk_open_tab.setChecked(True)
        h_bot_num_top.addWidget(self.chk_open_tab)

        self.lbl_editing_badge = QLabel(LanguageManager.tr("page_number_active_badge"))
        self.lbl_editing_badge.setStyleSheet(" font-size: 11px; font-weight: bold;")
        self.lbl_editing_badge.setVisible(False)
        h_bot_num_top.addWidget(self.lbl_editing_badge)

        h_bot_num_top.addStretch()

        self.btn_remove_number = QPushButton(LanguageManager.tr("page_number_btn_remove"))
        self.btn_remove_number.setObjectName("DeleteBtn")
        self.btn_remove_number.setStyleSheet("min-height: 28px; padding: 4px 12px;")
        self.btn_remove_number.clicked.connect(self.execute_remove_numbering)
        self.btn_remove_number.setVisible(False)
        h_bot_num_top.addWidget(self.btn_remove_number)
        v_bot_number.addLayout(h_bot_num_top)

        h_bot_num_actions = QHBoxLayout()
        h_bot_num_actions.setSpacing(8)
        btn_cancel_number = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel_number.setMinimumHeight(28)
        btn_cancel_number.clicked.connect(self.reject)
        self.btn_save_as_number = QPushButton(LanguageManager.tr("page_number_btn_save_as"))
        self.btn_save_as_number.setObjectName("SaveAsBtn")
        self.btn_save_as_number.setMinimumHeight(28)
        self.btn_save_as_number.clicked.connect(self.execute_save_as_numbering)
        self.btn_apply_number = QPushButton(LanguageManager.tr("page_number_btn_apply"))
        self.btn_apply_number.setObjectName("PrimaryBtn")
        self.btn_apply_number.setMinimumHeight(28)
        self.btn_apply_number.clicked.connect(self.execute_apply_numbering)

        h_bot_num_actions.addWidget(btn_cancel_number)
        h_bot_num_actions.addStretch()
        h_bot_num_actions.addWidget(self.btn_save_as_number)
        h_bot_num_actions.addWidget(self.btn_apply_number)
        v_bot_number.addLayout(h_bot_num_actions)

        v_num_left.addLayout(v_bot_number)

        h_number.addLayout(v_num_left, stretch=1)

        # Tab 1 Sağ: Canlı Önizleme Paneli
        grp_preview = QGroupBox(LanguageManager.tr("preview_title"))
        grp_preview.setFixedWidth(320)
        grp_preview.setStyleSheet("""
            QGroupBox {
                border: 1px solid rgba(255, 255, 255, 0.1);
                border-radius: 6px;
                margin-top: 10px;
                padding-top: 14px;
            }
        """)
        v_prev = QVBoxLayout(grp_preview)
        v_prev.setContentsMargins(10, 12, 10, 12)
        v_prev.setSpacing(8)

        self.lbl_preview_info = QLabel(LanguageManager.tr("preview_page_1"))
        self.lbl_preview_info.setStyleSheet(" font-size: 11px;")
        v_prev.addWidget(self.lbl_preview_info)

        frame_box = QFrame()
        frame_box.setObjectName("PreviewFrame")
        frame_box.setProperty("class", "preview-card")
        v_box = QVBoxLayout(frame_box)
        v_box.setContentsMargins(6, 6, 6, 6)
        v_box.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.lbl_preview = InteractivePageNumberPreviewLabel()
        self.lbl_preview.setFixedSize(250, 350)
        self.lbl_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_preview.setStyleSheet("border: 1px solid rgba(255, 255, 255, 0.15); border-radius: 2px;")
        self.lbl_preview.positionChanged.connect(self._on_preview_position_dragged)
        v_box.addWidget(self.lbl_preview)
        v_prev.addWidget(frame_box)

        # Sayfa Gezinti Çubuğu
        h_nav = QHBoxLayout()
        h_nav.setSpacing(6)
        self.btn_prev_page = QPushButton(LanguageManager.tr("preview_prev"))
        self.btn_prev_page.setProperty("class", "NavBtn")
        self.btn_prev_page.setIcon(ModernIconFactory.icon_chevron_left())
        self.btn_prev_page.setMinimumWidth(84)
        self.btn_prev_page.clicked.connect(self._preview_prev_page)

        self.spn_preview_page = QSpinBox()
        self.spn_preview_page.setObjectName("PreviewPageSpin")
        self.spn_preview_page.setButtonSymbols(QAbstractSpinBox.ButtonSymbols.NoButtons)
        self.spn_preview_page.setRange(1, 1)
        self.spn_preview_page.setValue(1)
        self.spn_preview_page.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.spn_preview_page.setFixedWidth(56)
        self.spn_preview_page.setStyleSheet("""
            QSpinBox#PreviewPageSpin {
                background-color: #1a1a1e;
                color: #f1f5f9;
                border: 1px solid rgba(255, 255, 255, 0.12);
                border-radius: 4px;
                padding: 2px 2px;
                font-size: 12px;
                font-weight: bold;
                text-align: center;
            }
            QSpinBox#PreviewPageSpin:focus {
                border: 1px solid #38bdf8;
            }
        """)
        self.spn_preview_page.valueChanged.connect(self._preview_page_spin_changed)

        self.lbl_total_preview_pages = QLabel("/ 1")
        self.lbl_total_preview_pages.setStyleSheet(" font-size: 11px; font-weight: 500;")

        self.btn_next_page = QPushButton(LanguageManager.tr("preview_next"))
        self.btn_next_page.setProperty("class", "NavBtn")
        self.btn_next_page.setIcon(ModernIconFactory.icon_chevron_right())
        self.btn_next_page.setMinimumWidth(84)
        self.btn_next_page.clicked.connect(self._preview_next_page)

        h_nav.addWidget(self.btn_prev_page)
        h_nav.addWidget(self.spn_preview_page)
        h_nav.addWidget(self.lbl_total_preview_pages)
        h_nav.addWidget(self.btn_next_page)
        v_prev.addLayout(h_nav)
        v_prev.addStretch()

        h_number.addWidget(grp_preview, stretch=0)

        self.tabs.addTab(tab_number, LanguageManager.tr("organize_tab_numbering"))
        layout.addWidget(self.tabs)

        # Tab 1 Sinyalleri
        self.cmb_num_type.currentIndexChanged.connect(self.update_preview)
        self.cmb_shape_style.currentIndexChanged.connect(self.update_preview)
        self.cmb_preset.currentIndexChanged.connect(self._on_preset_changed)
        self.txt_custom_format.textChanged.connect(self.update_preview)
        self.cmb_position.currentIndexChanged.connect(self._on_position_changed)
        self.spn_margin_x.valueChanged.connect(self.update_preview)
        self.spn_margin_y.valueChanged.connect(self.update_preview)
        self.spn_start_num.valueChanged.connect(self.update_preview)
        self.chk_skip_cover.toggled.connect(self.update_preview)
        self.rad_all_pages.toggled.connect(self._on_range_mode_changed)
        self.rad_custom_pages.toggled.connect(self._on_range_mode_changed)
        self.rad_thesis_mode.toggled.connect(self._on_range_mode_changed)
        self.txt_custom_pages.textChanged.connect(self.update_preview)
        self.spn_thesis_front_end.valueChanged.connect(self._on_thesis_front_end_changed)
        self.cmb_thesis_front_type.currentIndexChanged.connect(self.update_preview)
        self.chk_thesis_skip_cover.toggled.connect(self.update_preview)
        self.spn_thesis_front_start.valueChanged.connect(self.update_preview)
        self.cmb_thesis_body_type.currentIndexChanged.connect(self.update_preview)
        self.spn_thesis_body_start_num.valueChanged.connect(self.update_preview)
        self.cmb_font_family.currentIndexChanged.connect(self.update_preview)
        self.spn_font_size.valueChanged.connect(self.update_preview)
        self.cmb_font_style.currentIndexChanged.connect(self.update_preview)
        self.sld_opacity.valueChanged.connect(self.update_preview)

        # Aktif açık belge varsa yükle
        get_tab = getattr(self.main_window, 'get_current_tab_doc', lambda: None)
        tab_doc = get_tab() if callable(get_tab) else None
        if tab_doc and tab_doc.doc is not None:
            self.load_from_tab_doc(tab_doc)
        else:
            self.btn_apply_reorder.setEnabled(False)
            self.btn_save_as_reorder.setEnabled(False)
            self.btn_apply_number.setEnabled(False)
            self.btn_save_as_number.setEnabled(False)

        if initial_tab in (0, 1):
            self.tabs.setCurrentIndex(initial_tab)

    def _get_current_numbering_settings(self) -> dict:
        return {
            "num_type": self.cmb_num_type.currentData() or "arabic",
            "shape_style": self.cmb_shape_style.currentData() or "none",
            "preset": self.cmb_preset.currentData() or "{n}",
            "custom_format": self.txt_custom_format.text(),
            "position": self.cmb_position.currentData() or "bottom_center",
            "custom_pos": [self.custom_preview_pos.x(), self.custom_preview_pos.y()] if self.custom_preview_pos else None,
            "margin_x": self.spn_margin_x.value(),
            "margin_y": self.spn_margin_y.value(),
            "start_num": self.spn_start_num.value(),
            "skip_cover": self.chk_skip_cover.isChecked(),
            "page_range_mode": "thesis" if self.rad_thesis_mode.isChecked() else ("custom" if self.rad_custom_pages.isChecked() else "all"),
            "custom_pages": self.txt_custom_pages.text(),
            "thesis_front_end": self.spn_thesis_front_end.value(),
            "thesis_front_type": self.cmb_thesis_front_type.currentData() or "roman_lower",
            "thesis_skip_cover": self.chk_thesis_skip_cover.isChecked(),
            "thesis_front_start": self.spn_thesis_front_start.value(),
            "thesis_body_start_num": self.spn_thesis_body_start_num.value(),
            "thesis_body_type": self.cmb_thesis_body_type.currentData() or "arabic",
            "font_family": self.cmb_font_family.currentData() or "Helvetica",
            "font_size": self.spn_font_size.value(),
            "font_style": self.cmb_font_style.currentData() or "regular",
            "opacity": self.sld_opacity.value(),
            "color": self.number_color.name()
        }

    def _restore_saved_numbering_settings(self, s: dict):
        if not s:
            return
        self.blockSignals(True)
        try:
            if "num_type" in s:
                idx = self.cmb_num_type.findData(s["num_type"])
                if idx >= 0: self.cmb_num_type.setCurrentIndex(idx)
            if "shape_style" in s:
                idx = self.cmb_shape_style.findData(s["shape_style"])
                if idx >= 0: self.cmb_shape_style.setCurrentIndex(idx)
            if "preset" in s:
                idx = self.cmb_preset.findData(s["preset"])
                if idx >= 0:
                    self.cmb_preset.setCurrentIndex(idx)
                else:
                    idx_c = self.cmb_preset.findData("custom")
                    if idx_c >= 0: self.cmb_preset.setCurrentIndex(idx_c)
            if "custom_format" in s:
                self.txt_custom_format.setText(str(s["custom_format"]))
            if "position" in s:
                idx = self.cmb_position.findData(s["position"])
                if idx >= 0: self.cmb_position.setCurrentIndex(idx)
            if s.get("custom_pos"):
                self.custom_preview_pos = QPointF(float(s["custom_pos"][0]), float(s["custom_pos"][1]))
                if hasattr(self, 'lbl_preview'):
                    self.lbl_preview.norm_pos = self.custom_preview_pos
            else:
                self.custom_preview_pos = None
                if hasattr(self, 'lbl_preview'):
                    self.lbl_preview.norm_pos = None
            if "margin_x" in s:
                self.spn_margin_x.setValue(int(s["margin_x"]))
            if "margin_y" in s:
                self.spn_margin_y.setValue(int(s["margin_y"]))
            if "start_num" in s:
                self.spn_start_num.setValue(int(s["start_num"]))
            if "skip_cover" in s:
                self.chk_skip_cover.setChecked(bool(s["skip_cover"]))
            if s.get("page_range_mode") == "thesis":
                self.rad_thesis_mode.setChecked(True)
            elif s.get("page_range_mode") == "custom":
                self.rad_custom_pages.setChecked(True)
            else:
                self.rad_all_pages.setChecked(True)
            if "custom_pages" in s:
                self.txt_custom_pages.setText(str(s["custom_pages"]))
            if "thesis_front_end" in s:
                self.spn_thesis_front_end.setValue(int(s["thesis_front_end"]))
            if "thesis_front_type" in s:
                idx = self.cmb_thesis_front_type.findData(s["thesis_front_type"])
                if idx >= 0: self.cmb_thesis_front_type.setCurrentIndex(idx)
            if "thesis_skip_cover" in s:
                self.chk_thesis_skip_cover.setChecked(bool(s["thesis_skip_cover"]))
            if "thesis_front_start" in s:
                self.spn_thesis_front_start.setValue(int(s["thesis_front_start"]))
            if "thesis_body_start_num" in s:
                self.spn_thesis_body_start_num.setValue(int(s["thesis_body_start_num"]))
            if "thesis_body_type" in s:
                idx = self.cmb_thesis_body_type.findData(s["thesis_body_type"])
                if idx >= 0: self.cmb_thesis_body_type.setCurrentIndex(idx)
            if "font_family" in s:
                idx = self.cmb_font_family.findData(s["font_family"])
                if idx >= 0: self.cmb_font_family.setCurrentIndex(idx)
            if "font_size" in s:
                self.spn_font_size.setValue(int(s["font_size"]))
            if "font_style" in s:
                idx = self.cmb_font_style.findData(s["font_style"])
                if idx >= 0: self.cmb_font_style.setCurrentIndex(idx)
            if "opacity" in s:
                self.sld_opacity.setValue(int(s["opacity"]))
            if "color" in s:
                self.number_color = QColor(s["color"])
                self._update_color_btn()
        finally:
            self.blockSignals(False)
        self._on_preset_changed(self.cmb_preset.currentIndex())
        self._on_range_mode_changed()
        self.update_preview()

    def load_from_tab_doc(self, tab_doc):
        self.loaded_file_path = tab_doc.file_path
        self.is_active_tab_doc = True
        self.txt_file.setText(tab_doc.file_path or "Aktif Belge")
        self.populate_pages(tab_doc.doc)
        self.btn_apply_reorder.setEnabled(True)
        self.btn_save_as_reorder.setEnabled(True)
        self.btn_apply_number.setEnabled(True)
        self.btn_save_as_number.setEnabled(True)

        has_num = False
        if getattr(tab_doc, 'has_page_numbers', False):
            has_num = True
        elif hasattr(tab_doc.doc, 'embfile_names') and 'aurapdf_pagenum.json' in tab_doc.doc.embfile_names():
            has_num = True
            tab_doc.has_page_numbers = True

        if has_num:
            self.lbl_editing_badge.setVisible(True)
            self.btn_remove_number.setVisible(True)
            self.btn_apply_number.setText(LanguageManager.tr("page_number_btn_update"))
            saved = getattr(tab_doc, 'page_numbering_settings', None)
            if not saved and hasattr(tab_doc.doc, 'embfile_names') and 'aurapdf_pagenum.json' in tab_doc.doc.embfile_names():
                try:
                    import json
                    pdata = json.loads(tab_doc.doc.embfile_get('aurapdf_pagenum.json').decode('utf-8'))
                    saved = pdata.get('settings', {})
                    tab_doc.page_numbering_settings = saved
                except Exception:
                    saved = None
            if saved:
                self._restore_saved_numbering_settings(saved)
        else:
            self.lbl_editing_badge.setVisible(False)
            self.btn_remove_number.setVisible(False)
            self.btn_apply_number.setText(LanguageManager.tr("page_number_btn_apply"))

    def populate_pages(self, doc):
        self.list_widget.clear()
        for idx in range(len(doc)):
            try:
                page = doc[idx]
                rot = page.rotation % 360
                mat = fitz.Matrix(0.4, 0.4)
                pix = page.get_pixmap(matrix=mat)
                qimg = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
                icon = QIcon(QPixmap.fromImage(qimg))
            except Exception:
                rot = 0
                icon = QIcon()

            item = QListWidgetItem(icon, LanguageManager.tr("reorder_item_text", new_idx=idx + 1, orig_idx=idx + 1))
            item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable | Qt.ItemFlag.ItemIsDragEnabled | Qt.ItemFlag.ItemIsDropEnabled)
            item.setData(Qt.ItemDataRole.UserRole, idx)
            item.setData(Qt.ItemDataRole.UserRole + 1, rot)
            self.list_widget.addItem(item)

        if self.list_widget.count() > 0:
            self.list_widget.setCurrentRow(0)

        total = max(1, len(doc))
        self.preview_total_pages = total
        self.spn_preview_page.blockSignals(True)
        self.spn_preview_page.setRange(1, total)
        self.spn_preview_page.setValue(1)
        self.spn_preview_page.blockSignals(False)
        self.lbl_total_preview_pages.setText(f"/ {total}")
        if hasattr(self, 'spn_thesis_front_end'):
            self.spn_thesis_front_end.blockSignals(True)
            self.spn_thesis_front_end.setRange(1, max(1, total - 1))
            if self.spn_thesis_front_end.value() > max(1, total - 1):
                self.spn_thesis_front_end.setValue(min(3, max(1, total - 1)))
            self.spn_thesis_front_end.blockSignals(False)
            self._update_thesis_body_label()
        self.update_preview()

    def _on_rows_moved(self, *args):
        for i in range(self.list_widget.count()):
            item = self.list_widget.item(i)
            orig_idx = item.data(Qt.ItemDataRole.UserRole)
            item.setText(LanguageManager.tr("reorder_item_text", new_idx=i + 1, orig_idx=orig_idx + 1))
        self.update_preview()

    def _get_current_order(self) -> List[int]:
        return [self.list_widget.item(i).data(Qt.ItemDataRole.UserRole) for i in range(self.list_widget.count())]

    def _on_tab_changed(self, idx: int):
        if idx == 1:
            total = max(1, self.list_widget.count())
            self.preview_total_pages = total
            self.preview_current_page = max(0, min(self.preview_current_page, total - 1))
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setRange(1, total)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.lbl_total_preview_pages.setText(f"/ {total}")
            if hasattr(self, 'spn_thesis_front_end'):
                self.spn_thesis_front_end.blockSignals(True)
                self.spn_thesis_front_end.setRange(1, max(1, total - 1))
                self.spn_thesis_front_end.blockSignals(False)
                self._update_thesis_body_label()
            self.update_preview()

    def move_up(self):
        row = self.list_widget.currentRow()
        if row > 0:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(row - 1, item)
            self.list_widget.setCurrentRow(row - 1)
            self._on_rows_moved()

    def move_down(self):
        row = self.list_widget.currentRow()
        if 0 <= row < self.list_widget.count() - 1:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(row + 1, item)
            self.list_widget.setCurrentRow(row + 1)
            self._on_rows_moved()

    def move_top(self):
        row = self.list_widget.currentRow()
        if row > 0:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(0, item)
            self.list_widget.setCurrentRow(0)
            self._on_rows_moved()

    def move_bottom(self):
        row = self.list_widget.currentRow()
        if 0 <= row < self.list_widget.count() - 1:
            item = self.list_widget.takeItem(row)
            self.list_widget.insertItem(self.list_widget.count(), item)
            self.list_widget.setCurrentRow(self.list_widget.count() - 1)
            self._on_rows_moved()

    def reverse_order(self):
        items = [self.list_widget.takeItem(0) for _ in range(self.list_widget.count())]
        for it in reversed(items):
            self.list_widget.addItem(it)
        self._on_rows_moved()
        if self.list_widget.count() > 0:
            self.list_widget.setCurrentRow(0)

    def reset_order(self):
        tab_doc = self.main_window.get_current_tab_doc()
        if self.is_active_tab_doc and tab_doc and tab_doc.doc is not None:
            self.populate_pages(tab_doc.doc)
        elif self.loaded_file_path and os.path.isfile(self.loaded_file_path):
            try:
                doc = fitz.open(self.loaded_file_path)
                self.populate_pages(doc)
                doc.close()
            except Exception:
                pass
        else:
            items = [self.list_widget.takeItem(0) for _ in range(self.list_widget.count())]
            items.sort(key=lambda it: it.data(Qt.ItemDataRole.UserRole))
            for it in items:
                self.list_widget.addItem(it)
            self._on_rows_moved()
            if self.list_widget.count() > 0:
                self.list_widget.setCurrentRow(0)

    def delete_selected_page(self):
        row = self.list_widget.currentRow()
        if row < 0:
            return
        if self.list_widget.count() <= 1:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("reorder_warn_min_pages"))
            return

        item = self.list_widget.takeItem(row)
        del item
        self._on_rows_moved()
        new_row = min(row, self.list_widget.count() - 1)
        self.list_widget.setCurrentRow(new_row)

    def add_blank_page(self):
        row = self.list_widget.currentRow()
        insert_idx = (row + 1) if row >= 0 else self.list_widget.count()

        from PyQt6.QtWidgets import QMenu
        from PyQt6.QtGui import QCursor
        menu = QMenu(self)
        menu.setStyleSheet("""
            QMenu { background-color: #242426; color: #f1f5f9; border: 1px solid rgba(255, 255, 255, 0.12); border-radius: 6px; padding: 4px; }
            QMenu::item { padding: 8px 24px; border-radius: 4px; font-size: 13px; }
            QMenu::item:selected { background-color: #2563eb; color: #ffffff; }
        """)

        styles = [
            (LanguageManager.tr("new_page_choose_style"), "dialog"),
            ("Boş Sayfa", "blank"),
            ("Kareli Sayfa (Grid)", "grid"),
            ("Çizgili Sayfa", "lined"),
            ("Noktalı Sayfa", "dotted"),
            ("Cornell Not Sayfası", "cornell"),
            ("Müzik Portesi", "music"),
            ("Karanlık Mod", "dark"),
            ("Kapak Sayfası", "cover")
        ]

        for name, style_code in styles:
            action = menu.addAction(name)
            action.setData(style_code)

        chosen = menu.exec(QCursor.pos())
        if not chosen:
            return
        style = chosen.data()

        if style == "dialog":
            from src.ui.dialogs.new_page_dialog import NewPageDialog
            dlg = NewPageDialog(self.main_window, self)
            if not dlg.exec():
                return
            cfg = dlg.get_page_config()
            style = cfg.get("style", "blank")
            width = cfg.get("width", 595.32)
            height = cfg.get("height", 841.92)
            count = cfg.get("count", 1)
            if self.is_active_tab_doc:
                tab_doc = self.main_window.get_current_tab_doc()
                if tab_doc and tab_doc.doc is not None:
                    for i in range(count):
                        is_last = (i == count - 1)
                        tab_doc.insert_blank_page(
                            after_page=(row + i) if row >= 0 else len(tab_doc.doc) - 1,
                            width=width,
                            height=height,
                            style=style,
                            rerender=is_last,
                            notify_thumb=is_last
                        )
                    self.load_from_tab_doc(tab_doc)
                    self.list_widget.setCurrentRow(min(insert_idx, self.list_widget.count() - 1))
                    return

        if self.is_active_tab_doc:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc is not None:
                tab_doc.insert_blank_page(after_page=row if row >= 0 else len(tab_doc.doc) - 1, style=style)
                self.load_from_tab_doc(tab_doc)
                self.list_widget.setCurrentRow(min(insert_idx, self.list_widget.count() - 1))
                return

        if self.loaded_file_path and os.path.isfile(self.loaded_file_path):
            try:
                import fitz
                doc = fitz.open(self.loaded_file_path)
                width = 595.32
                height = 841.92
                rot = 0
                if len(doc) > 0:
                    ref_idx = max(0, min(row if row >= 0 else len(doc) - 1, len(doc) - 1))
                    ref_p = doc[ref_idx]
                    width = ref_p.mediabox.width
                    height = ref_p.mediabox.height
                    rot = ref_p.rotation
                new_p = doc.new_page(pno=insert_idx, width=width, height=height)
                if rot:
                    new_p.set_rotation(rot)

                if style and style != "blank":
                    from src.document.tab_document import PDFTabDocument
                    PDFTabDocument._draw_page_template(None, new_p, width, height, style)
                import tempfile
                with tempfile.NamedTemporaryFile(suffix="_blank.pdf", delete=False) as tf:
                    temp_p = tf.name
                doc.save(temp_p)
                doc.close()
                self.loaded_file_path = temp_p
                doc_new = fitz.open(temp_p)
                self.populate_pages(doc_new)
                doc_new.close()
                self.list_widget.setCurrentRow(min(insert_idx, self.list_widget.count() - 1))
            except Exception as e:
                QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def rotate_selected_page(self, delta: int = 90):
        items = self.list_widget.selectedItems()
        if not items:
            row = self.list_widget.currentRow()
            if row >= 0:
                items = [self.list_widget.item(row)]
        if not items:
            return

        doc = None
        should_close = False
        if self.is_active_tab_doc:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc is not None:
                doc = tab_doc.doc
        elif self.loaded_file_path and os.path.isfile(self.loaded_file_path):
            try:
                doc = fitz.open(self.loaded_file_path)
                should_close = True
            except Exception:
                doc = None

        for item in items:
            orig_idx = item.data(Qt.ItemDataRole.UserRole)
            cur_rot = item.data(Qt.ItemDataRole.UserRole + 1)
            if cur_rot is None:
                cur_rot = 0
            new_rot = (cur_rot + delta) % 360
            item.setData(Qt.ItemDataRole.UserRole + 1, new_rot)

            if doc is not None and 0 <= orig_idx < len(doc):
                try:
                    page = doc[orig_idx]
                    diff_rot = (new_rot - page.rotation) % 360
                    mat = fitz.Matrix(0.4, 0.4).prerotate(diff_rot)
                    pix = page.get_pixmap(matrix=mat)
                    qimg = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
                    item.setIcon(QIcon(QPixmap.fromImage(qimg)))
                except Exception:
                    pass

        if should_close and doc:
            doc.close()

        self.update_preview()

    def choose_file(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("reorder_choose_title"), "", "PDF (*.pdf)")
        if path and os.path.isfile(path):
            try:
                tab_doc = self.main_window.get_current_tab_doc()
                if tab_doc and tab_doc.file_path and os.path.abspath(tab_doc.file_path) == os.path.abspath(path):
                    self.load_from_tab_doc(tab_doc)
                    return

                doc = fitz.open(path)
                self.loaded_file_path = path
                self.txt_file.setText(path)
                self.is_active_tab_doc = False
                self.btn_apply_reorder.setEnabled(False)
                self.btn_save_as_reorder.setEnabled(True)
                self.btn_apply_number.setEnabled(False)
                self.btn_save_as_number.setEnabled(True)
                self.btn_remove_number.setVisible(False)

                has_emb = hasattr(doc, 'embfile_names') and 'aurapdf_pagenum.json' in doc.embfile_names()
                if has_emb:
                    self.lbl_editing_badge.setVisible(True)
                    try:
                        import json
                        pdata = json.loads(doc.embfile_get('aurapdf_pagenum.json').decode('utf-8'))
                        self._restore_saved_numbering_settings(pdata.get('settings', {}))
                    except Exception:
                        pass
                else:
                    self.lbl_editing_badge.setVisible(False)

                self.populate_pages(doc)
                doc.close()
            except Exception as e:
                QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def _update_color_btn(self):
        c = self.number_color.name()
        self.btn_color.setStyleSheet(f"background-color: {c}; border: 1px solid rgba(255, 255, 255, 0.25); border-radius: 3px;")

    def _choose_color(self):
        col = QColorDialog.getColor(self.number_color, self, LanguageManager.tr("page_number_font_color"))
        if col.isValid():
            self.number_color = col
            self._update_color_btn()
            self.update_preview()

    def _on_preset_changed(self, idx):
        data = self.cmb_preset.currentData()
        if data == "custom":
            self.txt_custom_format.setEnabled(True)
            self.txt_custom_format.setFocus()
        else:
            self.txt_custom_format.setEnabled(False)
            self.txt_custom_format.setText(str(data))
        self.update_preview()

    def _on_position_changed(self, idx):
        if self.cmb_position.currentData() != "custom":
            self.custom_preview_pos = None
            if hasattr(self, 'lbl_preview') and hasattr(self.lbl_preview, 'norm_pos'):
                self.lbl_preview.norm_pos = None
        self.update_preview()

    def _on_preview_position_dragged(self, pt: QPointF):
        self.custom_preview_pos = pt
        idx = self.cmb_position.findData("custom")
        if idx >= 0 and self.cmb_position.currentIndex() != idx:
            self.cmb_position.blockSignals(True)
            self.cmb_position.setCurrentIndex(idx)
            self.cmb_position.blockSignals(False)
        self.update_preview()

    def _preview_prev_page(self):
        if self.preview_current_page > 0:
            self.preview_current_page -= 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_next_page(self):
        if self.preview_current_page < self.preview_total_pages - 1:
            self.preview_current_page += 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_page_spin_changed(self, val):
        new_page = max(0, min(val - 1, self.preview_total_pages - 1))
        if new_page != self.preview_current_page:
            self.preview_current_page = new_page
            self.update_preview()

    def _on_range_mode_changed(self):
        is_thesis = self.rad_thesis_mode.isChecked()
        self.widget_std_range.setVisible(not is_thesis)
        self.widget_thesis.setVisible(is_thesis)
        if hasattr(self, 'widget_type_row'):
            self.widget_type_row.setVisible(not is_thesis)
        self.txt_custom_pages.setEnabled(self.rad_custom_pages.isChecked())
        self._update_thesis_body_label()
        self.update_preview()

    def _on_thesis_front_end_changed(self, val):
        self._update_thesis_body_label()
        self.update_preview()

    def _update_thesis_body_label(self):
        if not hasattr(self, 'spn_thesis_front_end') or not hasattr(self, 'lbl_thesis_body_range'):
            return
        front_end = self.spn_thesis_front_end.value()
        start = front_end + 1
        total = self.preview_total_pages
        self.lbl_thesis_body_range.setText(LanguageManager.tr("page_number_thesis_body_range", start=start, total=total))

    def _format_page_text(self, page_idx: int, total_pages: int) -> Optional[str]:
        if self.rad_thesis_mode.isChecked():
            front_end = self.spn_thesis_front_end.value()
            front_end_idx = min(total_pages - 1, max(0, front_end - 1))
            skip_cover = self.chk_thesis_skip_cover.isChecked()

            if page_idx <= front_end_idx:
                # 1. Bölüm: Ön Sayfalar (Roma Rakamı)
                if skip_cover and page_idx == 0:
                    return None
                start_front = self.spn_thesis_front_start.value()
                offset = (page_idx - 1) if skip_cover else page_idx
                n = start_front + offset
                num_type = self.cmb_thesis_front_type.currentData() or "roman_lower"
            else:
                # 2. Bölüm: Ana Gövde (Standart Rakam)
                start_body = self.spn_thesis_body_start_num.value()
                body_offset = page_idx - (front_end_idx + 1)
                n = start_body + body_offset
                num_type = self.cmb_thesis_body_type.currentData() or "arabic"

            formatted_n = format_number_by_type(n, num_type)
            formatted_total = format_number_by_type(total_pages, num_type) if num_type in ("zero2", "zero3") else str(total_pages)
            formatted_zero_n = f"{n:02d}"
            formatted_zero_total = f"{total_pages:02d}"

            data = self.cmb_preset.currentData()
            template = self.txt_custom_format.text() if data == "custom" else (str(data) if data else "{n}")
            return template.replace("{zero_n}", formatted_zero_n).replace("{zero_total}", formatted_zero_total).replace("{n}", formatted_n).replace("{total}", formatted_total)

        if self.chk_skip_cover.isChecked() and page_idx == 0:
            return None

        if self.rad_custom_pages.isChecked():
            allowed = parse_page_range(self.txt_custom_pages.text().strip(), total_pages)
            if page_idx not in allowed:
                return None

        start = self.spn_start_num.value()
        if self.chk_skip_cover.isChecked():
            n = start + (page_idx - 1)
        else:
            n = start + page_idx

        num_type = self.cmb_num_type.currentData() or "arabic"
        formatted_n = format_number_by_type(n, num_type)
        formatted_total = format_number_by_type(total_pages, num_type) if num_type in ("zero2", "zero3") else str(total_pages)
        formatted_zero_n = f"{n:02d}"
        formatted_zero_total = f"{total_pages:02d}"

        data = self.cmb_preset.currentData()
        template = self.txt_custom_format.text() if data == "custom" else (str(data) if data else "{n}")
        return template.replace("{zero_n}", formatted_zero_n).replace("{zero_total}", formatted_zero_total).replace("{n}", formatted_n).replace("{total}", formatted_total)

    def _get_fitz_font(self) -> str:
        family = self.cmb_font_family.currentData() or "Helvetica"
        style = self.cmb_font_style.currentData() or "regular"
        if family == "Helvetica":
            if style == "bold": return "Helvetica-Bold"
            if style == "italic": return "Helvetica-Oblique"
            return "Helvetica"
        elif family == "Times-Roman":
            if style == "bold": return "Times-Bold"
            if style == "italic": return "Times-Italic"
            return "Times-Roman"
        elif family == "Courier":
            if style == "bold": return "Courier-Bold"
            if style == "italic": return "Courier-Oblique"
            return "Courier"
        return "Helvetica"

    def _get_text_rect_and_align(self, page_w: float, page_h: float, font_size: float, page_rect: Optional[fitz.Rect] = None):
        mx = float(self.spn_margin_x.value())
        my = float(self.spn_margin_y.value())
        pos = self.cmb_position.currentData() or "bottom_center"
        box_h = max(26.0, font_size * 2.8)

        if page_rect is not None:
            x0, y0, x1, y1 = page_rect.x0, page_rect.y0, page_rect.x1, page_rect.y1
        else:
            x0, y0, x1, y1 = 0.0, 0.0, page_w, page_h

        if pos == "bottom_center":
            rect = fitz.Rect(x0 + mx, y1 - my - box_h, x1 - mx, y1 - my)
            align = fitz.TEXT_ALIGN_CENTER
        elif pos == "bottom_right":
            rect = fitz.Rect(x0 + mx, y1 - my - box_h, x1 - mx, y1 - my)
            align = fitz.TEXT_ALIGN_RIGHT
        elif pos == "bottom_left":
            rect = fitz.Rect(x0 + mx, y1 - my - box_h, x1 - mx, y1 - my)
            align = fitz.TEXT_ALIGN_LEFT
        elif pos == "top_center":
            rect = fitz.Rect(x0 + mx, y0 + my, x1 - mx, y0 + my + box_h)
            align = fitz.TEXT_ALIGN_CENTER
        elif pos == "top_right":
            rect = fitz.Rect(x0 + mx, y0 + my, x1 - mx, y0 + my + box_h)
            align = fitz.TEXT_ALIGN_RIGHT
        else: # top_left
            rect = fitz.Rect(x0 + mx, y0 + my, x1 - mx, y0 + my + box_h)
            align = fitz.TEXT_ALIGN_LEFT
        return rect, align

    def _render_base_pixmap(self, target_w: int = 250, target_h: int = 350) -> QPixmap:
        doc = None
        should_close = False
        if self.is_active_tab_doc:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc is not None:
                if getattr(tab_doc, 'has_page_numbers', False):
                    if getattr(tab_doc, 'pre_numbering_bytes', None):
                        try:
                            doc = fitz.open("pdf", tab_doc.pre_numbering_bytes)
                            should_close = True
                        except Exception:
                            doc = tab_doc.doc
                    else:
                        try:
                            doc = fitz.open("pdf", tab_doc.doc.tobytes())
                            remove_page_numbering_from_doc(doc)
                            should_close = True
                        except Exception:
                            doc = tab_doc.doc
                else:
                    doc = tab_doc.doc
        elif self.loaded_file_path and os.path.isfile(self.loaded_file_path):
            try:
                doc = fitz.open(self.loaded_file_path)
                should_close = True
                if hasattr(doc, 'embfile_names') and 'aurapdf_pagenum.json' in doc.embfile_names():
                    remove_page_numbering_from_doc(doc)
            except Exception:
                doc = None

        if doc is not None and len(doc) > 0:
            try:
                order = self._get_current_order()
                if order:
                    page_pos = min(self.preview_current_page, len(order) - 1)
                    real_page_idx = order[page_pos]
                    item = self.list_widget.item(page_pos)
                    cur_rot = item.data(Qt.ItemDataRole.UserRole + 1) if item else None
                else:
                    real_page_idx = min(self.preview_current_page, len(doc) - 1)
                    cur_rot = None
                page = doc[real_page_idx]
                r = page.rect
                diff_rot = (cur_rot - page.rotation) % 360 if cur_rot is not None else 0
                if diff_rot in (90, 270):
                    eff_w, eff_h = r.height, r.width
                else:
                    eff_w, eff_h = r.width, r.height
                self.preview_real_w = eff_w
                self.preview_real_h = eff_h
                scale = min(target_w / max(1.0, eff_w), target_h / max(1.0, eff_h))
                mat = fitz.Matrix(scale, scale).prerotate(diff_rot)
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
                qpix = QPixmap.fromImage(img)
                if should_close:
                    doc.close()
                return qpix
            except Exception:
                if should_close and doc:
                    doc.close()

        self.preview_real_w = 595.0
        self.preview_real_h = 842.0
        pix = QPixmap(target_w, target_h)
        pix.fill(QColor("#ffffff"))
        p = QPainter(pix)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setPen(QPen(QColor("#d1d5db"), 1))
        p.drawRect(0, 0, target_w - 1, target_h - 1)
        p.setPen(Qt.PenStyle.NoPen)
        p.setBrush(QBrush(QColor("#374151")))
        p.drawRoundedRect(22, 28, target_w - 90, 11, 2, 2)
        p.setBrush(QBrush(QColor("#9ca3af")))
        p.drawRoundedRect(22, 48, target_w - 140, 6, 2, 2)
        p.setBrush(QBrush(QColor("#e5e7eb")))
        y = 75
        while y < target_h - 35:
            w_line = target_w - 44 if (y // 14) % 4 != 0 else target_w - 100
            p.drawRoundedRect(22, y, w_line, 5, 2, 2)
            y += 13
        p.end()
        return pix

    def update_preview(self):
        base_pix = self._render_base_pixmap(target_w=250, target_h=350)
        preview_pix = base_pix.copy()
        pw = preview_pix.width()
        ph = preview_pix.height()

        real_w = getattr(self, 'preview_real_w', 595.0)
        real_h = getattr(self, 'preview_real_h', 842.0)
        scale_x = pw / max(1.0, real_w)
        scale_y = ph / max(1.0, real_h)

        text = self._format_page_text(self.preview_current_page, self.preview_total_pages)

        self.btn_prev_page.setEnabled(self.preview_current_page > 0)
        self.btn_next_page.setEnabled(self.preview_current_page < self.preview_total_pages - 1)
        self.lbl_preview_info.setText(LanguageManager.tr("preview_page_counter", cur=self.preview_current_page + 1, total=self.preview_total_pages))

        p = QPainter(preview_pix)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setRenderHint(QPainter.RenderHint.TextAntialiasing)

        if text is not None:
            raw_fs = float(self.spn_font_size.value())
            scaled_fs = max(7, int(raw_fs * scale_x))

            style = self.cmb_font_style.currentData() or "regular"
            is_bold = (style == "bold")
            is_italic = (style == "italic")
            family = self.cmb_font_family.currentData() or "Helvetica"
            qfamily = "Arial" if family == "Helvetica" else ("Times New Roman" if family == "Times-Roman" else "Courier New")

            font = QFont(qfamily, scaled_fs)
            font.setBold(is_bold)
            font.setItalic(is_italic)
            p.setFont(font)

            fm = QFontMetrics(font)
            text_w_preview = fm.horizontalAdvance(text)
            text_w_real = text_w_preview / max(0.001, scale_x)

            pos = self.cmb_position.currentData() or "bottom_center"
            mx = float(self.spn_margin_x.value())
            my = float(self.spn_margin_y.value())
            shape_style = self.cmb_shape_style.currentData() or "none"
            custom_norm = self.custom_preview_pos if pos == "custom" else None

            badge_real, text_real, align = get_badge_layout(
                real_w, real_h, text_w_real, raw_fs, pos, mx, my, shape_style,
                custom_norm_pos=custom_norm
            )

            rect_badge_prev = QRectF(
                badge_real.x0 * scale_x,
                badge_real.y0 * scale_y,
                badge_real.width * scale_x,
                badge_real.height * scale_y
            )
            rect_text_prev = QRectF(
                text_real.x0 * scale_x,
                text_real.y0 * scale_y,
                text_real.width * scale_x,
                text_real.height * scale_y
            )

            opacity = float(self.sld_opacity.value()) / 100.0
            text_color = QColor(self.number_color)
            text_color.setAlphaF(opacity)

            if shape_style == "dark_pill":
                p.setPen(QPen(QColor("#475569"), 1))
                bg_col = QColor(30, 41, 59, int(230 * opacity))
                p.setBrush(QBrush(bg_col))
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
                text_color = QColor(248, 250, 252, int(255 * opacity))
            elif shape_style == "accent_pill":
                p.setPen(QPen(QColor("#60a5fa"), 1))
                bg_col = QColor(37, 99, 235, int(230 * opacity))
                p.setBrush(QBrush(bg_col))
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
                text_color = QColor(255, 255, 255, int(255 * opacity))
            elif shape_style == "pill":
                p.setPen(QPen(QColor("#cbd5e1"), 1))
                bg_col = QColor(255, 255, 255, int(238 * opacity))
                p.setBrush(QBrush(bg_col))
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
            elif shape_style == "glassmorphism":
                p.setPen(QPen(QColor(220, 230, 245, int(220 * opacity)), 1.2))
                bg_col = QColor(245, 248, 255, int(160 * opacity))
                p.setBrush(QBrush(bg_col))
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
            elif shape_style == "modern_bar":
                p.setPen(QPen(QColor("#2563eb"), max(2.0, 2.4 * scale_x), Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap))
                p.drawLine(QPointF(rect_badge_prev.left() + 2.0 * scale_x, rect_badge_prev.top() + 2.5 * scale_y),
                           QPointF(rect_badge_prev.left() + 2.0 * scale_x, rect_badge_prev.bottom() - 2.5 * scale_y))
            elif shape_style == "bracket_modern":
                bx0 = rect_badge_prev.left()
                bx1 = rect_badge_prev.right()
                by0 = rect_badge_prev.top() + 2.0 * scale_y
                by1 = rect_badge_prev.bottom() - 2.0 * scale_y
                arm = min(5.0, (bx1 - bx0) * 0.12)
                p.setPen(QPen(QColor("#475569"), max(1.2, 1.4 * scale_x), Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap, Qt.PenJoinStyle.RoundJoin))
                p.drawLine(QPointF(bx0 + arm, by0), QPointF(bx0, by0))
                p.drawLine(QPointF(bx0, by0), QPointF(bx0, by1))
                p.drawLine(QPointF(bx0, by1), QPointF(bx0 + arm, by1))
                p.drawLine(QPointF(bx1 - arm, by0), QPointF(bx1, by0))
                p.drawLine(QPointF(bx1, by0), QPointF(bx1, by1))
                p.drawLine(QPointF(bx1, by1), QPointF(bx1 - arm, by1))
            elif shape_style == "gradient_pill":
                p.setPen(QPen(QColor("#8b5cf6"), 1))
                bg_col = QColor(109, 40, 217, int(235 * opacity))
                p.setBrush(QBrush(bg_col))
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
                text_color = QColor(255, 255, 255, int(255 * opacity))
            elif shape_style == "accent_dot":
                p.setPen(Qt.PenStyle.NoPen)
                p.setBrush(QBrush(QColor("#2563eb")))
                dot_rad = max(2.5, rect_badge_prev.height() * 0.18)
                p.drawEllipse(QPointF(rect_badge_prev.left() + dot_rad + 3.0 * scale_x, rect_badge_prev.center().y()), dot_rad, dot_rad)
            elif shape_style == "outline_pill":
                p.setPen(QPen(QColor("#64748b"), 1.2))
                p.setBrush(Qt.BrushStyle.NoBrush)
                r = rect_badge_prev.height() / 2.0
                p.drawRoundedRect(rect_badge_prev, r, r)
            elif shape_style == "box":
                p.setPen(QPen(QColor("#94a3b8"), 1))
                bg_col = QColor(255, 255, 255, int(238 * opacity))
                p.setBrush(QBrush(bg_col))
                p.drawRect(rect_badge_prev)
            elif shape_style == "circle":
                p.setPen(QPen(QColor("#94a3b8"), 1))
                bg_col = QColor(255, 255, 255, int(238 * opacity))
                p.setBrush(QBrush(bg_col))
                p.drawEllipse(rect_badge_prev)
            elif shape_style == "underline":
                p.setPen(QPen(QColor("#2563eb"), max(1.5, 2.0 * scale_x)))
                p.drawLine(QPointF(rect_badge_prev.left(), rect_badge_prev.bottom()), QPointF(rect_badge_prev.right(), rect_badge_prev.bottom()))

            p.setPen(QPen(text_color))
            if shape_style != "none":
                p.drawText(rect_badge_prev, Qt.AlignmentFlag.AlignCenter, text)
            else:
                qalign = Qt.AlignmentFlag.AlignVCenter
                if align == fitz.TEXT_ALIGN_CENTER:
                    qalign |= Qt.AlignmentFlag.AlignHCenter
                elif align == fitz.TEXT_ALIGN_RIGHT:
                    qalign |= Qt.AlignmentFlag.AlignRight
                else:
                    qalign |= Qt.AlignmentFlag.AlignLeft
                p.drawText(rect_text_prev, int(qalign), text)
        else:
            p.setPen(Qt.PenStyle.NoPen)
            p.setBrush(QBrush(QColor(0, 0, 0, 140)))
            p.drawRoundedRect(QRectF(15, ph - 32, pw - 30, 22), 4, 4)
            p.setPen(QPen(QColor("#ffffff")))
            p.setFont(QFont("Arial", 8, QFont.Weight.Bold))
            p.drawText(QRectF(15, ph - 32, pw - 30, 22), Qt.AlignmentFlag.AlignCenter, LanguageManager.tr("page_number_skip_cover"))

        p.end()
        self.lbl_preview.setPixmap(preview_pix)

    def execute_apply_reorder(self):
        tab_doc = self.main_window.get_current_tab_doc()
        if not tab_doc or tab_doc.doc is None:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("reorder_warn_no_doc"))
            return

        order = self._get_current_order()
        rotations = [self.list_widget.item(i).data(Qt.ItemDataRole.UserRole + 1) for i in range(self.list_widget.count())]
        tab_doc.reorder_pages(order, rotations=rotations, push_undo=True)
        if getattr(tab_doc, 'pre_numbering_bytes', None):
            try:
                pre_doc = fitz.open("pdf", tab_doc.pre_numbering_bytes)
                if order != list(range(len(pre_doc))):
                    pre_doc.select(order)
                for i, rot in enumerate(rotations):
                    if rot is not None:
                        pre_doc[i].set_rotation(rot % 360)
                tab_doc.pre_numbering_bytes = pre_doc.tobytes()
                pre_doc.close()
            except Exception:
                pass
        QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("reorder_success_apply"))
        self.accept()

    def execute_save_as_reorder(self):
        if not self.loaded_file_path or self.list_widget.count() == 0:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("reorder_warn_no_doc"))
            return

        order = self._get_current_order()
        rotations = [self.list_widget.item(i).data(Qt.ItemDataRole.UserRole + 1) for i in range(self.list_widget.count())]
        base_name = os.path.basename(self.loaded_file_path)
        save_path, _ = QFileDialog.getSaveFileName(
            self, LanguageManager.tr("reorder_save_title"), f"siralanmis_{base_name}", "PDF (*.pdf)"
        )
        if not save_path:
            return

        try:
            if self.is_active_tab_doc:
                tab_doc = self.main_window.get_current_tab_doc()
                doc = fitz.open("pdf", tab_doc.doc.tobytes())
                tab_doc.export_drawings_to_pdf(doc)
            else:
                doc = fitz.open(self.loaded_file_path)
            doc.select(order)
            for i, rot in enumerate(rotations):
                if rot is not None:
                    doc[i].set_rotation(rot % 360)
            doc.save(save_path, garbage=4, deflate=True)
            doc.close()

            reply = QMessageBox.question(
                self,
                LanguageManager.tr("success"),
                f"{LanguageManager.tr('reorder_success_save', path=save_path)}\n\n{LanguageManager.tr('reorder_ask_open')}",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes
            )
            if reply == QMessageBox.StandardButton.Yes:
                self.main_window.add_pdf_tab(save_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def execute_apply_numbering(self):
        tab_doc = self.main_window.get_current_tab_doc()
        if not tab_doc or tab_doc.doc is None:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("page_number_warn_no_doc"))
            return

        try:
            order = self._get_current_order()
            rotations = [self.list_widget.item(i).data(Qt.ItemDataRole.UserRole + 1) for i in range(self.list_widget.count())]
            old_bytes = tab_doc.doc.tobytes()
            old_drawings = [dict(d) for d in tab_doc.drawings]
            old_page = tab_doc.current_page_idx
            old_has_page_numbers = getattr(tab_doc, 'has_page_numbers', False)
            old_settings = getattr(tab_doc, 'page_numbering_settings', None)
            old_pre_bytes = getattr(tab_doc, 'pre_numbering_bytes', None)

            # Determine base document
            if old_has_page_numbers and old_pre_bytes:
                try:
                    base_doc = fitz.open("pdf", old_pre_bytes)
                except Exception:
                    base_doc = fitz.open("pdf", old_bytes)
                    remove_page_numbering_from_doc(base_doc)
            else:
                base_doc = fitz.open("pdf", old_bytes)
                if old_has_page_numbers:
                    remove_page_numbering_from_doc(base_doc)

            # Reorder & rotate base_doc to match current dialog state
            if order and order != list(range(len(base_doc))):
                base_doc.select(order)
            for i, rot in enumerate(rotations):
                if rot is not None:
                    base_doc[i].set_rotation(rot % 360)

            clean_reordered_bytes = base_doc.tobytes()

            settings = self._get_current_numbering_settings()
            shape_style = settings["shape_style"]
            fs = float(settings["font_size"])
            opacity = float(settings["opacity"]) / 100.0
            pos = settings["position"]
            mx = float(settings["margin_x"])
            my = float(settings["margin_y"])
            custom_norm = self.custom_preview_pos if pos == "custom" else None

            apply_numbering_to_doc(
                base_doc,
                self._format_page_text,
                self._get_fitz_font,
                fs, opacity, self.number_color, pos, mx, my, shape_style,
                custom_norm_pos=custom_norm,
                settings_dict=settings
            )

            new_bytes = base_doc.tobytes()
            base_doc.close()

            tab_doc.doc.close()
            tab_doc.doc = fitz.open("pdf", new_bytes)
            if hasattr(tab_doc, 'hide_managed_annotations'):
                tab_doc.hide_managed_annotations()
            tab_doc.has_page_numbers = True
            tab_doc.page_numbering_settings = settings
            tab_doc.pre_numbering_bytes = clean_reordered_bytes

            tab_doc.push_undo_action({
                "type": "reorder_pages",
                "old_bytes": old_bytes,
                "old_drawings": old_drawings,
                "old_page": old_page,
                "new_bytes": new_bytes,
                "new_drawings": [dict(d) for d in tab_doc.drawings],
                "new_page": old_page,
                "new_order": list(range(len(tab_doc.doc))),
                "old_has_page_numbers": old_has_page_numbers,
                "old_page_numbering_settings": old_settings,
                "old_pre_numbering_bytes": old_pre_bytes,
                "new_has_page_numbers": True,
                "new_page_numbering_settings": settings,
                "new_pre_numbering_bytes": clean_reordered_bytes
            })

            tab_doc.render_pages_init()
            for pinfo in tab_doc.page_positions.values():
                pinfo['rendered_zoom'] = 0.0
            tab_doc.rerender_visible_pages()
            tab_doc.go_to_page(old_page)
            tab_doc.scene.update()
            tab_doc.view.viewport().update()

            if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                self.main_window.thumbnails_dock.load_thumbnails_lazy()
                self.main_window.thumbnails_dock.select_page(old_page)

            tab_doc.set_modified(True)

            QTimer.singleShot(80, lambda: tab_doc.rerender_visible_pages() if tab_doc and tab_doc.doc else None)

            msg = LanguageManager.tr("page_number_success_update") if old_has_page_numbers else LanguageManager.tr("page_number_success_apply")
            QMessageBox.information(self, LanguageManager.tr("success"), msg)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("page_number_err_process", error=str(e)))

    def execute_remove_numbering(self):
        tab_doc = self.main_window.get_current_tab_doc()
        if not tab_doc or tab_doc.doc is None:
            return

        reply = QMessageBox.question(
            self,
            LanguageManager.tr("page_number_btn_remove"),
            LanguageManager.tr("page_number_confirm_remove"),
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        try:
            old_bytes = tab_doc.doc.tobytes()
            old_drawings = [dict(d) for d in tab_doc.drawings]
            old_page = tab_doc.current_page_idx
            old_has_page_numbers = getattr(tab_doc, 'has_page_numbers', False)
            old_settings = getattr(tab_doc, 'page_numbering_settings', None)
            old_pre_bytes = getattr(tab_doc, 'pre_numbering_bytes', None)

            if old_pre_bytes:
                clean_bytes = old_pre_bytes
            else:
                temp_doc = fitz.open("pdf", old_bytes)
                remove_page_numbering_from_doc(temp_doc)
                clean_bytes = temp_doc.tobytes()
                temp_doc.close()

            tab_doc.doc.close()
            tab_doc.doc = fitz.open("pdf", clean_bytes)
            if hasattr(tab_doc, 'hide_managed_annotations'):
                tab_doc.hide_managed_annotations()
            tab_doc.has_page_numbers = False
            tab_doc.page_numbering_settings = None
            tab_doc.pre_numbering_bytes = None

            tab_doc.push_undo_action({
                "type": "reorder_pages",
                "old_bytes": old_bytes,
                "old_drawings": old_drawings,
                "old_page": old_page,
                "new_bytes": clean_bytes,
                "new_drawings": [dict(d) for d in tab_doc.drawings],
                "new_page": old_page,
                "new_order": list(range(len(tab_doc.doc))),
                "old_has_page_numbers": old_has_page_numbers,
                "old_page_numbering_settings": old_settings,
                "old_pre_numbering_bytes": old_pre_bytes,
                "new_has_page_numbers": False,
                "new_page_numbering_settings": None,
                "new_pre_numbering_bytes": None
            })

            tab_doc.render_pages_init()
            for pinfo in tab_doc.page_positions.values():
                pinfo['rendered_zoom'] = 0.0
            tab_doc.rerender_visible_pages()
            tab_doc.go_to_page(old_page)
            tab_doc.scene.update()
            tab_doc.view.viewport().update()

            if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                self.main_window.thumbnails_dock.load_thumbnails_lazy()
                self.main_window.thumbnails_dock.select_page(old_page)

            tab_doc.set_modified(True)

            QTimer.singleShot(80, lambda: tab_doc.rerender_visible_pages() if tab_doc and tab_doc.doc else None)

            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("page_number_success_remove"))
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("page_number_err_process", error=str(e)))

    def execute_save_as_numbering(self):
        src = self.txt_file.text().strip()
        if not self.is_active_tab_doc and (not src or not os.path.isfile(src)):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("page_number_warn_no_doc"))
            return

        base_name = os.path.basename(self.loaded_file_path or "belge.pdf")
        name_part, _ = os.path.splitext(base_name)
        save_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("page_number_save_title"), f"{name_part}_numarali.pdf", "PDF (*.pdf)")
        if not save_path:
            return

        try:
            if self.is_active_tab_doc:
                tab_doc = self.main_window.get_current_tab_doc()
                if getattr(tab_doc, 'has_page_numbers', False) and getattr(tab_doc, 'pre_numbering_bytes', None):
                    doc = fitz.open("pdf", tab_doc.pre_numbering_bytes)
                else:
                    doc = fitz.open("pdf", tab_doc.doc.tobytes())
                    if getattr(tab_doc, 'has_page_numbers', False):
                        remove_page_numbering_from_doc(doc)
                tab_doc.export_drawings_to_pdf(doc)
            else:
                doc = fitz.open(self.loaded_file_path)
                remove_page_numbering_from_doc(doc)

            order = self._get_current_order()
            rotations = [self.list_widget.item(i).data(Qt.ItemDataRole.UserRole + 1) for i in range(self.list_widget.count())]
            if order and order != list(range(len(doc))):
                doc.select(order)
            for i, rot in enumerate(rotations):
                if rot is not None:
                    doc[i].set_rotation(rot % 360)

            settings = self._get_current_numbering_settings()
            shape_style = settings["shape_style"]
            fs = float(settings["font_size"])
            opacity = float(settings["opacity"]) / 100.0
            pos = settings["position"]
            mx = float(settings["margin_x"])
            my = float(settings["margin_y"])
            custom_norm = self.custom_preview_pos if pos == "custom" else None

            apply_numbering_to_doc(
                doc,
                self._format_page_text,
                self._get_fitz_font,
                fs, opacity, self.number_color, pos, mx, my, shape_style,
                custom_norm_pos=custom_norm,
                settings_dict=settings
            )

            doc.save(save_path, garbage=4, deflate=True)
            doc.close()

            reply = QMessageBox.question(
                self,
                LanguageManager.tr("success"),
                f"{LanguageManager.tr('page_number_success_save', path=save_path)}\n\n{LanguageManager.tr('page_number_ask_open')}",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.Yes
            )
            if reply == QMessageBox.StandardButton.Yes:
                self.main_window.add_pdf_tab(save_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("page_number_err_process", error=str(e)))


class PDFReorderPagesDialog(PDFOrganizePagesDialog):
    def __init__(self, main_window):
        super().__init__(main_window, initial_tab=0)


class PDFPageNumberDialog(PDFOrganizePagesDialog):
    def __init__(self, main_window):
        super().__init__(main_window, initial_tab=1)

def extract_pdf_hint(file_path: str) -> Optional[str]:
    """PDF dosyasının sonundaki gömülü parola ipucunu ayıklar (varsa)."""
    try:
        if not file_path or not os.path.isfile(file_path):
            return None
        size = os.path.getsize(file_path)
        with open(file_path, "rb") as f:
            f.seek(max(0, size - 8192))
            tail = f.read()
        m = re.search(rb"%AuraPDF_Hint:([A-Za-z0-9+/=]+)", tail)
        if m:
            b64 = m.group(1).decode("ascii", errors="ignore")
            return base64.b64decode(b64).decode("utf-8", errors="ignore")
    except Exception:
        pass
    return None


def parse_page_range(text: str, total_pages: int) -> List[int]:
    """Kullanıcının girdiği '1-5, 8, 11' gibi sayfa metnini 0-indeksli liste olarak döner."""
    pages = set()
    for part in text.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            tokens = part.split("-")
            try:
                start = max(1, int(tokens[0].strip()))
                end = min(total_pages, int(tokens[1].strip()))
                for p in range(start, end + 1):
                    pages.add(p - 1)
            except ValueError:
                continue
        else:
            try:
                p = int(part)
                if 1 <= p <= total_pages:
                    pages.add(p - 1)
            except ValueError:
                continue
    return sorted(list(pages))


# --- POWERPOINT (.PPTX) VE EXCEL (.XLSX) DÖNÜŞTÜRÜCÜ DİYALOGLARI ---

def map_pdf_rect_to_visual(r: fitz.Rect, rot: int, w_unrot: float, h_unrot: float) -> fitz.Rect:
    """PDF sayfasındaki döndürülmemiş koordinatı okuyucunun görsel oryantasyonuna dönüştürür."""
    rot = rot % 360
    if rot == 90:
        return fitz.Rect(h_unrot - r.y1, r.x0, h_unrot - r.y0, r.x1)
    elif rot == 180:
        return fitz.Rect(w_unrot - r.x1, h_unrot - r.y1, w_unrot - r.x0, h_unrot - r.y0)
    elif rot == 270:
        return fitz.Rect(r.y0, w_unrot - r.x1, r.y1, w_unrot - r.x0)
    return fitz.Rect(r.x0, r.y0, r.x1, r.y1)


class PPTXConversionThread(QThread):
    progress_signal = pyqtSignal(int, int, str)
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, input_pdf: str, output_path: str, mode: str = "editable", pages: Optional[List[int]] = None, temp_file_to_cleanup: Optional[str] = None):
        super().__init__()
        self.input_pdf = input_pdf
        self.output_path = output_path
        self.mode = mode
        self.pages = pages
        self.temp_file_to_cleanup = temp_file_to_cleanup
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            from pptx import Presentation
            from pptx.util import Pt, Inches
        except ImportError:
            self.finished_signal.emit(False, "python-pptx kütüphanesi sistemde yüklü değil.")
            return
        try:
            doc = fitz.open(self.input_pdf)
            prs = Presentation()
            if self.pages is None:
                self.pages = list(range(len(doc)))
            total = len(self.pages)
            if not self.pages:
                doc.close()
                self.finished_signal.emit(False, "Dönüştürülecek sayfa bulunamadı.")
                return

            first_page = doc[self.pages[0]]
            prs.slide_width = Pt(first_page.rect.width)
            prs.slide_height = Pt(first_page.rect.height)

            for idx, pno in enumerate(self.pages):
                if self._cancelled:
                    doc.close()
                    self.finished_signal.emit(False, "İşlem kullanıcı tarafından iptal edildi.")
                    return

                page = doc[pno]
                rect = page.rect
                rot = page.rotation % 360
                unrot_w = rect.height if rot in (90, 270) else rect.width
                unrot_h = rect.width if rot in (90, 270) else rect.height

                blank_layout = prs.slide_layouts[6]
                slide = prs.slides.add_slide(blank_layout)

                if self.mode == "highres":
                    pix = page.get_pixmap(dpi=150)
                    img_bytes = pix.tobytes("png")
                    slide.shapes.add_picture(io.BytesIO(img_bytes), 0, 0, Pt(rect.width), Pt(rect.height))

                for block in page.get_text("blocks"):
                    x0, y0, x1, y1, text, bno, btype = block
                    if btype == 0 and text.strip():
                        r_orig = fitz.Rect(x0, y0, x1, y1)
                        r_vis = map_pdf_rect_to_visual(r_orig, rot, unrot_w, unrot_h) if rot != 0 else r_orig
                        bw = max(20.0, r_vis.width)
                        bh = max(14.0, r_vis.height)
                        tb = slide.shapes.add_textbox(Pt(r_vis.x0), Pt(r_vis.y0), Pt(bw), Pt(bh))
                        tf = tb.text_frame
                        tf.word_wrap = True
                        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
                        lines = [l.strip() for l in text.split("\n") if l.strip()]
                        for l_idx, line_str in enumerate(lines):
                            p_para = tf.paragraphs[0] if l_idx == 0 else tf.add_paragraph()
                            p_para.text = line_str
                            p_para.font.size = Pt(11)

                if self.mode == "editable":
                    for img_info in page.get_images(full=True):
                        xref = img_info[0]
                        try:
                            base_image = doc.extract_image(xref)
                            image_bytes = base_image["image"]
                            rects = page.get_image_rects(xref)
                            if rects:
                                r_orig = rects[0]
                                r_vis = map_pdf_rect_to_visual(r_orig, rot, unrot_w, unrot_h) if rot != 0 else r_orig
                                if rot != 0:
                                    try:
                                        from PIL import Image
                                        pil_im = Image.open(io.BytesIO(image_bytes))
                                        pil_im = pil_im.rotate((360 - rot) % 360, expand=True)
                                        buf = io.BytesIO()
                                        pil_im.save(buf, format="PNG")
                                        image_bytes = buf.getvalue()
                                    except Exception:
                                        pass
                                slide.shapes.add_picture(io.BytesIO(image_bytes), Pt(r_vis.x0), Pt(r_vis.y0), Pt(r_vis.width), Pt(r_vis.height))
                        except Exception:
                            pass

                self.progress_signal.emit(idx + 1, total, f"Sayfa {pno + 1}/{len(doc)} slayta dönüştürülüyor...")

            prs.save(self.output_path)
            doc.close()
            self.finished_signal.emit(True, f"PowerPoint sunumu başarıyla oluşturuldu:\n{self.output_path}")
        except Exception as e:
            self.finished_signal.emit(False, f"PowerPoint dönüştürme hatası:\n{e}")
        finally:
            if self.temp_file_to_cleanup and os.path.isfile(self.temp_file_to_cleanup):
                try:
                    os.remove(self.temp_file_to_cleanup)
                except Exception:
                    pass


class PDFToPowerPointDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.thread = None

        self.setWindowTitle(LanguageManager.tr("pptx_title"))
        self.resize(600, 420)
        self.setMinimumSize(480, 320)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        layout.addWidget(QLabel(LanguageManager.tr("pptx_src_label")))
        h_src = QHBoxLayout()
        self.txt_source = QLineEdit()
        btn_browse_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_src.clicked.connect(self.browse_source)
        h_src.addWidget(self.txt_source)
        h_src.addWidget(btn_browse_src)
        layout.addLayout(h_src)

        layout.addWidget(QLabel(LanguageManager.tr("pptx_dest_label")))
        h_dest = QHBoxLayout()
        self.txt_dest = QLineEdit()
        btn_browse_dest = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_dest.clicked.connect(self.browse_dest)
        h_dest.addWidget(self.txt_dest)
        h_dest.addWidget(btn_browse_dest)
        layout.addLayout(h_dest)

        grp_mode = QGroupBox(LanguageManager.tr("pptx_mode_label"))
        v_mode = QVBoxLayout(grp_mode)
        self.rb_mode_editable = QRadioButton(LanguageManager.tr("pptx_mode_editable"))
        self.rb_mode_editable.setChecked(True)
        self.rb_mode_highres = QRadioButton(LanguageManager.tr("pptx_mode_highres"))
        v_mode.addWidget(self.rb_mode_editable)
        v_mode.addWidget(self.rb_mode_highres)
        layout.addWidget(grp_mode)

        grp_pages = QGroupBox(LanguageManager.tr("print_grp_pages"))
        v_pages = QVBoxLayout(grp_pages)
        self.rb_all = QRadioButton(LanguageManager.tr("word_pages_all"))
        self.rb_all.setChecked(True)
        self.rb_range = QRadioButton(LanguageManager.tr("word_pages_range"))
        h_range = QHBoxLayout()
        h_range.addWidget(self.rb_range)
        self.txt_range = QLineEdit()
        self.txt_range.setPlaceholderText("Örn: 1-5, 8")
        self.txt_range.setEnabled(False)
        self.rb_range.toggled.connect(self.txt_range.setEnabled)
        h_range.addWidget(self.txt_range)
        v_pages.addWidget(self.rb_all)
        v_pages.addLayout(h_range)
        layout.addWidget(grp_pages)

        lbl_engine = QLabel(LanguageManager.tr("pptx_engine_badge"))
        lbl_engine.setStyleSheet(" font-weight: bold;")
        layout.addWidget(lbl_engine)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        self.btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        self.btn_cancel.clicked.connect(self.reject)
        self.btn_convert = QPushButton(LanguageManager.tr("pptx_btn_run"))
        self.btn_convert.setObjectName("PrimaryBtn")
        self.btn_convert.clicked.connect(self.execute_conversion)
        h_btn.addWidget(self.btn_cancel)
        h_btn.addWidget(self.btn_convert)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.txt_source.setText(tab_doc.file_path)
            base = os.path.splitext(tab_doc.file_path)[0]
            self.txt_dest.setText(f"{base}.pptx")

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_source.setText(path)
            base = os.path.splitext(path)[0]
            self.txt_dest.setText(f"{base}.pptx")

    def browse_dest(self):
        cur = self.txt_dest.text()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), cur, "PowerPoint (*.pptx)")
        if path:
            self.txt_dest.setText(path)

    def execute_conversion(self):
        src = self.txt_source.text().strip()
        dest = self.txt_dest.text().strip()
        if not src or not os.path.isfile(src):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_no_file"))
            return
        if not dest:
            dest, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src.replace(".pdf", ".pptx"), "PowerPoint (*.pptx)")
            if not dest:
                return
            self.txt_dest.setText(dest)

        try:
            doc = fitz.open(src)
            total = len(doc)
            doc.close()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), f"PDF açılamadı:\n{e}")
            return

        pages = list(range(total)) if self.rb_all.isChecked() else parse_page_range(self.txt_range.text().strip(), total)
        if not pages:
            QMessageBox.warning(self, LanguageManager.tr("warning"), "Lütfen geçerli bir sayfa aralığı belirtin.")
            return

        mode = "editable" if self.rb_mode_editable.isChecked() else "highres"
        self.btn_convert.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, len(pages))
        self.progress_bar.setValue(0)

        temp_input_pdf = None
        if hasattr(self, 'main_window') and self.main_window:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc and getattr(tab_doc, 'file_path', None):
                try:
                    if os.path.abspath(tab_doc.file_path) == os.path.abspath(src):
                        import tempfile
                        with tempfile.NamedTemporaryFile(suffix="_pptx_active.pdf", delete=False) as tf:
                            temp_input_pdf = tf.name
                        temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                        tab_doc.export_drawings_to_pdf(temp_doc)
                        temp_doc.save(temp_input_pdf)
                        temp_doc.close()
                        src = temp_input_pdf
                except Exception:
                    temp_input_pdf = None

        self.thread = PPTXConversionThread(src, dest, mode, pages, temp_file_to_cleanup=temp_input_pdf)
        self.thread.progress_signal.connect(lambda cur, tot, msg: (self.progress_bar.setValue(cur), self.progress_bar.setFormat(f"{cur}/{tot} - {msg}")))
        self.thread.finished_signal.connect(self.on_finished)
        self.thread.start()

    def on_finished(self, success, message):
        self.btn_convert.setEnabled(True)
        self.progress_bar.setVisible(False)
        if success:
            dest = self.txt_dest.text().strip()
            reply = QMessageBox.question(self, LanguageManager.tr("success"),
                                         f"{LanguageManager.tr('pptx_success', path=dest)}\n\n{LanguageManager.tr('pptx_open_now')}",
                                         QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    os.startfile(dest)
                except Exception:
                    pass
            self.accept()
        else:
            QMessageBox.critical(self, LanguageManager.tr("error"), message)

    def closeEvent(self, event):
        if self.thread and self.thread.isRunning():
            self.thread.cancel()
            self.thread.wait(1000)
        super().closeEvent(event)


class ExcelConversionThread(QThread):
    progress_signal = pyqtSignal(int, int, str)
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, input_pdf: str, output_path: str, sheet_mode: str = "per_page", detect_tables: bool = True, pages: Optional[List[int]] = None, temp_file_to_cleanup: Optional[str] = None):
        super().__init__()
        self.input_pdf = input_pdf
        self.output_path = output_path
        self.sheet_mode = sheet_mode
        self.detect_tables = detect_tables
        self.pages = pages
        self.temp_file_to_cleanup = temp_file_to_cleanup
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            self.finished_signal.emit(False, "openpyxl kütüphanesi sistemde yüklü değil.")
            return
        try:
            doc = fitz.open(self.input_pdf)
            wb = openpyxl.Workbook()
            default_sheet = wb.active
            wb.remove(default_sheet)

            if self.pages is None:
                self.pages = list(range(len(doc)))
            total = len(self.pages)
            if self.sheet_mode == "single":
                ws = wb.create_sheet(title="PDF Verileri")
                current_row = 1

            for idx, pno in enumerate(self.pages):
                if self._cancelled:
                    doc.close()
                    self.finished_signal.emit(False, "İşlem kullanıcı tarafından iptal edildi.")
                    return

                page = doc[pno]
                rect = page.rect
                rot = page.rotation % 360
                unrot_w = rect.height if rot in (90, 270) else rect.width
                unrot_h = rect.width if rot in (90, 270) else rect.height

                if self.sheet_mode == "per_page":
                    ws = wb.create_sheet(title=f"Sayfa {pno + 1}")
                    current_row = 1
                else:
                    if current_row > 1:
                        current_row += 1
                    cell = ws.cell(row=current_row, column=1, value=f"=== SAYFA {pno + 1} ===")
                    cell.font = Font(bold=True, color="107C41")
                    current_row += 1

                tabs = page.find_tables() if self.detect_tables else None
                has_tables = tabs and len(tabs.tables) > 0

                if has_tables:
                    for tab in tabs:
                        extracted = tab.extract()
                        for r_data in extracted:
                            for c_idx, val in enumerate(r_data, start=1):
                                cell_val = str(val).strip() if val is not None else ""
                                cleaned = cell_val.replace(" ", "").replace(",", ".")
                                try:
                                    if ("." in cleaned or "," in cell_val) and cleaned.replace(".", "", 1).isdigit():
                                        ws.cell(row=current_row, column=c_idx, value=float(cleaned))
                                    elif cleaned.isdigit():
                                        ws.cell(row=current_row, column=c_idx, value=int(cleaned))
                                    else:
                                        ws.cell(row=current_row, column=c_idx, value=cell_val)
                                except Exception:
                                    ws.cell(row=current_row, column=c_idx, value=cell_val)
                            current_row += 1
                        current_row += 1
                else:
                    blocks = page.get_text("blocks")
                    mapped_blocks = []
                    for b in blocks:
                        if rot != 0:
                            r_orig = fitz.Rect(b[:4])
                            r_vis = map_pdf_rect_to_visual(r_orig, rot, unrot_w, unrot_h)
                            mapped_blocks.append((r_vis.x0, r_vis.y0, r_vis.x1, r_vis.y1, b[4], b[5], b[6]))
                        else:
                            mapped_blocks.append(b)
                    mapped_blocks.sort(key=lambda b: (round(b[1] / 12) * 12, b[0]))
                    for b in mapped_blocks:
                        text = b[4].strip()
                        if text:
                            for l in text.split("\n"):
                                if "\t" in l:
                                    cols = l.split("\t")
                                    for c_idx, val in enumerate(cols, start=1):
                                        ws.cell(row=current_row, column=c_idx, value=val.strip())
                                elif "   " in l:
                                    cols = [col.strip() for col in re.split(r'\s{2,}', l) if col.strip()]
                                    for c_idx, val in enumerate(cols, start=1):
                                        ws.cell(row=current_row, column=c_idx, value=val)
                                else:
                                    ws.cell(row=current_row, column=1, value=l.strip())
                                current_row += 1

                self.progress_signal.emit(idx + 1, total, f"Sayfa {pno + 1}/{len(doc)} Excel'e aktarılıyor...")

            wb.save(self.output_path)
            doc.close()
            self.finished_signal.emit(True, f"Excel çalışma kitabı başarıyla oluşturuldu:\n{self.output_path}")
        except Exception as e:
            self.finished_signal.emit(False, f"Excel dönüştürme hatası:\n{e}")
        finally:
            if self.temp_file_to_cleanup and os.path.isfile(self.temp_file_to_cleanup):
                try:
                    os.remove(self.temp_file_to_cleanup)
                except Exception:
                    pass


class PDFToExcelDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.thread = None

        self.setWindowTitle(LanguageManager.tr("excel_title"))
        self.resize(600, 430)
        self.setMinimumSize(480, 340)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        layout.addWidget(QLabel(LanguageManager.tr("excel_src_label")))
        h_src = QHBoxLayout()
        self.txt_source = QLineEdit()
        btn_browse_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_src.clicked.connect(self.browse_source)
        h_src.addWidget(self.txt_source)
        h_src.addWidget(btn_browse_src)
        layout.addLayout(h_src)

        layout.addWidget(QLabel(LanguageManager.tr("excel_dest_label")))
        h_dest = QHBoxLayout()
        self.txt_dest = QLineEdit()
        btn_browse_dest = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_browse_dest.clicked.connect(self.browse_dest)
        h_dest.addWidget(self.txt_dest)
        h_dest.addWidget(btn_browse_dest)
        layout.addLayout(h_dest)

        grp_sheet = QGroupBox(LanguageManager.tr("excel_sheet_mode"))
        v_sheet = QVBoxLayout(grp_sheet)
        self.rb_sheet_per_page = QRadioButton(LanguageManager.tr("excel_sheet_per_page"))
        self.rb_sheet_per_page.setChecked(True)
        self.rb_sheet_single = QRadioButton(LanguageManager.tr("excel_sheet_single"))
        self.chk_detect_tables = QCheckBox(LanguageManager.tr("excel_detect_tables"))
        self.chk_detect_tables.setChecked(True)
        v_sheet.addWidget(self.rb_sheet_per_page)
        v_sheet.addWidget(self.rb_sheet_single)
        v_sheet.addWidget(self.chk_detect_tables)
        layout.addWidget(grp_sheet)

        grp_pages = QGroupBox(LanguageManager.tr("print_grp_pages"))
        v_pages = QVBoxLayout(grp_pages)
        self.rb_all = QRadioButton(LanguageManager.tr("word_pages_all"))
        self.rb_all.setChecked(True)
        self.rb_range = QRadioButton(LanguageManager.tr("word_pages_range"))
        h_range = QHBoxLayout()
        h_range.addWidget(self.rb_range)
        self.txt_range = QLineEdit()
        self.txt_range.setPlaceholderText("Örn: 1-5, 8")
        self.txt_range.setEnabled(False)
        self.rb_range.toggled.connect(self.txt_range.setEnabled)
        h_range.addWidget(self.txt_range)
        v_pages.addWidget(self.rb_all)
        v_pages.addLayout(h_range)
        layout.addWidget(grp_pages)

        lbl_engine = QLabel(LanguageManager.tr("excel_engine_badge"))
        lbl_engine.setStyleSheet(" font-weight: bold;")
        layout.addWidget(lbl_engine)

        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        layout.addWidget(self.progress_bar)

        h_btn = QHBoxLayout()
        h_btn.addStretch()
        self.btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        self.btn_cancel.clicked.connect(self.reject)
        self.btn_convert = QPushButton(LanguageManager.tr("excel_btn_run"))
        self.btn_convert.setObjectName("PrimaryBtn")
        self.btn_convert.clicked.connect(self.execute_conversion)
        h_btn.addWidget(self.btn_cancel)
        h_btn.addWidget(self.btn_convert)
        layout.addLayout(h_btn)

        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            self.txt_source.setText(tab_doc.file_path)
            base = os.path.splitext(tab_doc.file_path)[0]
            self.txt_dest.setText(f"{base}.xlsx")

    def browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_source.setText(path)
            base = os.path.splitext(path)[0]
            self.txt_dest.setText(f"{base}.xlsx")

    def browse_dest(self):
        cur = self.txt_dest.text()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), cur, "Excel (*.xlsx)")
        if path:
            self.txt_dest.setText(path)

    def execute_conversion(self):
        src = self.txt_source.text().strip()
        dest = self.txt_dest.text().strip()
        if not src or not os.path.isfile(src):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("split_warn_no_file"))
            return
        if not dest:
            dest, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src.replace(".pdf", ".xlsx"), "Excel (*.xlsx)")
            if not dest:
                return
            self.txt_dest.setText(dest)

        try:
            doc = fitz.open(src)
            total = len(doc)
            doc.close()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), f"PDF açılamadı:\n{e}")
            return

        pages = list(range(total)) if self.rb_all.isChecked() else parse_page_range(self.txt_range.text().strip(), total)
        if not pages:
            QMessageBox.warning(self, LanguageManager.tr("warning"), "Lütfen geçerli bir sayfa aralığı belirtin.")
            return

        sheet_mode = "per_page" if self.rb_sheet_per_page.isChecked() else "single"
        detect_tables = self.chk_detect_tables.isChecked()

        self.btn_convert.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, len(pages))
        self.progress_bar.setValue(0)

        temp_input_pdf = None
        if hasattr(self, 'main_window') and self.main_window:
            tab_doc = self.main_window.get_current_tab_doc()
            if tab_doc and tab_doc.doc and getattr(tab_doc, 'file_path', None):
                try:
                    if os.path.abspath(tab_doc.file_path) == os.path.abspath(src):
                        import tempfile
                        with tempfile.NamedTemporaryFile(suffix="_excel_active.pdf", delete=False) as tf:
                            temp_input_pdf = tf.name
                        temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                        tab_doc.export_drawings_to_pdf(temp_doc)
                        temp_doc.save(temp_input_pdf)
                        temp_doc.close()
                        src = temp_input_pdf
                except Exception:
                    temp_input_pdf = None

        self.thread = ExcelConversionThread(src, dest, sheet_mode, detect_tables, pages, temp_file_to_cleanup=temp_input_pdf)
        self.thread.progress_signal.connect(lambda cur, tot, msg: (self.progress_bar.setValue(cur), self.progress_bar.setFormat(f"{cur}/{tot} - {msg}")))
        self.thread.finished_signal.connect(self.on_finished)
        self.thread.start()

    def on_finished(self, success, message):
        self.btn_convert.setEnabled(True)
        self.progress_bar.setVisible(False)
        if success:
            dest = self.txt_dest.text().strip()
            reply = QMessageBox.question(self, LanguageManager.tr("success"),
                                         f"{LanguageManager.tr('excel_success', path=dest)}\n\n{LanguageManager.tr('excel_open_now')}",
                                         QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    os.startfile(dest)
                except Exception:
                    pass
            self.accept()
        else:
            QMessageBox.critical(self, LanguageManager.tr("error"), message)

    def closeEvent(self, event):
        if self.thread and self.thread.isRunning():
            self.thread.cancel()
            self.thread.wait(1000)
        super().closeEvent(event)


class PDFEncryptDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.setWindowTitle(LanguageManager.tr("encrypt_dialog_title"))
        self.resize(680, 680)
        self.setMinimumSize(540, 500)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        self.tab_widget = QTabWidget(self)

        # --- TAB 1: Şifrele (Parola Ekle) ---
        tab_protect = QWidget()
        v_protect_root = QVBoxLayout(tab_protect)
        v_protect_root.setContentsMargins(8, 8, 8, 8)
        v_protect_root.setSpacing(8)

        protect_scroll = QScrollArea()
        protect_scroll.setWidgetResizable(True)
        protect_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        protect_scroll.setStyleSheet("""
            QScrollArea { border: none; background: transparent; }
            QScrollBar:vertical {
                border: none;
                background: #1a1a1a;
                width: 8px;
                margin: 0px;
                border-radius: 4px;
            }
            QScrollBar::handle:vertical {
                background: #444444;
                min-height: 24px;
                border-radius: 4px;
            }
            QScrollBar::handle:vertical:hover {
                background: #555555;
            }
        """)

        protect_content = QWidget()
        v_protect = QVBoxLayout(protect_content)
        v_protect.setContentsMargins(6, 6, 6, 6)
        v_protect.setSpacing(10)

        v_protect.addWidget(QLabel(LanguageManager.tr("encrypt_source_file")))
        h_src1 = QHBoxLayout()
        self.txt_protect_src = QLineEdit()
        self.txt_protect_src.setPlaceholderText("PDF Dosyası Seçin...")
        btn_src1 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src1.clicked.connect(lambda: self.browse_source(self.txt_protect_src, self.txt_protect_dst, "_sifreli.pdf"))
        h_src1.addWidget(self.txt_protect_src)
        h_src1.addWidget(btn_src1)
        v_protect.addLayout(h_src1)

        v_protect.addWidget(QLabel(LanguageManager.tr("encrypt_user_pw")))
        self.txt_protect_pw = QLineEdit()
        self.txt_protect_pw.setEchoMode(QLineEdit.EchoMode.Password)
        v_protect.addWidget(self.txt_protect_pw)

        v_protect.addWidget(QLabel(LanguageManager.tr("encrypt_confirm_pw")))
        self.txt_protect_confirm = QLineEdit()
        self.txt_protect_confirm.setEchoMode(QLineEdit.EchoMode.Password)
        v_protect.addWidget(self.txt_protect_confirm)

        self.chk_show_pw1 = QCheckBox(LanguageManager.tr("encrypt_show_pw"))
        self.chk_show_pw1.toggled.connect(self._toggle_show_pw1)
        v_protect.addWidget(self.chk_show_pw1)

        # Şifre İpucu (Password Hint)
        v_protect.addWidget(QLabel(LanguageManager.tr("encrypt_hint_label")))
        self.txt_protect_hint = QLineEdit()
        self.txt_protect_hint.setPlaceholderText(LanguageManager.tr("encrypt_hint_placeholder"))
        v_protect.addWidget(self.txt_protect_hint)

        # Acil Durum Kurtarma Anahtarı (Emergency Recovery Key)
        grp_recovery = QGroupBox(LanguageManager.tr("encrypt_recovery_group"))
        v_rec = QVBoxLayout(grp_recovery)
        self.chk_gen_recovery = QCheckBox(LanguageManager.tr("encrypt_recovery_chk"))
        self.chk_gen_recovery.setChecked(True)
        self.chk_gen_recovery.toggled.connect(self._toggle_recovery_key)
        v_rec.addWidget(self.chk_gen_recovery)

        h_rec_key = QHBoxLayout()
        self.txt_recovery_key = QLineEdit()
        self.txt_recovery_key.setPlaceholderText(LanguageManager.tr("encrypt_key_placeholder"))
        self.txt_recovery_key.setStyleSheet("font-family: Consolas, monospace; font-size: 13px; font-weight: bold; padding: 6px;")
        self._generate_recovery_key()
        self.btn_regen_key = QPushButton(LanguageManager.tr("encrypt_regen_key"))
        self.btn_regen_key.setToolTip(LanguageManager.tr("regen_key_tip"))
        self.btn_regen_key.clicked.connect(self._generate_recovery_key)
        self.btn_copy_key = QPushButton(LanguageManager.tr("encrypt_copy_key"))
        self.btn_copy_key.clicked.connect(self._copy_recovery_key)
        h_rec_key.addWidget(self.txt_recovery_key, 1)
        h_rec_key.addWidget(self.btn_regen_key)
        h_rec_key.addWidget(self.btn_copy_key)
        v_rec.addLayout(h_rec_key)

        lbl_rec_desc = QLabel(LanguageManager.tr("encrypt_recovery_desc"))
        lbl_rec_desc.setWordWrap(True)
        lbl_rec_desc.setStyleSheet(" font-size: 11px;")
        v_rec.addWidget(lbl_rec_desc)
        v_protect.addWidget(grp_recovery)

        h_algo = QHBoxLayout()
        h_algo.addWidget(QLabel(LanguageManager.tr("encrypt_algo")))
        self.cmb_algo = QComboBox()
        self.cmb_algo.addItem(LanguageManager.tr("encrypt_algo_256"), fitz.PDF_ENCRYPT_AES_256)
        self.cmb_algo.addItem(LanguageManager.tr("encrypt_algo_128"), fitz.PDF_ENCRYPT_AES_128)
        h_algo.addWidget(self.cmb_algo)
        v_protect.addLayout(h_algo)

        grp_perms = QGroupBox(LanguageManager.tr("encrypt_perm_group"))
        v_perms = QVBoxLayout(grp_perms)
        self.chk_perm_print = QCheckBox(LanguageManager.tr("encrypt_perm_print"))
        self.chk_perm_print.setChecked(True)
        self.chk_perm_copy = QCheckBox(LanguageManager.tr("encrypt_perm_copy"))
        self.chk_perm_copy.setChecked(True)
        v_perms.addWidget(self.chk_perm_print)
        v_perms.addWidget(self.chk_perm_copy)
        v_protect.addWidget(grp_perms)

        v_protect.addWidget(QLabel(LanguageManager.tr("encrypt_target_file")))
        h_dst1 = QHBoxLayout()
        self.txt_protect_dst = QLineEdit()
        btn_dst1 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst1.clicked.connect(lambda: self.browse_target(self.txt_protect_dst, "_sifreli.pdf"))
        h_dst1.addWidget(self.txt_protect_dst)
        h_dst1.addWidget(btn_dst1)
        v_protect.addLayout(h_dst1)

        v_protect.addStretch()

        protect_scroll.setWidget(protect_content)
        v_protect_root.addWidget(protect_scroll, 1)

        btn_protect_exec = QPushButton(LanguageManager.tr("encrypt_btn_protect"))
        btn_protect_exec.setObjectName("PrimaryBtn")
        btn_protect_exec.setMinimumHeight(32)
        btn_protect_exec.clicked.connect(self.execute_encrypt)
        v_protect_root.addWidget(btn_protect_exec)

        self.tab_widget.addTab(tab_protect, LanguageManager.tr("encrypt_tab_protect"))

        # --- TAB 2: Şifre Kaldır (Korumayı Çöz) ---
        tab_unlock = QWidget()
        v_unlock = QVBoxLayout(tab_unlock)
        v_unlock.setContentsMargins(14, 14, 14, 14)
        v_unlock.setSpacing(10)

        v_unlock.addWidget(QLabel(LanguageManager.tr("encrypt_source_file")))
        h_src2 = QHBoxLayout()
        self.txt_unlock_src = QLineEdit()
        self.txt_unlock_src.setPlaceholderText("Şifreli PDF Dosyası Seçin...")
        btn_src2 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src2.clicked.connect(lambda: self.browse_source(self.txt_unlock_src, self.txt_unlock_dst, "_sifresiz.pdf"))
        h_src2.addWidget(self.txt_unlock_src)
        h_src2.addWidget(btn_src2)
        v_unlock.addLayout(h_src2)

        v_unlock.addWidget(QLabel(LanguageManager.tr("encrypt_current_pw")))
        self.txt_unlock_pw = QLineEdit()
        self.txt_unlock_pw.setEchoMode(QLineEdit.EchoMode.Password)
        v_unlock.addWidget(self.txt_unlock_pw)

        self.chk_show_pw2 = QCheckBox(LanguageManager.tr("encrypt_show_pw"))
        self.chk_show_pw2.toggled.connect(self._toggle_show_pw2)
        v_unlock.addWidget(self.chk_show_pw2)

        v_unlock.addWidget(QLabel(LanguageManager.tr("encrypt_target_file")))
        h_dst2 = QHBoxLayout()
        self.txt_unlock_dst = QLineEdit()
        btn_dst2 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst2.clicked.connect(lambda: self.browse_target(self.txt_unlock_dst, "_sifresiz.pdf"))
        h_dst2.addWidget(self.txt_unlock_dst)
        h_dst2.addWidget(btn_dst2)
        v_unlock.addLayout(h_dst2)

        v_unlock.addStretch()

        btn_unlock_exec = QPushButton(LanguageManager.tr("encrypt_btn_unlock"))
        btn_unlock_exec.setObjectName("PrimaryBtn")
        btn_unlock_exec.setMinimumHeight(32)
        btn_unlock_exec.clicked.connect(self.execute_decrypt)
        v_unlock.addWidget(btn_unlock_exec)

        self.tab_widget.addTab(tab_unlock, LanguageManager.tr("encrypt_tab_unlock"))

        layout.addWidget(self.tab_widget)

        # Alt Butonlar
        h_bot = QHBoxLayout()
        self.chk_open_tab = QCheckBox(LanguageManager.tr("encrypt_chk_open"))
        self.chk_open_tab.setChecked(True)
        h_bot.addWidget(self.chk_open_tab)
        h_bot.addStretch()
        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.setMinimumHeight(30)
        btn_cancel.clicked.connect(self.reject)
        h_bot.addWidget(btn_cancel)
        layout.addLayout(h_bot)

        # Aktif açık sekme varsa başlangıçta hazırla
        if (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            p = tab_doc.file_path
            self.txt_protect_src.setText(p)
            base_dir = os.path.dirname(p)
            base_name, _ = os.path.splitext(os.path.basename(p))
            self.txt_protect_dst.setText(os.path.join(base_dir, f"{base_name}_sifreli.pdf"))
            self.txt_unlock_src.setText(p)
            self.txt_unlock_dst.setText(os.path.join(base_dir, f"{base_name}_sifresiz.pdf"))

    def _generate_recovery_key(self):
        k1 = secrets.token_hex(2).upper()
        k2 = secrets.token_hex(2).upper()
        k3 = secrets.token_hex(2).upper()
        self.txt_recovery_key.setText(f"AURA-{k1}-{k2}-{k3}")

    def _toggle_recovery_key(self, checked: bool):
        self.txt_recovery_key.setEnabled(checked)
        self.btn_regen_key.setEnabled(checked)
        self.btn_copy_key.setEnabled(checked)

    def _copy_recovery_key(self):
        key = self.txt_recovery_key.text()
        if key:
            QApplication.clipboard().setText(key)
            QMessageBox.information(self, LanguageManager.tr("info"), LanguageManager.tr("encrypt_key_copied"))

    def _toggle_show_pw1(self, checked: bool):
        mode = QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
        self.txt_protect_pw.setEchoMode(mode)
        self.txt_protect_confirm.setEchoMode(mode)

    def _toggle_show_pw2(self, checked: bool):
        mode = QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
        self.txt_unlock_pw.setEchoMode(mode)

    def browse_source(self, line_src: QLineEdit, line_dst: QLineEdit, suffix: str):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            line_src.setText(path)
            base_dir = os.path.dirname(path)
            base_name, _ = os.path.splitext(os.path.basename(path))
            line_dst.setText(os.path.join(base_dir, f"{base_name}{suffix}"))

    def browse_target(self, line_dst: QLineEdit, default_suffix: str):
        current_val = line_dst.text().strip()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), current_val or f"belge{default_suffix}", "PDF (*.pdf)")
        if path:
            line_dst.setText(path)

    def execute_encrypt(self):
        src_path = self.txt_protect_src.text().strip()
        dst_path = self.txt_protect_dst.text().strip()
        pw1 = self.txt_protect_pw.text()
        pw2 = self.txt_protect_confirm.text()

        if not src_path or not os.path.isfile(src_path):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return
        if not pw1:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_pw"))
            return
        if pw1 != pw2:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_pw_mismatch"))
            return
        if not dst_path:
            dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src_path.replace(".pdf", "_sifreli.pdf"), "PDF (*.pdf)")
            if not dst_path:
                return
            self.txt_protect_dst.setText(dst_path)

        try:
            doc = fitz.open(src_path)
            if doc.is_encrypted:
                from PyQt6.QtWidgets import QInputDialog
                curr_pw, ok = QInputDialog.getText(
                    self, LanguageManager.tr("password_required_title"),
                    LanguageManager.tr("encrypt_current_pw"), QLineEdit.EchoMode.Password
                )
                if not ok:
                    doc.close()
                    return
                if doc.authenticate(curr_pw) <= 0:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_wrong_pw"))
                    return

            perms = fitz.PDF_PERM_ACCESSIBILITY
            if self.chk_perm_print.isChecked():
                perms |= fitz.PDF_PERM_PRINT
            if self.chk_perm_copy.isChecked():
                perms |= fitz.PDF_PERM_COPY

            algo = self.cmb_algo.currentData() or fitz.PDF_ENCRYPT_AES_256

            if self.chk_gen_recovery.isChecked():
                if not self.txt_recovery_key.text().strip():
                    self._generate_recovery_key()
                owner_pw = self.txt_recovery_key.text().strip()
            else:
                owner_pw = pw1

            doc.save(
                dst_path,
                encryption=algo,
                user_pw=pw1,
                owner_pw=owner_pw,
                permissions=perms
            )
            doc.close()

            # İpucu varsa dosya sonuna PDF spec uyumlu güvenli bir yorum olarak ekle
            hint_val = self.txt_protect_hint.text().strip()
            if hint_val:
                try:
                    b64 = base64.b64encode(hint_val.encode("utf-8")).decode("ascii")
                    with open(dst_path, "ab") as f:
                        f.write(f"\n%AuraPDF_Hint:{b64}\n".encode("ascii"))
                except Exception:
                    pass

            # Kurtarma anahtarı metin dosyasını otomatik olarak PDF'in bulunduğu dizine kaydet
            saved_txt_msg = ""
            rec_key = self.txt_recovery_key.text().strip()
            if self.chk_gen_recovery.isChecked() and rec_key:
                try:
                    dst_dir = os.path.dirname(os.path.abspath(dst_path))
                    base_name, _ = os.path.splitext(os.path.basename(dst_path))
                    txt_path = os.path.join(dst_dir, f"{base_name}_Kurtarma_Anahtari.txt")
                    
                    sha256_hash = hashlib.sha256()
                    md5_hash = hashlib.md5()
                    file_size = 0
                    if os.path.isfile(dst_path):
                        with open(dst_path, "rb") as f:
                            while chunk := f.read(65536):
                                sha256_hash.update(chunk)
                                md5_hash.update(chunk)
                        file_size = os.path.getsize(dst_path)
                    sha256_hex = sha256_hash.hexdigest()
                    md5_hex = md5_hash.hexdigest()

                    if file_size < 1024:
                        size_str = f"{file_size} bayt"
                    elif file_size < 1024 * 1024:
                        size_str = f"{file_size:,} bayt ({file_size / 1024:.1f} KB)".replace(",", ".")
                    else:
                        size_str = f"{file_size:,} bayt ({file_size / (1024 * 1024):.2f} MB)".replace(",", ".")
                    
                    content = (
                        "============================================================\n"
                        "              AURAPDF - GÜVENLİK VE KURTARMA BİLGİSİ\n"
                        "============================================================\n\n"
                        "[BELGE BİLGİLERİ]\n"
                        f"Orijinal Dosya Adı : {os.path.basename(dst_path)}\n"
                        f"Kayıt Tarihi       : {time.strftime('%Y-%m-%d %H:%M:%S')}\n"
                        f"Dosya Boyutu       : {size_str}\n"
                        f"Tam Dosya Yolu     : {dst_path}\n\n"
                        "[DİJİTAL PARMAK İZİ / DOSYA HASH DEĞERLERİ]\n"
                        f"SHA-256 Özeti      : {sha256_hex}\n"
                        f"MD5 Özeti          : {md5_hex}\n\n"
                        "ÖNEMLİ GÜVENCE NOTU:\n"
                        "Bu belgenin dosya adı veya konumu gelecekte değiştirilse dahi, yukarıda yer\n"
                        "alan benzersiz SHA-256 dijital parmak izi (hash) sayesinde bu kurtarma anahtarının\n"
                        "hangi PDF belgesine ait olduğunu her zaman AuraPDF üzerinden teyit edebilirsiniz.\n\n"
                        "------------------------------------------------------------\n"
                        "ACİL DURUM KURTARMA ANAHTARI (EMERGENCY RECOVERY KEY):\n"
                        f"{rec_key}\n"
                        "------------------------------------------------------------\n\n"
                        f"Parola İpucu       : {hint_val if hint_val else 'Belirtilmedi'}\n\n"
                        "NASIL KULLANILIR?\n"
                        "1. Belgenin normal açma parolasını unutursanız AuraPDF uygulamasını açın.\n"
                        "2. Araçlar (Toolbox) menüsünden 'Şifre Kurtarma Asistanı'nı seçin.\n"
                        "3. 'Kurtarma Anahtarı ile Aç' sekmesinde belgenizi seçin (veya bu .txt dosyasını yükleyin).\n"
                        "4. 'Kurtarma Dosyası (.txt) Yükle' butonu ile bu dosyayı seçtiğinizde, AuraPDF dosya hash değerini\n"
                        "   ve anahtarınızı otomatik doğrulayacaktır.\n"
                        "   (Ya da belgeyi açarken parola yerine doğrudan bu kurtarma anahtarını yazabilirsiniz.)\n\n"
                        "DİKKAT: Bu dosyayı güvenli bir yerde saklayınız.\n"
                        "============================================================\n"
                    )
                    with open(txt_path, "w", encoding="utf-8") as f:
                        f.write(content)
                    
                    saved_txt_msg = f"\n\n{LanguageManager.tr('encrypt_txt_saved', path=txt_path)}"
                except Exception as ex:
                    print(f"Kurtarma anahtarı txt dosyası kaydedilemedi: {ex}")

            info_msg = LanguageManager.tr("encrypt_success_protect", path=dst_path) + saved_txt_msg
            QMessageBox.information(self, LanguageManager.tr("success"), info_msg)
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                self.main_window.add_pdf_tab(dst_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_save", error=e))

    def execute_decrypt(self):
        src_path = self.txt_unlock_src.text().strip()
        dst_path = self.txt_unlock_dst.text().strip()
        pw = self.txt_unlock_pw.text()

        if not src_path or not os.path.isfile(src_path):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return
        if not dst_path:
            dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src_path.replace(".pdf", "_sifresiz.pdf"), "PDF (*.pdf)")
            if not dst_path:
                return
            self.txt_unlock_dst.setText(dst_path)

        try:
            doc = fitz.open(src_path)
            if doc.is_encrypted:
                if doc.authenticate(pw) <= 0:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_wrong_pw"))
                    return
            doc.save(dst_path, garbage=4, deflate=True)
            doc.close()

            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("encrypt_success_unlock", path=dst_path))
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                self.main_window.add_pdf_tab(dst_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_save", error=e))


# --- POPÜLER ŞİFRELER SÖZLÜĞÜ (TOP PASSWORDS LIST) ---
COMMON_PASSWORDS = [
    "123456", "123456789", "qwerty", "password", "12345", "12345678", "111111", "123123",
    "1234567", "1234", "000000", "admin", "welcome", "login", "iloveyou", "secret",
    "master", "starwars", "pass123", "password1", "123321", "sunshine", "football",
    "princess", "computer", "orange", "monkey", "shadow", "superman", "testing",
    "istanbul", "ankara", "izmir", "turkiye", "gs1905", "fb1907", "bjk1903",
    "1905", "1907", "1903", "1453", "2020", "2021", "2022", "2023", "2024", "2025", "2026",
    "qazwsx", "asdfgh", "zxcvbn", "112233", "654321", "987654321", "12344321", "pass", "root",
    "admin123", "123456a", "user", "guest", "default", "parola", "sifre", "gizli",
    "102030", "1111", "0000", "2222", "3333", "4444", "5555", "6666", "7777", "8888", "9999",
    "abc123", "abcdef", "dragon", "baseball", "myspace1", "michelle", "charlie", "andrew"
]


class PasswordRecoveryWorker(QThread):
    progress_signal = pyqtSignal(int, int, float, str)  # current, total, speed, candidate
    found_signal = pyqtSignal(str)
    finished_signal = pyqtSignal(bool, str)

    def __init__(self, src_path: str, dst_path: str, method: str, mask: str = ""):
        super().__init__()
        self.src_path = src_path
        self.dst_path = dst_path
        self.method = method
        self.mask = mask
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            doc = fitz.open(self.src_path)
            if not doc.is_encrypted:
                doc.close()
                self.finished_signal.emit(False, "Belge zaten şifreli değil.")
                return

            candidates = []
            if self.method == "pin4":
                candidates = [f"{i:04d}" for i in range(10000)]
            elif self.method == "pin6":
                candidates = [f"{i:06d}" for i in range(1000000)]
            elif self.method == "dict":
                candidates = list(COMMON_PASSWORDS)
            elif self.method == "mask":
                candidates = self._generate_mask_candidates(self.mask)

            total = len(candidates)
            if total == 0:
                doc.close()
                self.finished_signal.emit(False, "Denebilecek parola şablonu oluşturulamadı.")
                return

            t0 = time.time()
            tested = 0

            for cand in candidates:
                if self._cancelled:
                    doc.close()
                    self.finished_signal.emit(False, LanguageManager.tr("recovery_status_stopped"))
                    return

                tested += 1
                # PyMuPDF authenticate C çağrısı
                if doc.authenticate(cand) > 0:
                    # Parola bulundu!
                    if self.dst_path:
                        try:
                            doc.save(self.dst_path, garbage=4, deflate=True)
                        except Exception:
                            pass
                    doc.close()
                    self.found_signal.emit(cand)
                    self.finished_signal.emit(True, cand)
                    return

                if tested % 30 == 0 or tested == total:
                    elapsed = max(0.001, time.time() - t0)
                    speed = tested / elapsed
                    self.progress_signal.emit(tested, total, speed, cand)

            doc.close()
            self.finished_signal.emit(False, LanguageManager.tr("recovery_status_not_found"))
        except Exception as e:
            self.finished_signal.emit(False, str(e))

    def _generate_mask_candidates(self, mask: str) -> List[str]:
        # '?' -> 0-9, '*' -> a-z
        results = [""]
        for ch in mask:
            if ch == '?':
                new_res = []
                for prefix in results:
                    for d in "0123456789":
                        new_res.append(prefix + d)
                        if len(new_res) > 200000:
                            return new_res
                results = new_res
            elif ch == '*':
                new_res = []
                for prefix in results:
                    for l in "abcdefghijklmnopqrstuvwxyz":
                        new_res.append(prefix + l)
                        if len(new_res) > 200000:
                            return new_res
                results = new_res
            else:
                results = [prefix + ch for prefix in results]
        return results


class PDFPasswordRecoveryDialog(QDialog):
    def __init__(self, main_window, initial_file: str = ""):
        super().__init__(main_window)
        self.main_window = main_window
        self.initial_file = initial_file
        self.worker: Optional[PasswordRecoveryWorker] = None
        self.found_pw = None
        self.current_doc_hash = ""
        self.loaded_txt_hash = None

        self.setWindowTitle(LanguageManager.tr("recovery_dialog_title"))
        self.resize(660, 580)
        self.setMinimumSize(540, 460)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        self.tab_widget = QTabWidget(self)

        # --- TAB 1: Kurtarma Anahtarı ile Aç ---
        tab_key = QWidget()
        v_key = QVBoxLayout(tab_key)
        v_key.setContentsMargins(14, 14, 14, 14)
        v_key.setSpacing(10)

        v_key.addWidget(QLabel(LanguageManager.tr("recovery_source_file")))
        h_src1 = QHBoxLayout()
        self.txt_key_src = QLineEdit()
        self.txt_key_src.textChanged.connect(self._on_source_file_changed)
        btn_src1 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src1.clicked.connect(lambda: self.browse_source(self.txt_key_src, self.txt_key_dst))
        h_src1.addWidget(self.txt_key_src)
        h_src1.addWidget(btn_src1)
        v_key.addLayout(h_src1)

        # Hash ve Dijital Parmak İzi Doğrulama Kartı
        frame_hash = QFrame()
        frame_hash.setStyleSheet("QFrame { background: rgba(128, 128, 128, 0.08); border-radius: 6px; padding: 6px 8px; }")
        v_hash_card = QVBoxLayout(frame_hash)
        v_hash_card.setContentsMargins(6, 6, 6, 6)
        v_hash_card.setSpacing(4)

        h_hash_line = QHBoxLayout()
        h_hash_line.addWidget(QLabel(LanguageManager.tr("recovery_file_hash")))
        self.lbl_src_hash = QLabel(LanguageManager.tr("recovery_no_file_selected"))
        self.lbl_src_hash.setStyleSheet("font-family: Consolas, monospace; font-size: 11px;")
        self.lbl_src_hash.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        h_hash_line.addWidget(self.lbl_src_hash, 1)

        self.btn_copy_hash = QPushButton(LanguageManager.tr("recovery_copy_hash"))
        self.btn_copy_hash.setStyleSheet("padding: 2px 8px; font-size: 11px;")
        self.btn_copy_hash.setEnabled(False)
        self.btn_copy_hash.clicked.connect(self._copy_src_hash)
        h_hash_line.addWidget(self.btn_copy_hash)
        v_hash_card.addLayout(h_hash_line)

        self.lbl_hash_match_status = QLabel("")
        self.lbl_hash_match_status.setStyleSheet("font-size: 11px; font-weight: 500;")
        self.lbl_hash_match_status.setVisible(False)
        v_hash_card.addWidget(self.lbl_hash_match_status)
        v_key.addWidget(frame_hash)

        # Anahtar Başlığı ve TXT'den Yükle Butonu
        h_key_head = QHBoxLayout()
        h_key_head.addWidget(QLabel(LanguageManager.tr("recovery_key_label")))
        h_key_head.addStretch()
        btn_load_txt = QPushButton(LanguageManager.tr("recovery_btn_load_txt"))
        btn_load_txt.setObjectName("SecondaryBtn")
        btn_load_txt.clicked.connect(self.load_from_recovery_file)
        h_key_head.addWidget(btn_load_txt)
        v_key.addLayout(h_key_head)

        self.txt_key_input = QLineEdit()
        self.txt_key_input.setPlaceholderText(LanguageManager.tr("recovery_key_placeholder"))
        self.txt_key_input.setStyleSheet("font-family: Consolas, monospace; font-size: 13px; font-weight: bold; ")
        v_key.addWidget(self.txt_key_input)

        v_key.addWidget(QLabel(LanguageManager.tr("recovery_target_file")))
        h_dst1 = QHBoxLayout()
        self.txt_key_dst = QLineEdit()
        btn_dst1 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst1.clicked.connect(lambda: self.browse_target(self.txt_key_dst))
        h_dst1.addWidget(self.txt_key_dst)
        h_dst1.addWidget(btn_dst1)
        v_key.addLayout(h_dst1)

        v_key.addStretch()

        btn_unlock_key = QPushButton(LanguageManager.tr("recovery_btn_unlock_key"))
        btn_unlock_key.setObjectName("SuccessBtn")
        btn_unlock_key.clicked.connect(self.execute_key_unlock)
        v_key.addWidget(btn_unlock_key)

        self.tab_widget.addTab(tab_key, LanguageManager.tr("recovery_tab_key"))

        # --- TAB 2: Parola Bulma Asistanı (PIN & Sözlük) ---
        tab_brute = QWidget()
        v_brute = QVBoxLayout(tab_brute)
        v_brute.setContentsMargins(14, 14, 14, 14)
        v_brute.setSpacing(10)

        v_brute.addWidget(QLabel(LanguageManager.tr("recovery_source_file")))
        h_src2 = QHBoxLayout()
        self.txt_brute_src = QLineEdit()
        btn_src2 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src2.clicked.connect(lambda: self.browse_source(self.txt_brute_src, self.txt_brute_dst))
        h_src2.addWidget(self.txt_brute_src)
        h_src2.addWidget(btn_src2)
        v_brute.addLayout(h_src2)

        h_method = QHBoxLayout()
        h_method.addWidget(QLabel(LanguageManager.tr("recovery_method_label")))
        self.cmb_method = QComboBox()
        self.cmb_method.addItem(LanguageManager.tr("recovery_method_pin4"), "pin4")
        self.cmb_method.addItem(LanguageManager.tr("recovery_method_dict"), "dict")
        self.cmb_method.addItem(LanguageManager.tr("recovery_method_mask"), "mask")
        self.cmb_method.addItem(LanguageManager.tr("recovery_method_pin6"), "pin6")
        self.cmb_method.currentIndexChanged.connect(self._on_method_changed)
        h_method.addWidget(self.cmb_method)
        v_brute.addLayout(h_method)

        # Mask input container (gizli başlar)
        self.widget_mask = QWidget()
        v_mask = QVBoxLayout(self.widget_mask)
        v_mask.setContentsMargins(0, 0, 0, 0)
        v_mask.setSpacing(4)
        v_mask.addWidget(QLabel(LanguageManager.tr("recovery_mask_label")))
        self.txt_mask = QLineEdit("202?")
        self.txt_mask.setPlaceholderText("Örn: 202? veya pass****")
        v_mask.addWidget(self.txt_mask)
        self.widget_mask.setVisible(False)
        v_brute.addWidget(self.widget_mask)

        v_brute.addWidget(QLabel(LanguageManager.tr("recovery_target_file")))
        h_dst2 = QHBoxLayout()
        self.txt_brute_dst = QLineEdit()
        btn_dst2 = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst2.clicked.connect(lambda: self.browse_target(self.txt_brute_dst))
        h_dst2.addWidget(self.txt_brute_dst)
        h_dst2.addWidget(btn_dst2)
        v_brute.addLayout(h_dst2)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        v_brute.addWidget(self.progress_bar)

        self.lbl_status = QLabel(LanguageManager.tr("recovery_status_ready"))
        self.lbl_status.setStyleSheet(" font-size: 11px;")
        v_brute.addWidget(self.lbl_status)

        v_brute.addStretch()

        self.btn_start_stop = QPushButton(LanguageManager.tr("recovery_btn_start"))
        self.btn_start_stop.setObjectName("PrimaryBtn")
        self.btn_start_stop.clicked.connect(self._toggle_bruteforce)
        v_brute.addWidget(self.btn_start_stop)

        self.tab_widget.addTab(tab_brute, LanguageManager.tr("recovery_tab_bruteforce"))

        layout.addWidget(self.tab_widget)

        # Alt Butonlar
        h_bot = QHBoxLayout()
        self.chk_open_tab = QCheckBox(LanguageManager.tr("encrypt_chk_open"))
        self.chk_open_tab.setChecked(True)
        h_bot.addWidget(self.chk_open_tab)
        h_bot.addStretch()
        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.clicked.connect(self.close)
        h_bot.addWidget(btn_cancel)
        layout.addLayout(h_bot)

        # Başlangıç dosyası ayarla
        p = self.initial_file or ""
        if not p and (tab_doc := self.main_window.get_current_tab_doc()) and tab_doc.file_path:
            p = tab_doc.file_path
        if p:
            self.set_source_file(p)

    def set_source_file(self, p: str):
        self.txt_key_src.setText(p)
        self.txt_brute_src.setText(p)
        base_dir = os.path.dirname(p)
        base_name, _ = os.path.splitext(os.path.basename(p))
        dst = os.path.join(base_dir, f"{base_name}_kurtarildi.pdf")
        self.txt_key_dst.setText(dst)
        self.txt_brute_dst.setText(dst)
        self._on_source_file_changed()

    def _compute_sha256(self, filepath: str) -> str:
        try:
            if not filepath or not os.path.isfile(filepath):
                return ""
            h = hashlib.sha256()
            with open(filepath, "rb") as f:
                while chunk := f.read(65536):
                    h.update(chunk)
            return h.hexdigest().lower()
        except Exception:
            return ""

    def _on_source_file_changed(self, text: str = ""):
        path = self.txt_key_src.text().strip()
        if path and os.path.isfile(path):
            hash_val = self._compute_sha256(path)
            self.current_doc_hash = hash_val
            if hash_val:
                self.lbl_src_hash.setText(hash_val)
                self.btn_copy_hash.setEnabled(True)
            else:
                self.lbl_src_hash.setText(LanguageManager.tr("recovery_no_file_selected"))
                self.btn_copy_hash.setEnabled(False)
        else:
            self.current_doc_hash = ""
            self.lbl_src_hash.setText(LanguageManager.tr("recovery_no_file_selected"))
            self.btn_copy_hash.setEnabled(False)
        self._check_hash_match()

    def _check_hash_match(self):
        if getattr(self, 'loaded_txt_hash', None) and getattr(self, 'current_doc_hash', None):
            if self.loaded_txt_hash.lower() == self.current_doc_hash.lower():
                self.lbl_hash_match_status.setText(f"✓ {LanguageManager.tr('recovery_hash_match')}")
                self.lbl_hash_match_status.setStyleSheet("color: #10b981; font-size: 11px; font-weight: bold;")
                self.lbl_hash_match_status.setVisible(True)
            else:
                short_txt = self.loaded_txt_hash[:12] + "..."
                short_doc = self.current_doc_hash[:12] + "..."
                self.lbl_hash_match_status.setText(f"⚠ {LanguageManager.tr('recovery_hash_mismatch', txt_hash=short_txt, doc_hash=short_doc)}")
                self.lbl_hash_match_status.setStyleSheet("color: #f59e0b; font-size: 11px; font-weight: bold;")
                self.lbl_hash_match_status.setVisible(True)
        else:
            self.lbl_hash_match_status.setVisible(False)

    def _copy_src_hash(self):
        if getattr(self, 'current_doc_hash', None):
            QApplication.clipboard().setText(self.current_doc_hash)
            QMessageBox.information(self, LanguageManager.tr("confirm_title"), LanguageManager.tr("recovery_hash_copied"))

    def load_from_recovery_file(self):
        cur_src = self.txt_key_src.text().strip()
        init_dir = os.path.dirname(cur_src) if cur_src and os.path.isfile(cur_src) else ""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            LanguageManager.tr("recovery_btn_load_txt"),
            init_dir,
            LanguageManager.tr("recovery_txt_filter")
        )
        if not file_path:
            return
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        except Exception as e:
            QMessageBox.warning(self, LanguageManager.tr("error"), str(e))
            return

        key = ""
        match_aura = re.search(r"AURA-[A-Z0-9]{4}-[A-Z0-9]{4}-[A-Z0-9]{4}", content)
        if match_aura:
            key = match_aura.group(0)
        else:
            match_gen = re.search(r"(?:ACİL DURUM KURTARMA ANAHTARI|EMERGENCY RECOVERY KEY)[:\s\n\-]+([^\n\r]+)", content, re.IGNORECASE)
            if match_gen:
                candidate = match_gen.group(1).strip()
                if candidate and not candidate.startswith("---") and not candidate.startswith("==="):
                    key = candidate

        match_sha = re.search(r"SHA-256\s*(?:Özeti|Hash|Digest)?\s*[:=]\s*([a-fA-F0-9]{64})", content, re.IGNORECASE)
        if match_sha:
            self.loaded_txt_hash = match_sha.group(1).strip().lower()
        else:
            self.loaded_txt_hash = None

        if key:
            self.txt_key_input.setText(key)

        if not key and not self.loaded_txt_hash:
            QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("recovery_txt_read_err"))
            return

        cur_pdf = self.txt_key_src.text().strip()
        if not cur_pdf or not os.path.isfile(cur_pdf):
            candidate_pdf = file_path.replace("_Kurtarma_Anahtari.txt", ".pdf")
            if not os.path.isfile(candidate_pdf):
                match_name = re.search(r"(?:Orijinal Dosya Adı|Dosya Adı)\s*[:=]\s*(.+)", content, re.IGNORECASE)
                if match_name:
                    fname = match_name.group(1).strip()
                    cand2 = os.path.join(os.path.dirname(file_path), fname)
                    if os.path.isfile(cand2):
                        candidate_pdf = cand2
            if os.path.isfile(candidate_pdf):
                self.set_source_file(candidate_pdf)

        self._check_hash_match()
        if key:
            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("recovery_txt_loaded_ok"))

    def _on_method_changed(self):
        method = self.cmb_method.currentData()
        self.widget_mask.setVisible(method == "mask")

    def browse_source(self, line_src: QLineEdit, line_dst: QLineEdit):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            line_src.setText(path)
            base_dir = os.path.dirname(path)
            base_name, _ = os.path.splitext(os.path.basename(path))
            line_dst.setText(os.path.join(base_dir, f"{base_name}_kurtarildi.pdf"))

    def browse_target(self, line_dst: QLineEdit):
        current_val = line_dst.text().strip()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), current_val or "kurtarilan_belge.pdf", "PDF (*.pdf)")
        if path:
            line_dst.setText(path)

    def execute_key_unlock(self):
        src_path = self.txt_key_src.text().strip()
        dst_path = self.txt_key_dst.text().strip()
        key = self.txt_key_input.text().strip()

        if not src_path or not os.path.isfile(src_path):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return
        if not key:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("recovery_err_invalid_key"))
            return
        if not dst_path:
            dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src_path.replace(".pdf", "_kurtarildi.pdf"), "PDF (*.pdf)")
            if not dst_path:
                return
            self.txt_key_dst.setText(dst_path)

        try:
            doc = fitz.open(src_path)
            if doc.is_encrypted:
                if doc.authenticate(key) <= 0:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("recovery_err_invalid_key"))
                    return
            doc.save(dst_path, garbage=4, deflate=True)
            doc.close()

            self.found_pw = key
            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("recovery_success_saved", path=dst_path))
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                self.main_window.add_pdf_tab(dst_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def _toggle_bruteforce(self):
        if self.worker and self.worker.isRunning():
            self.worker.cancel()
            self.btn_start_stop.setText(LanguageManager.tr("recovery_btn_start"))
            self.btn_start_stop.setObjectName("PrimaryBtn")
            self.btn_start_stop.setStyle(self.btn_start_stop.style())
            return

        src_path = self.txt_brute_src.text().strip()
        dst_path = self.txt_brute_dst.text().strip()
        if not src_path or not os.path.isfile(src_path):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return
        if not dst_path:
            dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src_path.replace(".pdf", "_kurtarildi.pdf"), "PDF (*.pdf)")
            if not dst_path:
                return
            self.txt_brute_dst.setText(dst_path)

        method = self.cmb_method.currentData()
        mask = self.txt_mask.text().strip() if method == "mask" else ""

        self.progress_bar.setValue(0)
        self.lbl_status.setText("Başlatılıyor...")
        self.btn_start_stop.setText(LanguageManager.tr("recovery_btn_stop"))
        self.btn_start_stop.setObjectName("DangerBtn")
        self.btn_start_stop.setStyle(self.btn_start_stop.style())

        self.worker = PasswordRecoveryWorker(src_path, dst_path, method, mask)
        self.worker.progress_signal.connect(self._on_worker_progress)
        self.worker.found_signal.connect(self._on_worker_found)
        self.worker.finished_signal.connect(self._on_worker_finished)
        self.worker.start()

    def _on_worker_progress(self, current: int, total: int, speed: float, candidate: str):
        pct = int((current / max(1, total)) * 100)
        self.progress_bar.setValue(pct)
        rem_sec = int((total - current) / max(0.1, speed))
        rem_str = f"{rem_sec // 60} dk {rem_sec % 60} sn" if rem_sec >= 60 else f"{rem_sec} sn"
        self.lbl_status.setText(LanguageManager.tr("recovery_status_running", tested=current, speed=speed, remaining=rem_str))

    def _on_worker_found(self, pw: str):
        self.found_pw = pw
        self.progress_bar.setValue(100)
        self.lbl_status.setText(LanguageManager.tr("recovery_status_found", pw=pw))

    def _on_worker_finished(self, success: bool, msg: str):
        self.btn_start_stop.setText(LanguageManager.tr("recovery_btn_start"))
        self.btn_start_stop.setObjectName("PrimaryBtn")
        self.btn_start_stop.setStyle(self.btn_start_stop.style())

        if success:
            dst_path = self.txt_brute_dst.text().strip()
            QMessageBox.information(
                self,
                LanguageManager.tr("success"),
                f"{LanguageManager.tr('recovery_status_found', pw=msg)}\n\n" + LanguageManager.tr("recovery_success_saved", path=dst_path)
            )
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab') and os.path.isfile(dst_path):
                self.main_window.add_pdf_tab(dst_path)
            self.accept()
        else:
            if msg and msg != LanguageManager.tr("recovery_status_stopped"):
                QMessageBox.warning(self, LanguageManager.tr("warning"), msg)

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            self.worker.cancel()
            self.worker.wait(1000)
        super().closeEvent(event)


class InteractiveWatermarkPreviewLabel(QLabel):
    positionChanged = pyqtSignal(QPointF)
    scaleChanged = pyqtSignal(int)
    angleChanged = pyqtSignal(int)
    widthChanged = pyqtSignal(int)
    heightChanged = pyqtSignal(int)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.norm_pos = QPointF(0.5, 0.5)  # Normalized center (0.5, 0.5)
        self.is_dragging = False
        self.is_resizing = False
        self.is_rotating = False
        self.active_handle = None
        self.watermark_rect = QRectF()
        self.watermark_center = QPointF()
        self.current_scale_value = 50
        self.current_angle = 0
        self.init_dist_from_center = 50.0
        self.init_scale = 50
        self.lock_aspect = True
        self.preview_page_w = 595.0
        self.preview_page_h = 842.0
        self.setMouseTracking(True)
        self.setCursor(Qt.CursorShape.OpenHandCursor)
        self.setToolTip(LanguageManager.tr("watermark_drag_hint"))

    def _event_to_norm_pos(self, event) -> Optional[QPointF]:
        pix = self.pixmap()
        if not pix or pix.isNull():
            return None
        pw = pix.width()
        ph = pix.height()
        ox = (self.width() - pw) / 2.0
        oy = (self.height() - ph) / 2.0
        pos = event.position() if hasattr(event, "position") else event.pos()
        px = pos.x()
        py = pos.y()
        nx = max(0.05, min(0.95, (px - ox) / max(1.0, pw)))
        ny = max(0.05, min(0.95, (py - oy) / max(1.0, ph)))
        return QPointF(nx, ny)

    def _hit_test_handle(self, pos) -> Optional[str]:
        if not self.watermark_rect or not self.watermark_rect.isValid() or self.watermark_rect.isEmpty():
            return None
        cx = self.watermark_center.x() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().x()
        cy = self.watermark_center.y() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().y()
        angle = getattr(self, 'current_angle', 0)
        rad = math.radians(angle)
        dx = pos.x() - cx
        dy = pos.y() - cy
        unrot_x = cx + dx * math.cos(rad) - dy * math.sin(rad)
        unrot_y = cy + dx * math.sin(rad) + dy * math.cos(rad)
        unrot_pt = QPointF(unrot_x, unrot_y)

        r = self.watermark_rect
        # 1. Rotation handle (located 22px above top center)
        rot_pt = QPointF(r.center().x(), r.top() - 22.0)
        if math.hypot(unrot_pt.x() - rot_pt.x(), unrot_pt.y() - rot_pt.y()) <= 13.0:
            return 'rotate'

        # 2. Side handles (en-boy kilitli değilken orta tutamaçlar)
        if not getattr(self, 'lock_aspect', True):
            side_handles = {
                'ml': QPointF(r.left(), r.center().y()),
                'mr': QPointF(r.right(), r.center().y()),
                'mt': QPointF(r.center().x(), r.top()),
                'mb': QPointF(r.center().x(), r.bottom()),
            }
            for name, pt in side_handles.items():
                if math.hypot(unrot_pt.x() - pt.x(), unrot_pt.y() - pt.y()) <= 13.0:
                    return name

        # 3. Corner handles
        corners = {
            'tl': QPointF(r.left(), r.top()),
            'tr': QPointF(r.right(), r.top()),
            'bl': QPointF(r.left(), r.bottom()),
            'br': QPointF(r.right(), r.bottom()),
        }
        for name, pt in corners.items():
            if math.hypot(unrot_pt.x() - pt.x(), unrot_pt.y() - pt.y()) <= 14.0:
                return name
        return None

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            pos = event.position() if hasattr(event, "position") else event.pos()
            handle = self._hit_test_handle(pos)
            if handle == 'rotate':
                self.is_rotating = True
                self.active_handle = 'rotate'
                self.setCursor(Qt.CursorShape.ClosedHandCursor)
            elif handle:
                self.is_resizing = True
                self.active_handle = handle
                cx = self.watermark_center.x() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().x()
                cy = self.watermark_center.y() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().y()
                self.init_dist_from_center = max(10.0, math.hypot(pos.x() - cx, pos.y() - cy))
                self.init_scale = self.current_scale_value
            else:
                self.is_dragging = True
                self.setCursor(Qt.CursorShape.ClosedHandCursor)
                new_pos = self._event_to_norm_pos(event)
                if new_pos is not None:
                    self.norm_pos = new_pos
                    self.positionChanged.emit(self.norm_pos)
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        pos = event.position() if hasattr(event, "position") else event.pos()
        if self.is_rotating and (event.buttons() & Qt.MouseButton.LeftButton):
            cx = self.watermark_center.x() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().x()
            cy = self.watermark_center.y() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().y()
            dx = pos.x() - cx
            dy = pos.y() - cy
            rad = math.atan2(dy, dx) + math.pi / 2.0
            deg = -math.degrees(rad)
            norm_angle = int(round((deg + 180.0) % 360.0 - 180.0))
            if event.modifiers() & Qt.KeyboardModifier.ShiftModifier:
                norm_angle = int(round(norm_angle / 15.0) * 15)
                if norm_angle > 180: norm_angle -= 360
            self.current_angle = norm_angle
            self.angleChanged.emit(norm_angle)
        elif self.is_resizing and (event.buttons() & Qt.MouseButton.LeftButton):
            cx = self.watermark_center.x() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().x()
            cy = self.watermark_center.y() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().y()
            angle = getattr(self, 'current_angle', 0)
            rad = math.radians(angle)
            dx = pos.x() - cx
            dy = pos.y() - cy
            unrot_x = cx + dx * math.cos(rad) - dy * math.sin(rad)
            unrot_y = cy + dx * math.sin(rad) + dy * math.cos(rad)

            pix = self.pixmap()
            pw = float(pix.width()) if pix else 260.0
            ph = float(pix.height()) if pix else 360.0
            real_w = float(getattr(self, 'preview_page_w', 595.0))
            real_h = float(getattr(self, 'preview_page_h', 842.0))

            if self.active_handle in ('ml', 'mr'):
                pixel_w = max(10.0, 2.0 * abs(unrot_x - cx))
                new_w_pt = max(10, min(3000, int(round(pixel_w * (real_w / max(1.0, pw))))))
                self.widthChanged.emit(new_w_pt)
            elif self.active_handle in ('mt', 'mb'):
                pixel_h = max(10.0, 2.0 * abs(unrot_y - cy))
                new_h_pt = max(10, min(3000, int(round(pixel_h * (real_h / max(1.0, ph))))))
                self.heightChanged.emit(new_h_pt)
            elif not getattr(self, 'lock_aspect', True):
                pixel_w = max(10.0, 2.0 * abs(unrot_x - cx))
                pixel_h = max(10.0, 2.0 * abs(unrot_y - cy))
                new_w_pt = max(10, min(3000, int(round(pixel_w * (real_w / max(1.0, pw))))))
                new_h_pt = max(10, min(3000, int(round(pixel_h * (real_h / max(1.0, ph))))))
                self.widthChanged.emit(new_w_pt)
                self.heightChanged.emit(new_h_pt)
            else:
                cur_dist = math.hypot(pos.x() - cx, pos.y() - cy)
                ratio = cur_dist / max(1.0, self.init_dist_from_center)
                new_scale = max(5, min(300, int(round(self.init_scale * ratio))))
                self.scaleChanged.emit(new_scale)
        elif self.is_dragging and (event.buttons() & Qt.MouseButton.LeftButton):
            new_pos = self._event_to_norm_pos(event)
            if new_pos is not None:
                self.norm_pos = new_pos
                self.positionChanged.emit(self.norm_pos)
        else:
            handle = self._hit_test_handle(pos)
            if handle == 'rotate':
                self.setCursor(Qt.CursorShape.PointingHandCursor)
                self.setToolTip(LanguageManager.tr("watermark_rotate_hint"))
            elif handle in ('ml', 'mr'):
                self.setCursor(Qt.CursorShape.SizeHorCursor)
                self.setToolTip("")
            elif handle in ('mt', 'mb'):
                self.setCursor(Qt.CursorShape.SizeVerCursor)
                self.setToolTip("")
            elif handle in ('tl', 'br'):
                self.setCursor(Qt.CursorShape.SizeFDiagCursor)
                self.setToolTip("")
            elif handle in ('tr', 'bl'):
                self.setCursor(Qt.CursorShape.SizeBDiagCursor)
                self.setToolTip("")
            else:
                cx = self.watermark_center.x() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().x()
                cy = self.watermark_center.y() if (hasattr(self, 'watermark_center') and not self.watermark_center.isNull()) else self.watermark_rect.center().y()
                angle = getattr(self, 'current_angle', 0)
                rad = math.radians(angle)
                dx = pos.x() - cx
                dy = pos.y() - cy
                unrot_x = cx + dx * math.cos(rad) - dy * math.sin(rad)
                unrot_y = cy + dx * math.sin(rad) + dy * math.cos(rad)
                unrot_pt = QPointF(unrot_x, unrot_y)
                if self.watermark_rect and self.watermark_rect.isValid() and self.watermark_rect.contains(unrot_pt):
                    self.setCursor(Qt.CursorShape.OpenHandCursor)
                    self.setToolTip(LanguageManager.tr("watermark_drag_hint"))
                else:
                    self.setCursor(Qt.CursorShape.ArrowCursor)
                    self.setToolTip("")
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.is_rotating = False
            self.is_resizing = False
            self.is_dragging = False
            self.active_handle = None
            self.setCursor(Qt.CursorShape.OpenHandCursor)
            if self.norm_pos is not None:
                self.positionChanged.emit(self.norm_pos)
        super().mouseReleaseEvent(event)

    def mouseDoubleClickEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.reset_to_center()
        super().mouseDoubleClickEvent(event)

    def wheelEvent(self, event):
        delta = event.angleDelta().y()
        if delta != 0:
            step = 5 if delta > 0 else -5
            new_scale = max(5, min(300, self.current_scale_value + step))
            self.scaleChanged.emit(new_scale)
            event.accept()
            return
        super().wheelEvent(event)

    def reset_to_center(self):
        self.norm_pos = QPointF(0.5, 0.5)
        self.positionChanged.emit(self.norm_pos)


def get_watermark_fitz_font(font_key: str = "helv") -> fitz.Font:
    """
    Filigran için Türkçe karakterleri %100 destekleyen font nesnesini döndürür.
    Sistemdeki TrueType fontları (Arial, Times New Roman vb.) kullanarak
    CIDFontType2 Unicode font sağlar ve Adobe Acrobat CFF font ayıklama hatalarını önler.
    """
    import os
    from src.ui.direct_text_editor import get_unicode_font_path
    fpath = get_unicode_font_path(font_key, bold=True, italic=False)
    if fpath and os.path.isfile(fpath):
        try:
            return fitz.Font(fontfile=fpath)
        except Exception:
            pass
    try:
        fpath_arial = get_unicode_font_path("Arial", bold=True, italic=False)
        if fpath_arial and os.path.isfile(fpath_arial):
            return fitz.Font(fontfile=fpath_arial)
    except Exception:
        pass
    try:
        return fitz.Font("helv")
    except Exception:
        return None


def calculate_safe_watermark_fontsize(text: str, angle_deg: float, page_w: float, page_h: float, margin: float = 35.0, font_key: str = "helv") -> float:
    if not text:
        return 50.0
    rad = math.radians(abs(angle_deg))
    avail_w = max(20.0, page_w - 2 * margin)
    avail_h = max(20.0, page_h - 2 * margin)
    import os
    from src.ui.direct_text_editor import get_unicode_font_path
    wm_font_file = get_unicode_font_path(font_key, bold=True, italic=False)
    try:
        if wm_font_file and os.path.isfile(wm_font_file):
            len100 = fitz.Font(fontfile=wm_font_file).text_length(text, fontsize=100.0)
        else:
            fitz_font = get_watermark_fitz_font(font_key)
            len100 = fitz_font.text_length(text, fontsize=100.0) if fitz_font else fitz.get_text_length(text, fontsize=100.0)
    except Exception:
        len100 = fitz.get_text_length(text, fontsize=100.0)
    unit_w = max(1.0, len100 / 100.0)
    unit_h = 0.85
    bb_w = unit_w * math.cos(rad) + unit_h * math.sin(rad)
    bb_h = unit_w * math.sin(rad) + unit_h * math.cos(rad)
    max_fs = min(avail_w / max(0.001, bb_w), avail_h / max(0.001, bb_h))
    return max(8.0, max_fs)


def remove_watermark_from_doc(doc: fitz.Document) -> bool:
    """
    Belgedeki AuraPDF filigranını (metin veya görsel) sayfa içerik akışlarından (content streams)
    ve gömülü metadatadan güvenli ve kalıcı bir şekilde temizler.
    Bu işlem hem filigranı oluşturan bilgisayarda hem de dosyanın aktarıldığı farklı bilgisayarlarda
    (PC1 -> PC2) çalışır. Orijinal sayfa metinlerine veya çizimlerine asla zarar vermez.
    """
    removed = False

    # 1. Önce belgede etiketli (% AuraPDF Watermark Stream) akış olup olmadığını kontrol et
    has_tagged_stream = False
    for p in doc:
        try:
            for x in p.get_contents():
                st = doc.xref_stream(x)
                if b'% AuraPDF Watermark Stream' in st:
                    has_tagged_stream = True
                    break
        except Exception:
            pass
        if has_tagged_stream:
            break

    settings = {}
    if hasattr(doc, 'embfile_names') and 'aurapdf_watermark.json' in doc.embfile_names():
        try:
            import json
            wdata = json.loads(doc.embfile_get('aurapdf_watermark.json').decode('utf-8'))
            settings = wdata.get('settings', {})
        except Exception as e:
            print("Error reading aurapdf_watermark.json:", e)

    is_text = (settings.get('type', 'text') == 'text')
    text = settings.get('text', '').strip()
    hex_tokens = []
    if not has_tagged_stream and is_text and text:
        # Sadece eski (etiketsiz) belgelerde tam metin eşleşmesi kullan (tek tek kelimeler değil)
        try:
            hex_tokens.append(''.join(f'{ord(c):02x}' for c in text).encode('ascii').lower())
        except Exception:
            pass
        try:
            hex_tokens.append(text.encode('utf-8').lower())
        except Exception:
            pass

    total_pages = len(doc)
    for p_idx in range(total_pages):
        page = doc[p_idx]
        try:
            contents = page.get_contents()
        except Exception:
            contents = []
        if not contents:
            continue

        remaining = []
        page_changed = False
        for x in contents:
            try:
                s = doc.xref_stream(x)
            except Exception:
                remaining.append(x)
                continue

            is_wm = False
            if has_tagged_stream:
                # Etiketli filigran varsa SADECE bu etiketi taşıyan akışı temizle (%100 hassasiyet)
                if b'% AuraPDF Watermark Stream' in s:
                    is_wm = True
            else:
                # Eski sürüm AuraPDF belgeleri için güvenli geri uyumluluk fallback'i
                if hex_tokens and any(tok in s.lower() for tok in hex_tokens):
                    if b'BT' in s and b'ET' in s and (b'cm' in s or b'/fitzca' in s or b'/F_AuraPDF_WM' in s or b'/F_WM' in s):
                        is_wm = True
                elif not is_text and b'/fzImg' in s and b'Do' in s and len(s) < 400:
                    is_wm = True
                elif is_text and (b'/F_AuraPDF_WM' in s or b'/F_WM' in s) and b'BT' in s and b'ET' in s:
                    is_wm = True

            if is_wm:
                page_changed = True
            else:
                remaining.append(x)

        if page_changed:
            try:
                if not remaining:
                    doc.xref_set_key(page.xref, 'Contents', 'null')
                elif len(remaining) == 1:
                    page.set_contents(remaining[0])
                else:
                    page.set_contents(remaining)
                removed = True
            except Exception as e:
                print(f"Error updating page {p_idx} contents:", e)

    if hasattr(doc, 'embfile_names') and 'aurapdf_watermark.json' in doc.embfile_names():
        try:
            doc.embfile_del('aurapdf_watermark.json')
            removed = True
        except Exception:
            pass

    return removed


def apply_watermark_to_doc(doc: fitz.Document, settings: dict):
    is_text = (settings.get("type", "text") == "text")
    total_pages = len(doc)
    pages_mode = settings.get("pages_mode", "all")
    if pages_mode == "first":
        pages_to_process = [0]
    elif pages_mode == "current":
        curr_p = int(settings.get("current_page_idx", 0))
        pages_to_process = [max(0, min(curr_p, total_pages - 1))]
    elif pages_mode == "custom":
        pages_to_process = parse_page_range(settings.get("custom_pages", ""), total_pages)
    else:
        pages_to_process = list(range(total_pages))

    overlay = (settings.get("layer", "overlay") == "overlay")
    nx = float(settings.get("norm_pos_x", 0.5))
    ny = float(settings.get("norm_pos_y", 0.5))

    if is_text:
        text = settings.get("text", "").strip() or "AURAPDF"
        raw_fs = float(settings.get("font_size", 54))
        angle = float(settings.get("angle", 45))
        opacity = float(settings.get("opacity", 30)) / 100.0
        color_hex = settings.get("color", "#c0392b")
        col = QColor(color_hex)
        r_norm, g_norm, b_norm = col.redF(), col.greenF(), col.blueF()
        autofit = bool(settings.get("autofit", True))

        font_key = settings.get("font_family", "helv")
        import os
        from src.ui.direct_text_editor import get_unicode_font_path
        wm_font_file = get_unicode_font_path(font_key, bold=True, italic=False)
        if not wm_font_file or not os.path.isfile(wm_font_file):
            wm_font_file = get_unicode_font_path("Arial", bold=True, italic=False)

        for page_idx in pages_to_process:
            if page_idx < 0 or page_idx >= total_pages:
                continue
            page = doc[page_idx]
            if page.rotation != 0:
                page.remove_rotation()
            rect = page.rect
            cx = rect.x0 + rect.width * nx
            cy = rect.y0 + rect.height * ny

            if autofit:
                safe_fs = calculate_safe_watermark_fontsize(text, angle, rect.width, rect.height, margin=35.0, font_key=font_key)
                eff_fs = min(raw_fs, safe_fs)
            else:
                hard_max_fs = calculate_safe_watermark_fontsize(text, angle, rect.width, rect.height, margin=10.0, font_key=font_key) * 1.05
                eff_fs = min(raw_fs, hard_max_fs)

            wm_font_name = "F_AuraPDF_WM"
            if wm_font_file and os.path.isfile(wm_font_file):
                try:
                    page.insert_font(fontname=wm_font_name, fontfile=wm_font_file)
                except Exception:
                    wm_font_name = "helv"
            else:
                wm_font_name = "helv"

            try:
                if wm_font_file and os.path.isfile(wm_font_file):
                    text_len = fitz.Font(fontfile=wm_font_file).text_length(text, fontsize=eff_fs)
                else:
                    text_len = fitz.get_text_length(text, fontsize=eff_fs)
            except Exception:
                text_len = fitz.get_text_length(text, fontsize=eff_fs)
            pt = fitz.Point(cx - text_len / 2.0, cy + eff_fs / 3.0)

            before_contents = set(page.get_contents())
            page.insert_text(
                pt,
                text,
                fontname=wm_font_name,
                fontfile=wm_font_file if wm_font_name == "F_AuraPDF_WM" else None,
                fontsize=eff_fs,
                color=(r_norm, g_norm, b_norm),
                fill_opacity=opacity,
                morph=(fitz.Point(cx, cy), fitz.Matrix(angle)),
                overlay=overlay
            )
            after_contents = page.get_contents()
            new_xrefs = [x for x in after_contents if x not in before_contents]
            for x in new_xrefs:
                try:
                    orig_st = doc.xref_stream(x)
                    doc.update_stream(x, b'% AuraPDF Watermark Stream\n' + orig_st)
                except Exception:
                    pass

        try:
            if hasattr(doc, 'subset_fonts'):
                doc.subset_fonts()
        except Exception:
            pass
    else:
        img_path = settings.get("img_path", "").strip()
        if img_path and os.path.isfile(img_path):
            base_img = QImage(img_path)
            if not base_img.isNull():
                scale = float(settings.get("scale", 50)) / 100.0
                opacity = float(settings.get("opacity", 35)) / 100.0
                angle = float(settings.get("angle", 0))

                if angle != 0:
                    rot_matrix = QTransform().rotate(-angle)
                    rotated_base = base_img.transformed(rot_matrix, Qt.TransformationMode.SmoothTransformation)
                else:
                    rotated_base = base_img

                res_img = QImage(rotated_base.size(), QImage.Format.Format_ARGB32)
                res_img.fill(Qt.GlobalColor.transparent)
                p = QPainter(res_img)
                p.setOpacity(opacity)
                p.drawImage(0, 0, rotated_base)
                p.end()

                ba = QByteArray()
                buf = QBuffer(ba)
                buf.open(QIODevice.OpenModeFlag.WriteOnly)
                res_img.save(buf, "PNG")
                img_stream = bytes(ba)

                orig_w = max(1, base_img.width())
                for page_idx in pages_to_process:
                    if page_idx < 0 or page_idx >= total_pages:
                        continue
                    page = doc[page_idx]
                    if page.rotation != 0:
                        page.remove_rotation()
                    rect = page.rect
                    cx = rect.x0 + rect.width * nx
                    cy = rect.y0 + rect.height * ny

                    lock_aspect = settings.get("lock_aspect", True)
                    if lock_aspect:
                        scale_factor = (rect.width * scale) / float(orig_w)
                        bb_w = rotated_base.width() * scale_factor
                        bb_h = rotated_base.height() * scale_factor
                    else:
                        pw_ref = float(settings.get("page_w", 595.0))
                        ph_ref = float(settings.get("page_h", 842.0))
                        custom_w = float(settings.get("img_w", 298.0))
                        custom_h = float(settings.get("img_h", 200.0))
                        bb_w = (custom_w / max(1.0, pw_ref)) * rect.width
                        bb_h = (custom_h / max(1.0, ph_ref)) * rect.height

                    img_rect = fitz.Rect(cx - bb_w / 2.0, cy - bb_h / 2.0, cx + bb_w / 2.0, cy + bb_h / 2.0)
                    before_contents = set(page.get_contents())
                    page.insert_image(img_rect, stream=img_stream, overlay=overlay)
                    after_contents = page.get_contents()
                    new_xrefs = [x for x in after_contents if x not in before_contents]
                    for x in new_xrefs:
                        try:
                            orig_st = doc.xref_stream(x)
                            doc.update_stream(x, b'% AuraPDF Watermark Stream\n' + orig_st)
                        except Exception:
                            pass

    # Embed watermark metadata
    import json
    wdata = json.dumps({"version": 2, "settings": settings}, ensure_ascii=False).encode('utf-8')
    try:
        if hasattr(doc, 'embfile_names'):
            if 'aurapdf_watermark.json' in doc.embfile_names():
                doc.embfile_del('aurapdf_watermark.json')
            doc.embfile_add('aurapdf_watermark.json', wdata, filename='aurapdf_watermark.json', ufilename='aurapdf_watermark.json', desc='AuraPDF Watermark Metadata')
    except Exception:
        pass


class PDFWatermarkDialog(QDialog):
    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.watermark_color = QColor("#c0392b")

        self.setWindowTitle(LanguageManager.tr("watermark_dialog_title"))
        self.resize(990, 840)
        self.setMinimumSize(700, 500)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        root = QHBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(12)

        left_container = QWidget()
        left_v = QVBoxLayout(left_container)
        left_v.setContentsMargins(0, 0, 0, 0)
        left_v.setSpacing(8)

        self.lbl_editing_badge = ModernInfoCard(
            text=LanguageManager.tr("watermark_editing_badge"),
            parent=self
        )
        self.lbl_editing_badge.setVisible(False)
        left_v.addWidget(self.lbl_editing_badge)

        # Kaynak PDF
        grp_src = QGroupBox(LanguageManager.tr("encrypt_source_file"))
        grp_src.setStyleSheet("QGroupBox{font-weight:600;padding-top:4px;}")
        h_src_v = QVBoxLayout(grp_src)
        h_src_v.setContentsMargins(8, 8, 8, 8)
        h_src_l = QHBoxLayout()
        self.txt_src = QLineEdit()
        self.txt_src.setPlaceholderText(LanguageManager.tr("split_btn_browse") + "...")
        btn_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src.setFixedWidth(80)
        btn_src.clicked.connect(self._browse_source)
        h_src_l.addWidget(self.txt_src)
        h_src_l.addWidget(btn_src)
        h_src_v.addLayout(h_src_l)
        left_v.addWidget(grp_src)

        # Sekmeler
        self.tab_widget = QTabWidget()
        self.tab_widget.setStyleSheet("QTabWidget::pane{border:none;}")

        # ---- TAB 1: Metin ----
        tab_text = QWidget()
        tab_text_scroll = QScrollArea()
        tab_text_scroll.setWidgetResizable(True)
        tab_text_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        tab_text_scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        tab_text_scroll.setStyleSheet("QScrollArea{background:transparent;border:none;}")
        tab_text_inner = QWidget()
        tab_text_inner.setStyleSheet("background:transparent;")
        tv = QVBoxLayout(tab_text_inner)
        tv.setContentsMargins(4, 8, 4, 8)
        tv.setSpacing(10)

        grp_wtext = QGroupBox(LanguageManager.tr("watermark_text_label"))
        grp_wtext.setStyleSheet("QGroupBox{font-weight:600;}")
        v_wtext = QVBoxLayout(grp_wtext)
        v_wtext.setContentsMargins(8, 8, 8, 8)
        v_wtext.setSpacing(6)
        self.txt_watermark = QLineEdit(LanguageManager.tr("watermark_preset_confidential"))
        self.txt_watermark.setMinimumHeight(32)
        v_wtext.addWidget(self.txt_watermark)
        lbl_pres = QLabel(LanguageManager.tr("watermark_presets_label"))
        lbl_pres.setStyleSheet("font-size:11px;font-weight:500;")
        v_wtext.addWidget(lbl_pres)
        grid_pre = QGridLayout()
        grid_pre.setSpacing(5)
        preset_keys = [
            "watermark_preset_confidential", "watermark_preset_draft",
            "watermark_preset_copy",         "watermark_preset_urgent",
            "watermark_preset_approved",     "watermark_preset_private",
        ]
        for i, pk in enumerate(preset_keys):
            pname = LanguageManager.tr(pk)
            b = QPushButton(pname)
            b.setProperty("class", "PresetBtn")
            b.setStyleSheet("QPushButton{padding:2px 8px;font-size:11px;min-height:22px;}")
            b.clicked.connect(lambda _, t=pname: (self.txt_watermark.setText(t), self.update_preview()))
            grid_pre.addWidget(b, i // 3, i % 3)
        v_wtext.addLayout(grid_pre)
        tv.addWidget(grp_wtext)

        grp_style = QGroupBox(LanguageManager.tr("watermark_font_family"))
        grp_style.setStyleSheet("QGroupBox{font-weight:600;}")
        v_style = QVBoxLayout(grp_style)
        v_style.setContentsMargins(8, 8, 8, 8)
        v_style.setSpacing(6)
        h_font_row = QHBoxLayout()
        h_font_row.addWidget(QLabel(LanguageManager.tr("watermark_font_family")))
        self.cmb_font_family = QComboBox()
        self.cmb_font_family.addItem("Helvetica / Arial (Sans-Serif)", "helv")
        self.cmb_font_family.addItem("Times-Roman (Serif)", "tiro")
        self.cmb_font_family.addItem("Courier (Monospace)", "couri")
        self.cmb_font_family.currentIndexChanged.connect(self.update_preview)
        h_font_row.addWidget(self.cmb_font_family, stretch=1)
        v_style.addLayout(h_font_row)
        h_color_size = QHBoxLayout()
        h_color_size.addWidget(QLabel(LanguageManager.tr("watermark_color_label")))
        self.btn_color = QPushButton("      ")
        self.btn_color.setFixedWidth(46)
        self.btn_color.setFixedHeight(26)
        self._update_color_btn()
        self.btn_color.clicked.connect(self._pick_color)
        h_color_size.addWidget(self.btn_color)
        h_color_size.addSpacing(12)
        h_color_size.addWidget(QLabel(LanguageManager.tr("watermark_font_size")))
        self.spn_fontsize = ModernSpinBox(min_val=12, max_val=200, default_val=54, suffix=" pt", parent=self)
        self.spn_fontsize.setFixedWidth(92)
        self.spn_fontsize.valueChanged.connect(self.update_preview)
        h_color_size.addWidget(self.spn_fontsize)
        h_color_size.addSpacing(8)
        self.chk_autofit = QCheckBox(LanguageManager.tr("watermark_autofit"))
        self.chk_autofit.setChecked(True)
        self.chk_autofit.setStyleSheet("")
        self.chk_autofit.toggled.connect(self.update_preview)
        h_color_size.addWidget(self.chk_autofit)
        h_color_size.addStretch()
        v_style.addLayout(h_color_size)
        tv.addWidget(grp_style)

        grp_rot = QGroupBox(LanguageManager.tr("watermark_rotation_label"))
        grp_rot.setStyleSheet("QGroupBox{font-weight:600;}")
        v_rot = QVBoxLayout(grp_rot)
        v_rot.setContentsMargins(8, 8, 8, 8)
        v_rot.setSpacing(6)
        h_rot = QHBoxLayout()
        self.cmb_rotation = QComboBox()
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_45_up"),   45)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_45_down"), -45)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_0"),       0)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_90_up"),   90)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_90_down"), -90)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_180"),     180)
        self.cmb_rotation.addItem(LanguageManager.tr("watermark_rot_custom"),  "custom")
        self.cmb_rotation.currentIndexChanged.connect(self._on_rotation_changed)
        h_rot.addWidget(self.cmb_rotation, stretch=1)
        self.spn_custom_angle = ModernSpinBox(min_val=-180, max_val=180, default_val=45, suffix="\u00b0", parent=self)
        self.spn_custom_angle.setFixedWidth(92)
        self.spn_custom_angle.setVisible(False)
        self.spn_custom_angle.valueChanged.connect(self.update_preview)
        self.sld_custom_angle = QSlider(Qt.Orientation.Horizontal)
        self.sld_custom_angle.setRange(-180, 180)
        self.sld_custom_angle.setValue(45)
        self.sld_custom_angle.setVisible(False)
        self.sld_custom_angle.valueChanged.connect(
            lambda v: (self.spn_custom_angle.setValue(v) if self.spn_custom_angle.value() != v else None, self.update_preview()))
        self.spn_custom_angle.valueChanged.connect(
            lambda v: self.sld_custom_angle.setValue(v) if self.sld_custom_angle.value() != v else None)
        h_rot.addWidget(self.sld_custom_angle)
        h_rot.addWidget(self.spn_custom_angle)
        v_rot.addLayout(h_rot)
        h_op = QHBoxLayout()
        self.lbl_opacity = QLabel(LanguageManager.tr("watermark_opacity_label", val=30))
        self.lbl_opacity.setFixedWidth(120)
        self.sld_opacity = QSlider(Qt.Orientation.Horizontal)
        self.sld_opacity.setRange(5, 100)
        self.sld_opacity.setValue(30)
        self.sld_opacity.valueChanged.connect(
            lambda v: (self.lbl_opacity.setText(LanguageManager.tr("watermark_opacity_label", val=v)), self.update_preview()))
        h_op.addWidget(self.lbl_opacity)
        h_op.addWidget(self.sld_opacity)
        v_rot.addLayout(h_op)
        tv.addWidget(grp_rot)

        grp_pages_layer = QGroupBox(LanguageManager.tr("watermark_page_range_group"))
        grp_pages_layer.setStyleSheet("QGroupBox{font-weight:600;}")
        pg_v = QVBoxLayout(grp_pages_layer)
        pg_v.setContentsMargins(8, 8, 8, 8)
        pg_v.setSpacing(6)
        h_pg = QHBoxLayout()
        self.rad_all_pages = QRadioButton(LanguageManager.tr("watermark_pages_all"))
        self.rad_all_pages.setChecked(True)
        tab_doc_init = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        cur_p_idx = getattr(tab_doc_init, 'current_page_idx', getattr(tab_doc_init, 'current_page', 0)) if tab_doc_init else 0
        cur_p_init = (cur_p_idx + 1) if (tab_doc_init and tab_doc_init.doc and len(tab_doc_init.doc) > 0) else 1
        self.rad_current_page = QRadioButton(f"{LanguageManager.tr('watermark_pages_current', default='Yalnızca Mevcut Sayfa')} ({cur_p_init})")
        self.rad_first_page = QRadioButton(LanguageManager.tr("watermark_pages_first"))
        self.rad_custom_pages = QRadioButton(LanguageManager.tr("watermark_pages_custom"))
        self.txt_custom_pages = QLineEdit()
        self.txt_custom_pages.setPlaceholderText("1-5, 8, 12")
        self.txt_custom_pages.setEnabled(False)
        self.rad_custom_pages.toggled.connect(self.txt_custom_pages.setEnabled)
        h_pg.addWidget(self.rad_all_pages)
        h_pg.addWidget(self.rad_current_page)
        h_pg.addWidget(self.rad_first_page)
        h_pg.addWidget(self.rad_custom_pages)
        h_pg.addWidget(self.txt_custom_pages, stretch=1)
        pg_v.addLayout(h_pg)
        h_layer = QHBoxLayout()
        self.rad_layer_overlay  = QRadioButton(LanguageManager.tr("watermark_layer_overlay"))
        self.rad_layer_underlay = QRadioButton(LanguageManager.tr("watermark_layer_underlay"))
        self.rad_layer_overlay.setChecked(True)
        self.rad_layer_overlay.toggled.connect(self.update_preview)
        self.rad_layer_underlay.toggled.connect(self.update_preview)
        lbl_layer_title = QLabel(LanguageManager.tr("watermark_layer_group") + ":")
        lbl_layer_title.setStyleSheet("font-weight:500;")
        h_layer.addWidget(lbl_layer_title)
        h_layer.addWidget(self.rad_layer_overlay)
        h_layer.addWidget(self.rad_layer_underlay)
        h_layer.addStretch()
        pg_v.addLayout(h_layer)
        tv.addWidget(grp_pages_layer)
        tv.addStretch()

        tab_text_scroll.setWidget(tab_text_inner)
        tl = QVBoxLayout(tab_text)
        tl.setContentsMargins(0, 0, 0, 0)
        tl.addWidget(tab_text_scroll)
        self.tab_widget.addTab(tab_text, LanguageManager.tr("watermark_tab_text"))

        # ---- TAB 2: Gorsel ----
        tab_img = QWidget()
        tab_img_scroll = QScrollArea()
        tab_img_scroll.setWidgetResizable(True)
        tab_img_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        tab_img_scroll.setFrameShape(QScrollArea.Shape.NoFrame)
        tab_img_scroll.setStyleSheet("QScrollArea{background:transparent;border:none;}")
        tab_img_inner = QWidget()
        tab_img_inner.setStyleSheet("background:transparent;")
        iv = QVBoxLayout(tab_img_inner)
        iv.setContentsMargins(4, 8, 4, 8)
        iv.setSpacing(10)

        grp_imgfile = QGroupBox(LanguageManager.tr("watermark_image_file"))
        grp_imgfile.setStyleSheet("QGroupBox{font-weight:600;}")
        imgf_v = QVBoxLayout(grp_imgfile)
        imgf_v.setContentsMargins(8, 8, 8, 8)
        h_ifile = QHBoxLayout()
        self.txt_img_src = QLineEdit()
        self.txt_img_src.setPlaceholderText("PNG, JPG, BMP...")
        btn_img_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_img_src.setFixedWidth(80)
        btn_img_src.clicked.connect(self._browse_image)
        h_ifile.addWidget(self.txt_img_src)
        h_ifile.addWidget(btn_img_src)
        imgf_v.addLayout(h_ifile)
        iv.addWidget(grp_imgfile)

        grp_imgsize = QGroupBox(LanguageManager.tr("watermark_scale_label"))
        grp_imgsize.setStyleSheet("QGroupBox{font-weight:600;}")
        v_is = QVBoxLayout(grp_imgsize)
        v_is.setContentsMargins(8, 8, 8, 8)
        v_is.setSpacing(6)
        h_scale = QHBoxLayout()
        self.lbl_img_scale = QLabel(LanguageManager.tr("watermark_image_scale", val=50))
        self.lbl_img_scale.setFixedWidth(100)
        self.sld_img_scale = QSlider(Qt.Orientation.Horizontal)
        self.sld_img_scale.setRange(5, 300)
        self.sld_img_scale.setValue(50)
        self.spn_img_scale = ModernSpinBox(min_val=5, max_val=300, default_val=50, suffix=" %", parent=self)
        self.spn_img_scale.setFixedWidth(92)
        self.sld_img_scale.valueChanged.connect(self._on_slider_scale_changed)
        self.spn_img_scale.valueChanged.connect(self._on_spin_scale_changed)
        h_scale.addWidget(self.lbl_img_scale)
        h_scale.addWidget(self.sld_img_scale)
        h_scale.addWidget(self.spn_img_scale)
        v_is.addLayout(h_scale)
        g_sp = QGridLayout()
        g_sp.setSpacing(5)
        preset_scales = [
            ("watermark_preset_small", 15), ("watermark_preset_normal", 30),
            ("watermark_preset_medium", 50), ("watermark_preset_large",  80),
            ("watermark_preset_full",  100),
        ]
        for i, (pk, pv) in enumerate(preset_scales):
            b = QPushButton(LanguageManager.tr(pk))
            b.setProperty("class", "PresetBtn")
            b.setStyleSheet("QPushButton{padding:2px 8px;font-size:11px;min-height:22px;}")
            b.clicked.connect(lambda _, v=pv: self.spn_img_scale.setValue(v))
            g_sp.addWidget(b, i // 3, i % 3)
        v_is.addLayout(g_sp)
        h_dims = QHBoxLayout()
        h_dims.addWidget(QLabel(LanguageManager.tr("watermark_img_w")))
        self.spn_img_w = ModernSpinBox(min_val=10, max_val=3000, default_val=298, suffix=" pt", parent=self)
        self.spn_img_w.setFixedWidth(105)
        self.spn_img_w.valueChanged.connect(self._on_dim_w_changed)
        h_dims.addWidget(self.spn_img_w)
        h_dims.addSpacing(10)
        h_dims.addWidget(QLabel(LanguageManager.tr("watermark_img_h")))
        self.spn_img_h = ModernSpinBox(min_val=10, max_val=3000, default_val=200, suffix=" pt", parent=self)
        self.spn_img_h.setFixedWidth(105)
        self.spn_img_h.valueChanged.connect(self._on_dim_h_changed)
        h_dims.addWidget(self.spn_img_h)
        h_dims.addSpacing(10)
        self.chk_lock_aspect = QCheckBox(LanguageManager.tr("watermark_lock_ratio"))
        self.chk_lock_aspect.setChecked(True)
        self.chk_lock_aspect.setStyleSheet("")
        self.chk_lock_aspect.toggled.connect(self._on_lock_aspect_toggled)
        h_dims.addWidget(self.chk_lock_aspect)
        h_dims.addStretch()
        v_is.addLayout(h_dims)
        iv.addWidget(grp_imgsize)

        grp_irot = QGroupBox(LanguageManager.tr("watermark_rotation_group"))
        grp_irot.setStyleSheet("QGroupBox{font-weight:600;}")
        v_irot = QVBoxLayout(grp_irot)
        v_irot.setContentsMargins(8, 8, 8, 8)
        v_irot.setSpacing(6)
        h_irot = QHBoxLayout()
        self.lbl_img_rotation = QLabel(LanguageManager.tr("watermark_angle_label", val=0))
        self.lbl_img_rotation.setFixedWidth(100)
        self.sld_img_rotation = QSlider(Qt.Orientation.Horizontal)
        self.sld_img_rotation.setRange(-180, 180)
        self.sld_img_rotation.setValue(0)
        self.spn_img_rotation = ModernSpinBox(min_val=-180, max_val=180, default_val=0, suffix="\u00b0", parent=self)
        self.spn_img_rotation.setFixedWidth(92)
        self.sld_img_rotation.valueChanged.connect(self._on_slider_img_rot_changed)
        self.spn_img_rotation.valueChanged.connect(self._on_spin_img_rot_changed)
        h_irot.addWidget(self.lbl_img_rotation)
        h_irot.addWidget(self.sld_img_rotation)
        h_irot.addWidget(self.spn_img_rotation)
        v_irot.addLayout(h_irot)
        g_rp = QGridLayout()
        g_rp.setSpacing(5)
        rot_presets = [
            ("watermark_rot_0", 0),           ("watermark_rot_45_up", 45),
            ("watermark_rot_45_down", -45),   ("watermark_rot_90_up", 90),
            ("watermark_rot_90_down", -90),   ("watermark_rot_180", 180),
        ]
        for i, (rk, rv) in enumerate(rot_presets):
            b = QPushButton(LanguageManager.tr(rk))
            b.setProperty("class", "PresetBtn")
            b.setStyleSheet("QPushButton{padding:2px 8px;font-size:11px;min-height:22px;}")
            b.clicked.connect(lambda _, v=rv: self.spn_img_rotation.setValue(v))
            g_rp.addWidget(b, i // 3, i % 3)
        v_irot.addLayout(g_rp)
        h_iop = QHBoxLayout()
        self.lbl_img_op = QLabel(LanguageManager.tr("watermark_opacity_label", val=35))
        self.lbl_img_op.setFixedWidth(100)
        self.sld_img_op = QSlider(Qt.Orientation.Horizontal)
        self.sld_img_op.setRange(5, 100)
        self.sld_img_op.setValue(35)
        self.sld_img_op.valueChanged.connect(
            lambda v: (self.lbl_img_op.setText(LanguageManager.tr("watermark_opacity_label", val=v)), self.update_preview()))
        h_iop.addWidget(self.lbl_img_op)
        h_iop.addWidget(self.sld_img_op)
        v_irot.addLayout(h_iop)
        iv.addWidget(grp_irot)
        iv.addStretch()

        tab_img_scroll.setWidget(tab_img_inner)
        il = QVBoxLayout(tab_img)
        il.setContentsMargins(0, 0, 0, 0)
        il.addWidget(tab_img_scroll)
        self.tab_widget.addTab(tab_img, LanguageManager.tr("watermark_tab_image"))

        left_v.addWidget(self.tab_widget, stretch=1)

        # Hedef PDF
        grp_dst = QGroupBox(LanguageManager.tr("encrypt_target_file"))
        grp_dst.setStyleSheet("QGroupBox{font-weight:600;padding-top:4px;}")
        h_dst_v = QVBoxLayout(grp_dst)
        h_dst_v.setContentsMargins(8, 8, 8, 8)
        h_dst_l = QHBoxLayout()
        self.txt_dst = QLineEdit()
        btn_dst = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst.setFixedWidth(80)
        btn_dst.clicked.connect(self._browse_target)
        h_dst_l.addWidget(self.txt_dst)
        h_dst_l.addWidget(btn_dst)
        h_dst_v.addLayout(h_dst_l)
        left_v.addWidget(grp_dst)

        # Alt Butonlar
        v_bot = QVBoxLayout()
        v_bot.setSpacing(6)
        v_bot.setContentsMargins(0, 2, 0, 0)
        h_bot_top = QHBoxLayout()
        self.chk_open_tab = QCheckBox(LanguageManager.tr("encrypt_chk_open"))
        self.chk_open_tab.setChecked(True)
        h_bot_top.addWidget(self.chk_open_tab)
        h_bot_top.addStretch()
        self.btn_remove_watermark = QPushButton(LanguageManager.tr("watermark_btn_remove"))
        self.btn_remove_watermark.setStyleSheet(
            "QPushButton { background-color: rgba(239, 68, 68, 0.15); color: #fca5a5; border: 1px solid rgba(239, 68, 68, 0.35); "
            "border-radius: 4px; font-weight: bold; padding: 5px 12px; min-height: 26px; } "
            "QPushButton:hover { background-color: rgba(239, 68, 68, 0.28); border-color: #ef4444; color: #ffffff; }"
        )
        self.btn_remove_watermark.clicked.connect(self.execute_remove_watermark)
        self.btn_remove_watermark.setVisible(False)
        h_bot_top.addWidget(self.btn_remove_watermark)
        v_bot.addLayout(h_bot_top)
        h_bot_actions = QHBoxLayout()
        h_bot_actions.setSpacing(8)
        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.setMinimumHeight(30)
        btn_cancel.clicked.connect(self.reject)
        h_bot_actions.addWidget(btn_cancel)
        h_bot_actions.addStretch()
        self.btn_save_as = QPushButton(LanguageManager.tr("watermark_btn_save_as"))
        self.btn_save_as.setMinimumHeight(30)
        self.btn_save_as.clicked.connect(self.execute_save_as_watermark)
        h_bot_actions.addWidget(self.btn_save_as)
        self.btn_add_new = QPushButton(LanguageManager.tr("watermark_btn_add_new"))
        self.btn_add_new.setMinimumHeight(30)
        self.btn_add_new.clicked.connect(lambda: self.execute_apply_watermark(mode="add_new"))
        self.btn_add_new.setVisible(False)
        h_bot_actions.addWidget(self.btn_add_new)
        self.btn_apply = QPushButton(LanguageManager.tr("watermark_btn_apply_doc"))
        self.btn_apply.setObjectName("PrimaryBtn")
        self.btn_apply.setMinimumHeight(30)
        self.btn_apply.clicked.connect(lambda: self.execute_apply_watermark(mode="auto"))
        h_bot_actions.addWidget(self.btn_apply)
        v_bot.addLayout(h_bot_actions)
        left_v.addLayout(v_bot)

        # Sag: Canli Onizleme
        grp_preview = QGroupBox(LanguageManager.tr("preview_title"))
        grp_preview.setFixedWidth(295)
        grp_preview.setStyleSheet(
            "QGroupBox { border: 1px solid rgba(255, 255, 255, 0.1); "
            "border-radius: 6px; margin-top: 10px; padding-top: 14px; }"
        )
        v_prev = QVBoxLayout(grp_preview)
        v_prev.setContentsMargins(10, 12, 10, 12)
        v_prev.setSpacing(8)
        self.lbl_preview_info = QLabel(LanguageManager.tr("preview_page_1"))
        self.lbl_preview_info.setStyleSheet("font-size:11px;")
        v_prev.addWidget(self.lbl_preview_info)
        frame_box = QFrame()
        frame_box.setObjectName("PreviewFrame")
        frame_box.setProperty("class", "preview-card")
        v_box = QVBoxLayout(frame_box)
        v_box.setContentsMargins(6, 6, 6, 6)
        v_box.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_preview = InteractiveWatermarkPreviewLabel()
        self.lbl_preview.setFixedSize(242, 338)
        self.lbl_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_preview.setStyleSheet("border: 1px solid rgba(255, 255, 255, 0.15); border-radius: 2px;")
        self.lbl_preview.positionChanged.connect(lambda pos: self.update_preview())
        v_box.addWidget(self.lbl_preview)
        v_prev.addWidget(frame_box)
        self.preview_current_page = getattr(tab_doc_init, 'current_page_idx', getattr(tab_doc_init, 'current_page', 0)) if tab_doc_init else 0
        self.preview_total_pages  = 1
        h_nav = QHBoxLayout()
        h_nav.setSpacing(6)
        self.btn_prev_page = QPushButton(LanguageManager.tr("preview_prev"))
        self.btn_prev_page.setProperty("class", "NavBtn")
        self.btn_prev_page.setIcon(ModernIconFactory.icon_chevron_left())
        self.btn_prev_page.setMinimumWidth(84)
        self.btn_prev_page.clicked.connect(self._preview_prev_page)
        self.spn_preview_page = QSpinBox()
        self.spn_preview_page.setObjectName("PreviewPageSpin")
        self.spn_preview_page.setButtonSymbols(QAbstractSpinBox.ButtonSymbols.NoButtons)
        self.spn_preview_page.setRange(1, 1)
        self.spn_preview_page.setValue(1)
        self.spn_preview_page.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.spn_preview_page.setFixedWidth(52)
        self.spn_preview_page.setStyleSheet(
            "QSpinBox#PreviewPageSpin { background-color: #1a1a1e; color: #f1f5f9; border: 1px solid rgba(255, 255, 255, 0.12); "
            "border-radius: 4px; padding: 2px; font-size: 12px; font-weight: bold; } "
            "QSpinBox#PreviewPageSpin:focus { border: 1px solid #38bdf8; }"
        )
        self.spn_preview_page.valueChanged.connect(self._preview_page_spin_changed)
        self.lbl_total_preview_pages = QLabel("/ 1")
        self.lbl_total_preview_pages.setStyleSheet("font-size:11px;font-weight:500;")
        self.btn_next_page = QPushButton(LanguageManager.tr("preview_next"))
        self.btn_next_page.setProperty("class", "NavBtn")
        self.btn_next_page.setIcon(ModernIconFactory.icon_chevron_right())
        self.btn_next_page.setMinimumWidth(84)
        self.btn_next_page.clicked.connect(self._preview_next_page)
        h_nav.addWidget(self.btn_prev_page)
        h_nav.addWidget(self.spn_preview_page)
        h_nav.addWidget(self.lbl_total_preview_pages)
        h_nav.addWidget(self.btn_next_page)
        v_prev.addLayout(h_nav)
        lbl_drag_hint = QLabel(LanguageManager.tr("watermark_drag_hint"))
        lbl_drag_hint.setStyleSheet("font-size:10px;font-weight:500;")
        lbl_drag_hint.setWordWrap(True)
        lbl_drag_hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        v_prev.addWidget(lbl_drag_hint)
        btn_reset_pos = QPushButton(LanguageManager.tr("watermark_reset_pos"))
        btn_reset_pos.setProperty("class", "PresetBtn")
        btn_reset_pos.clicked.connect(self.lbl_preview.reset_to_center)
        v_prev.addWidget(btn_reset_pos)
        v_prev.addStretch()

        root.addWidget(left_container, stretch=1)
        root.addWidget(grp_preview,    stretch=0)

        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        if tab_doc and tab_doc.doc and len(tab_doc.doc) > 0:
            if getattr(tab_doc, 'file_path', None):
                p = tab_doc.file_path
                self.txt_src.setText(p)
                base_dir = os.path.dirname(p)
                base_name, _ = os.path.splitext(os.path.basename(p))
                self.txt_dst.setText(os.path.join(base_dir, f"{base_name}_filigranli.pdf"))
            else:
                doc_title = getattr(tab_doc, 'custom_title', None) or LanguageManager.tr("new_blank_doc")
                self.txt_src.setText(f"[{doc_title}]")
                desktop = os.path.join(os.path.expanduser("~"), "Desktop")
                self.txt_dst.setText(os.path.join(desktop, f"{doc_title}_filigranli.pdf"))

        self.txt_src.textChanged.connect(self._on_source_changed)
        self.txt_watermark.textChanged.connect(self.update_preview)
        self.spn_fontsize.valueChanged.connect(self.update_preview)
        self.sld_opacity.valueChanged.connect(self.update_preview)
        self.txt_img_src.textChanged.connect(self._on_img_src_changed)
        self.sld_img_op.valueChanged.connect(self.update_preview)
        self.tab_widget.currentChanged.connect(self.update_preview)
        self.lbl_preview.scaleChanged.connect(self._on_preview_scale_changed)
        self.lbl_preview.angleChanged.connect(self._on_preview_angle_changed)
        self.lbl_preview.widthChanged.connect(self._on_preview_width_changed)
        self.lbl_preview.heightChanged.connect(self._on_preview_height_changed)

        self.has_existing_watermark = False
        tab_doc2 = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        if tab_doc2 and (getattr(tab_doc2, 'has_watermark', False) or getattr(tab_doc2, 'watermark_settings', None)):
            self.has_existing_watermark = True
            self.lbl_editing_badge.setVisible(True)
            self.btn_remove_watermark.setVisible(True)
            self.btn_add_new.setVisible(True)
            self.btn_apply.setText(LanguageManager.tr("watermark_btn_update"))
            if tab_doc2.watermark_settings:
                self._restore_saved_watermark_settings(tab_doc2.watermark_settings)

        self.update_preview()

    def _is_active_tab_doc(self, src_path=None) -> bool:
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        if not tab_doc or tab_doc.doc is None or len(tab_doc.doc) == 0:
            return False
        if src_path is None:
            src_path = self.txt_src.text().strip()
        if getattr(tab_doc, 'file_path', None) and os.path.isfile(src_path):
            return os.path.abspath(tab_doc.file_path) == os.path.abspath(src_path)
        if not getattr(tab_doc, 'file_path', None) and (src_path.startswith("[") or not src_path):
            return True
        return False

    def _on_preview_width_changed(self, new_w):
        val = max(10, min(3000, new_w))
        self.spn_img_w.blockSignals(True)
        self.spn_img_w.setValue(val)
        self.spn_img_w.blockSignals(False)
        self._on_dim_w_changed(val)

    def _on_preview_height_changed(self, new_h):
        val = max(10, min(3000, new_h))
        self.spn_img_h.blockSignals(True)
        self.spn_img_h.setValue(val)
        self.spn_img_h.blockSignals(False)
        self._on_dim_h_changed(val)

    def _on_preview_scale_changed(self, val):
        if self.tab_widget.currentIndex() == 0:
            fs = max(8, min(144, val))
            self.spn_fontsize.blockSignals(True)
            self.spn_fontsize.setValue(fs)
            self.spn_fontsize.blockSignals(False)
            self.update_preview()
        else:
            self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=val))
            self.spn_img_scale.blockSignals(True)
            self.spn_img_scale.setValue(val)
            self.spn_img_scale.blockSignals(False)
            self.sld_img_scale.blockSignals(True)
            self.sld_img_scale.setValue(val)
            self.sld_img_scale.blockSignals(False)
            self.lbl_preview.current_scale_value = val
            self._sync_dims_from_scale(val)
            self.update_preview()

    def _on_preview_angle_changed(self, angle: int):
        is_text = (self.tab_widget.currentIndex() == 0)
        if is_text:
            matched = False
            for idx in range(self.cmb_rotation.count()):
                data = self.cmb_rotation.itemData(idx)
                if data == angle:
                    self.cmb_rotation.blockSignals(True)
                    self.cmb_rotation.setCurrentIndex(idx)
                    self.cmb_rotation.blockSignals(False)
                    self.spn_custom_angle.setVisible(False)
                    if hasattr(self, 'sld_custom_angle'):
                        self.sld_custom_angle.setVisible(False)
                    matched = True
                    break
            if not matched:
                custom_idx = self.cmb_rotation.findData("custom")
                if custom_idx != -1 and self.cmb_rotation.currentIndex() != custom_idx:
                    self.cmb_rotation.blockSignals(True)
                    self.cmb_rotation.setCurrentIndex(custom_idx)
                    self.cmb_rotation.blockSignals(False)
                self.spn_custom_angle.setVisible(True)
                if hasattr(self, 'sld_custom_angle'):
                    self.sld_custom_angle.setVisible(True)
                self.spn_custom_angle.blockSignals(True)
                self.spn_custom_angle.setValue(angle)
                self.spn_custom_angle.blockSignals(False)
                if hasattr(self, 'sld_custom_angle'):
                    self.sld_custom_angle.blockSignals(True)
                    self.sld_custom_angle.setValue(angle)
                    self.sld_custom_angle.blockSignals(False)
        else:
            if hasattr(self, 'spn_img_rotation'):
                self.spn_img_rotation.blockSignals(True)
                self.spn_img_rotation.setValue(angle)
                self.spn_img_rotation.blockSignals(False)
            if hasattr(self, 'sld_img_rotation'):
                self.sld_img_rotation.blockSignals(True)
                self.sld_img_rotation.setValue(angle)
                self.sld_img_rotation.blockSignals(False)
            if hasattr(self, 'lbl_img_rotation'):
                self.lbl_img_rotation.setText(LanguageManager.tr("watermark_angle_label", val=angle))
        self.update_preview()

    def _get_current_watermark_settings(self) -> dict:
        is_text = (self.tab_widget.currentIndex() == 0)
        norm_pos = getattr(self.lbl_preview, 'norm_pos', QPointF(0.5, 0.5))

        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        curr_p_idx = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0)) if tab_doc else 0

        pages_mode = "all"
        if hasattr(self, 'rad_current_page') and self.rad_current_page.isChecked():
            pages_mode = "current"
        elif self.rad_first_page.isChecked():
            pages_mode = "first"
        elif self.rad_custom_pages.isChecked():
            pages_mode = "custom"

        settings = {
            "type": "text" if is_text else "image",
            "layer": "overlay" if self.rad_layer_overlay.isChecked() else "underlay",
            "pages_mode": pages_mode,
            "current_page_idx": curr_p_idx,
            "custom_pages": self.txt_custom_pages.text().strip(),
            "norm_pos_x": float(norm_pos.x()),
            "norm_pos_y": float(norm_pos.y()),
        }

        if is_text:
            settings.update({
                "text": self.txt_watermark.text().strip(),
                "font_family": self.cmb_font_family.currentData() if hasattr(self, 'cmb_font_family') else "helv",
                "font_size": self.spn_fontsize.value(),
                "angle": self._get_current_angle(),
                "opacity": self.sld_opacity.value(),
                "color": self.watermark_color.name(),
                "autofit": self.chk_autofit.isChecked(),
            })
        else:
            settings.update({
                "img_path": self.txt_img_src.text().strip(),
                "scale": self.sld_img_scale.value(),
                "opacity": self.sld_img_op.value(),
                "angle": self._get_current_img_angle(),
                "lock_aspect": self.chk_lock_aspect.isChecked(),
                "img_w": self.spn_img_w.value(),
                "img_h": self.spn_img_h.value(),
                "page_w": getattr(self, 'preview_page_w', 595.0),
                "page_h": getattr(self, 'preview_page_h', 842.0),
            })
        return settings

    def _restore_saved_watermark_settings(self, s: dict):
        if not isinstance(s, dict):
            return
        wm_type = s.get("type", "text")
        if wm_type == "image":
            self.tab_widget.setCurrentIndex(1)
            if "img_path" in s:
                self.txt_img_src.setText(str(s["img_path"]))
            elif "image_path" in s:
                self.txt_img_src.setText(str(s["image_path"]))
            scale_val = int(s.get("scale", s.get("img_scale", 50)))
            self.sld_img_scale.setValue(scale_val)
            self.spn_img_scale.setValue(scale_val)
            op_val = int(s.get("opacity", s.get("img_opacity", 35)))
            self.sld_img_op.setValue(op_val)
            angle_val = int(s.get("angle", s.get("img_angle", 0)))
            if hasattr(self, 'spn_img_rotation'):
                self.spn_img_rotation.setValue(angle_val)
            if hasattr(self, 'sld_img_rotation'):
                self.sld_img_rotation.setValue(angle_val)
            if "lock_aspect" in s:
                self.chk_lock_aspect.setChecked(bool(s["lock_aspect"]))
            if "img_w" in s:
                self.spn_img_w.setValue(int(s["img_w"]))
            if "img_h" in s:
                self.spn_img_h.setValue(int(s["img_h"]))
        else:
            self.tab_widget.setCurrentIndex(0)
            if "text" in s:
                self.txt_watermark.setText(str(s["text"]))
            if "font_family" in s and hasattr(self, 'cmb_font_family'):
                idx = self.cmb_font_family.findData(s["font_family"])
                if idx >= 0:
                    self.cmb_font_family.setCurrentIndex(idx)
            fs_val = int(s.get("font_size", s.get("fontsize", 54)))
            self.spn_fontsize.setValue(fs_val)
            if "autofit" in s:
                self.chk_autofit.setChecked(bool(s["autofit"]))
            if "color" in s:
                self.watermark_color = QColor(s["color"])
                self._update_color_btn()
            if "opacity" in s:
                self.sld_opacity.setValue(int(s["opacity"]))
            angle_val = int(s.get("angle", 45))
            matched = False
            for idx in range(self.cmb_rotation.count()):
                if self.cmb_rotation.itemData(idx) == angle_val:
                    self.cmb_rotation.setCurrentIndex(idx)
                    matched = True
                    break
            if not matched:
                custom_idx = self.cmb_rotation.findData("custom")
                if custom_idx != -1:
                    self.cmb_rotation.setCurrentIndex(custom_idx)
                self.spn_custom_angle.setValue(angle_val)
                if hasattr(self, 'sld_custom_angle'):
                    self.sld_custom_angle.setValue(angle_val)

        nx = float(s.get("norm_pos_x", 0.5))
        ny = float(s.get("norm_pos_y", 0.5))
        if "norm_pos" in s and isinstance(s["norm_pos"], (list, tuple)) and len(s["norm_pos"]) >= 2:
            nx = float(s["norm_pos"][0])
            ny = float(s["norm_pos"][1])
        self.lbl_preview.norm_pos = QPointF(nx, ny)

        if s.get("layer") == "underlay":
            self.rad_layer_underlay.setChecked(True)
        else:
            self.rad_layer_overlay.setChecked(True)

        pages_mode = s.get("pages_mode", "all")
        if pages_mode == "first":
            self.rad_first_page.setChecked(True)
        elif pages_mode == "custom":
            self.rad_custom_pages.setChecked(True)
            self.txt_custom_pages.setText(str(s.get("custom_pages", "")))
        else:
            self.rad_all_pages.setChecked(True)

    def _on_slider_scale_changed(self, val):
        self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=val))
        self.spn_img_scale.blockSignals(True)
        self.spn_img_scale.setValue(val)
        self.spn_img_scale.blockSignals(False)
        self.lbl_preview.current_scale_value = val
        self._sync_dims_from_scale(val)
        self.update_preview()

    def _on_spin_scale_changed(self, val):
        self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=val))
        self.sld_img_scale.blockSignals(True)
        self.sld_img_scale.setValue(val)
        self.sld_img_scale.blockSignals(False)
        self.lbl_preview.current_scale_value = val
        self._sync_dims_from_scale(val)
        self.update_preview()

    def _get_image_aspect(self) -> float:
        img_path = self.txt_img_src.text().strip()
        if img_path and os.path.isfile(img_path):
            img = QImage(img_path)
            if not img.isNull() and img.width() > 0:
                return float(img.height()) / float(img.width())
        return 0.75

    def _sync_dims_from_scale(self, scale_val):
        page_w = getattr(self, 'preview_page_w', 595.0)
        aspect = self._get_image_aspect()
        target_w = max(10, int(round(page_w * (scale_val / 100.0))))
        target_h = max(10, int(round(target_w * aspect)))
        self.spn_img_w.blockSignals(True)
        self.spn_img_w.setValue(target_w)
        self.spn_img_w.blockSignals(False)
        self.spn_img_h.blockSignals(True)
        self.spn_img_h.setValue(target_h)
        self.spn_img_h.blockSignals(False)

    def _on_dim_w_changed(self, new_w):
        page_w = getattr(self, 'preview_page_w', 595.0)
        scale_pct = max(5, min(300, int(round((new_w / max(1.0, page_w)) * 100.0))))
        aspect = self._get_image_aspect()
        if self.chk_lock_aspect.isChecked():
            new_h = max(10, int(round(new_w * aspect)))
            self.spn_img_h.blockSignals(True)
            self.spn_img_h.setValue(new_h)
            self.spn_img_h.blockSignals(False)
        self.spn_img_scale.blockSignals(True)
        self.spn_img_scale.setValue(scale_pct)
        self.spn_img_scale.blockSignals(False)
        self.sld_img_scale.blockSignals(True)
        self.sld_img_scale.setValue(scale_pct)
        self.sld_img_scale.blockSignals(False)
        self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=scale_pct))
        self.lbl_preview.current_scale_value = scale_pct
        self.update_preview()

    def _on_dim_h_changed(self, new_h):
        aspect = self._get_image_aspect()
        page_w = getattr(self, 'preview_page_w', 595.0)
        if self.chk_lock_aspect.isChecked():
            new_w = max(10, int(round(new_h / max(0.01, aspect))))
            self.spn_img_w.blockSignals(True)
            self.spn_img_w.setValue(new_w)
            self.spn_img_w.blockSignals(False)
        else:
            new_w = self.spn_img_w.value()
        scale_pct = max(5, min(300, int(round((new_w / max(1.0, page_w)) * 100.0))))
        self.spn_img_scale.blockSignals(True)
        self.spn_img_scale.setValue(scale_pct)
        self.spn_img_scale.blockSignals(False)
        self.sld_img_scale.blockSignals(True)
        self.sld_img_scale.setValue(scale_pct)
        self.sld_img_scale.blockSignals(False)
        self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=scale_pct))
        self.lbl_preview.current_scale_value = scale_pct
        self.update_preview()

    def _on_lock_aspect_toggled(self, checked: bool):
        self.lbl_preview.lock_aspect = checked
        if checked:
            aspect = self._get_image_aspect()
            new_h = max(10, int(round(self.spn_img_w.value() * aspect)))
            self.spn_img_h.blockSignals(True)
            self.spn_img_h.setValue(new_h)
            self.spn_img_h.blockSignals(False)
            page_w = getattr(self, 'preview_page_w', 595.0)
            scale_pct = max(5, min(300, int(round((self.spn_img_w.value() / max(1.0, page_w)) * 100.0))))
            self.spn_img_scale.blockSignals(True)
            self.spn_img_scale.setValue(scale_pct)
            self.spn_img_scale.blockSignals(False)
            self.sld_img_scale.blockSignals(True)
            self.sld_img_scale.setValue(scale_pct)
            self.sld_img_scale.blockSignals(False)
            self.lbl_img_scale.setText(LanguageManager.tr("watermark_image_scale", val=scale_pct))
            self.lbl_preview.current_scale_value = scale_pct
        self.update_preview()

    def _on_img_src_changed(self):
        self._sync_dims_from_scale(self.spn_img_scale.value())
        self.update_preview()

    def _on_rotation_changed(self, idx):
        data = self.cmb_rotation.currentData()
        is_custom = (data == "custom")
        self.spn_custom_angle.setVisible(is_custom)
        if hasattr(self, 'sld_custom_angle'):
            self.sld_custom_angle.setVisible(is_custom)
        if not is_custom and isinstance(data, (int, float)):
            self.spn_custom_angle.blockSignals(True)
            self.spn_custom_angle.setValue(int(data))
            self.spn_custom_angle.blockSignals(False)
            if hasattr(self, 'sld_custom_angle'):
                self.sld_custom_angle.blockSignals(True)
                self.sld_custom_angle.setValue(int(data))
                self.sld_custom_angle.blockSignals(False)
        self.update_preview()

    def _get_current_angle(self) -> int:
        data = self.cmb_rotation.currentData()
        if data == "custom":
            return self.spn_custom_angle.value()
        return int(data) if data is not None else 45

    def _get_current_img_angle(self) -> int:
        return self.spn_img_rotation.value() if hasattr(self, 'spn_img_rotation') else 0

    def _on_slider_img_rot_changed(self, val: int):
        if hasattr(self, 'spn_img_rotation') and self.spn_img_rotation.value() != val:
            self.spn_img_rotation.blockSignals(True)
            self.spn_img_rotation.setValue(val)
            self.spn_img_rotation.blockSignals(False)
        if hasattr(self, 'lbl_img_rotation'):
            self.lbl_img_rotation.setText(LanguageManager.tr("watermark_angle_label", val=val))
        self.update_preview()

    def _on_spin_img_rot_changed(self, val: int):
        if hasattr(self, 'sld_img_rotation') and self.sld_img_rotation.value() != val:
            self.sld_img_rotation.blockSignals(True)
            self.sld_img_rotation.setValue(val)
            self.sld_img_rotation.blockSignals(False)
        if hasattr(self, 'lbl_img_rotation'):
            self.lbl_img_rotation.setText(LanguageManager.tr("watermark_angle_label", val=val))
        self.update_preview()

    def _calculate_safe_fontsize(self, text: str, angle_deg: float, page_w: float, page_h: float, margin: float = 35.0, font_key: str = "helv") -> float:
        if not text:
            return 50.0
        rad = math.radians(abs(angle_deg))
        avail_w = max(20.0, page_w - 2 * margin)
        avail_h = max(20.0, page_h - 2 * margin)
        fitz_font = get_watermark_fitz_font(font_key)
        try:
            len100 = fitz_font.text_length(text, fontsize=100.0)
        except Exception:
            len100 = fitz.get_text_length(text, fontsize=100.0)
        unit_w = max(1.0, len100 / 100.0)
        unit_h = 0.85
        bb_w = unit_w * math.cos(rad) + unit_h * math.sin(rad)
        bb_h = unit_w * math.sin(rad) + unit_h * math.cos(rad)
        max_fs = min(avail_w / max(0.001, bb_w), avail_h / max(0.001, bb_h))
        return max(8.0, max_fs)

    def _on_source_changed(self):
        self.preview_current_page = 0
        self.update_preview()

    def _browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_src.setText(path)
            base_dir = os.path.dirname(path)
            base_name, _ = os.path.splitext(os.path.basename(path))
            self.txt_dst.setText(os.path.join(base_dir, f"{base_name}_filigranli.pdf"))
            self.preview_current_page = 0
            self.update_preview()

    def _browse_target(self):
        cur = self.txt_dst.text().strip()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), cur or "belge_filigranli.pdf", "PDF (*.pdf)")
        if path:
            self.txt_dst.setText(path)

    def _browse_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "Görsel Seç", "", "Görseller (*.png *.jpg *.jpeg *.bmp *.webp)")
        if path:
            self.txt_img_src.setText(path)
            self._sync_dims_from_scale(self.spn_img_scale.value())
            self.update_preview()

    def _update_color_btn(self):
        hex_c = self.watermark_color.name()
        self.btn_color.setStyleSheet(f"background-color: {hex_c}; border: 1px solid rgba(255, 255, 255, 0.25); border-radius: 4px; min-width: 36px;")

    def _pick_color(self):
        color = QColorDialog.getColor(self.watermark_color, self, "Filigran Rengi Seç")
        if color.isValid():
            self.watermark_color = color
            self._update_color_btn()
            self.update_preview()

    def _preview_prev_page(self):
        if self.preview_current_page > 0:
            self.preview_current_page -= 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_next_page(self):
        if self.preview_current_page < self.preview_total_pages - 1:
            self.preview_current_page += 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_page_spin_changed(self, val):
        target = val - 1
        if 0 <= target < self.preview_total_pages and target != self.preview_current_page:
            self.preview_current_page = target
            self.update_preview()

    def _render_base_page(self, src_path, target_w=260, target_h=360):
        if self._is_active_tab_doc(src_path):
            try:
                tab_doc = self.main_window.get_current_tab_doc()
                doc = tab_doc.doc
                total = len(doc)
                self.preview_total_pages = max(1, total)
                if self.preview_current_page >= self.preview_total_pages:
                    self.preview_current_page = self.preview_total_pages - 1
                self.spn_preview_page.blockSignals(True)
                self.spn_preview_page.setRange(1, self.preview_total_pages)
                self.spn_preview_page.setValue(self.preview_current_page + 1)
                self.spn_preview_page.blockSignals(False)
                self.lbl_total_preview_pages.setText(f"/ {self.preview_total_pages}")
                self.btn_prev_page.setEnabled(self.preview_current_page > 0)
                self.btn_next_page.setEnabled(self.preview_current_page < self.preview_total_pages - 1)
                self.spn_preview_page.setEnabled(self.preview_total_pages > 1)

                page = doc[self.preview_current_page]
                rect = page.rect
                self.preview_page_w = rect.width
                self.preview_page_h = rect.height
                scale = min(target_w / max(1.0, rect.width), target_h / max(1.0, rect.height))
                mat = fitz.Matrix(scale, scale)
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
                qpix = QPixmap.fromImage(img)
                if hasattr(self, 'lbl_preview_info'):
                    self.lbl_preview_info.setText(LanguageManager.tr("preview_page_counter", cur=self.preview_current_page + 1, total=self.preview_total_pages))
                return qpix
            except Exception:
                pass
        elif src_path and os.path.isfile(src_path):
            try:
                doc = fitz.open(src_path)
                if not doc.is_encrypted or doc.authenticate("") > 0:
                    total = len(doc)
                    self.preview_total_pages = max(1, total)
                    if self.preview_current_page >= self.preview_total_pages:
                        self.preview_current_page = self.preview_total_pages - 1
                    self.spn_preview_page.blockSignals(True)
                    self.spn_preview_page.setRange(1, self.preview_total_pages)
                    self.spn_preview_page.setValue(self.preview_current_page + 1)
                    self.spn_preview_page.blockSignals(False)
                    self.lbl_total_preview_pages.setText(f"/ {self.preview_total_pages}")
                    self.btn_prev_page.setEnabled(self.preview_current_page > 0)
                    self.btn_next_page.setEnabled(self.preview_current_page < self.preview_total_pages - 1)
                    self.spn_preview_page.setEnabled(self.preview_total_pages > 1)

                    page = doc[self.preview_current_page]
                    rect = page.rect
                    self.preview_page_w = rect.width
                    self.preview_page_h = rect.height
                    scale = min(target_w / max(1.0, rect.width), target_h / max(1.0, rect.height))
                    mat = fitz.Matrix(scale, scale)
                    pix = page.get_pixmap(matrix=mat, alpha=False)
                    img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888)
                    qpix = QPixmap.fromImage(img)
                    doc.close()
                    if hasattr(self, 'lbl_preview_info'):
                        self.lbl_preview_info.setText(LanguageManager.tr("preview_page_counter", cur=self.preview_current_page + 1, total=self.preview_total_pages))
                    return qpix
                doc.close()
            except Exception:
                pass

        if hasattr(self, 'lbl_preview_info'):
            self.lbl_preview_info.setText(LanguageManager.tr("preview_sample"))
        if hasattr(self, 'btn_prev_page'):
            self.btn_prev_page.setEnabled(False)
            self.btn_next_page.setEnabled(False)
            self.spn_preview_page.setEnabled(False)
            self.lbl_total_preview_pages.setText("/ 1")

        pix = QPixmap(target_w, target_h)
        pix.fill(QColor("#ffffff"))
        p = QPainter(pix)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        
        # Sayfa kenarlığı
        p.setPen(QPen(QColor("#d1d5db"), 1))
        p.drawRect(0, 0, target_w - 1, target_h - 1)
        
        # Başlık ve iskelet metin çizgileri
        p.setPen(Qt.PenStyle.NoPen)
        p.setBrush(QBrush(QColor("#374151")))
        p.drawRoundedRect(22, 28, target_w - 90, 11, 2, 2)
        
        p.setBrush(QBrush(QColor("#9ca3af")))
        p.drawRoundedRect(22, 48, target_w - 140, 6, 2, 2)
        
        p.setBrush(QBrush(QColor("#e5e7eb")))
        y = 75
        while y < target_h - 35:
            w_line = target_w - 44 if (y // 14) % 4 != 0 else target_w - 100
            p.drawRoundedRect(22, y, w_line, 5, 2, 2)
            y += 13
            
        p.end()
        return pix

    def update_preview(self):
        src_path = self.txt_src.text().strip()
        base_pix = self._render_base_page(src_path, target_w=260, target_h=360)
        
        self.lbl_preview.lock_aspect = self.chk_lock_aspect.isChecked()
        self.lbl_preview.preview_page_w = getattr(self, 'preview_page_w', 595.0)
        self.lbl_preview.preview_page_h = getattr(self, 'preview_page_h', 842.0)

        preview_pix = base_pix.copy()
        pw = preview_pix.width()
        ph = preview_pix.height()

        norm_pos = getattr(self.lbl_preview, 'norm_pos', QPointF(0.5, 0.5))
        cx = pw * norm_pos.x()
        cy = ph * norm_pos.y()
        
        p = QPainter(preview_pix)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setRenderHint(QPainter.RenderHint.TextAntialiasing)
        p.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
        
        is_text = self.tab_widget.currentIndex() == 0
        underlay = self.rad_layer_underlay.isChecked()
        if underlay:
            p.setCompositionMode(QPainter.CompositionMode.CompositionMode_Multiply)
        else:
            p.setCompositionMode(QPainter.CompositionMode.CompositionMode_SourceOver)
            
        if is_text:
            text = self.txt_watermark.text().strip() or "AURAPDF"
            raw_fs = self.spn_fontsize.value()
            angle = self._get_current_angle()
            opacity = float(self.sld_opacity.value()) / 100.0

            font_key = self.cmb_font_family.currentData() if hasattr(self, 'cmb_font_family') else "helv"
            if font_key == "tiro":
                qfont_family = "Times New Roman"
            elif font_key == "couri":
                qfont_family = "Courier New"
            else:
                qfont_family = "Segoe UI"

            # Safe font size math
            real_w = getattr(self, 'preview_page_w', 595.0)
            scaled_fs = max(8, int(raw_fs * (pw / max(1.0, real_w))))
            if self.chk_autofit.isChecked():
                preview_safe_fs = self._calculate_safe_fontsize(text, angle, pw, ph, margin=15.0, font_key=font_key)
                scaled_fs = min(scaled_fs, int(preview_safe_fs))
            
            col = QColor(self.watermark_color)
            col.setAlphaF(opacity)
            p.setPen(QPen(col))
            
            font = QFont(qfont_family, scaled_fs, QFont.Weight.Bold)
            p.setFont(font)
            self.lbl_preview.current_angle = angle
            self.lbl_preview.current_scale_value = raw_fs
            
            p.save()
            p.translate(cx, cy)
            p.rotate(-angle)
            rect = QRectF(-pw, -ph, pw * 2, ph * 2)
            p.drawText(rect, Qt.AlignmentFlag.AlignCenter, text)

            fm = QFontMetrics(font)
            tw = fm.horizontalAdvance(text)
            th = fm.height()
            bw = tw + 16
            bh = th + 8

            # Bounding box
            dash_pen = QPen(QColor("#388bfd"), 1.2, Qt.PenStyle.DashLine)
            p.setPen(dash_pen)
            p.setBrush(Qt.BrushStyle.NoBrush)
            p.drawRoundedRect(QRectF(-bw / 2.0, -bh / 2.0, bw, bh), 3, 3)

            # 4 corner handles
            handle_pen = QPen(QColor("#1a73e8"), 1.5)
            handle_brush = QBrush(QColor("#ffffff"))
            p.setPen(handle_pen)
            p.setBrush(handle_brush)
            corners = [
                QPointF(-bw / 2.0, -bh / 2.0),
                QPointF(bw / 2.0, -bh / 2.0),
                QPointF(-bw / 2.0, bh / 2.0),
                QPointF(bw / 2.0, bh / 2.0),
            ]
            for cpt in corners:
                p.drawEllipse(cpt, 3.0, 3.0)

            # Rotation stem & circular handle (22px above top center)
            stem_pen = QPen(QColor("#388bfd"), 1.5)
            p.setPen(stem_pen)
            p.drawLine(QPointF(0.0, -bh / 2.0), QPointF(0.0, -bh / 2.0 - 22.0))

            rot_handle_pen = QPen(QColor("#1a73e8"), 1.5)
            rot_handle_brush = QBrush(QColor("#ffffff"))
            p.setPen(rot_handle_pen)
            p.setBrush(rot_handle_brush)
            p.drawEllipse(QPointF(0.0, -bh / 2.0 - 22.0), 5.0, 5.0)

            p.setPen(Qt.PenStyle.NoPen)
            p.setBrush(QBrush(QColor("#1a73e8")))
            p.drawEllipse(QPointF(0.0, -bh / 2.0 - 22.0), 2.0, 2.0)

            p.restore()

            ox = (self.lbl_preview.width() - pw) / 2.0
            oy = (self.lbl_preview.height() - ph) / 2.0
            self.lbl_preview.watermark_center = QPointF(ox + cx, oy + cy)
            self.lbl_preview.watermark_rect = QRectF(ox + cx - bw / 2.0, oy + cy - bh / 2.0, bw, bh)
        else:
            img_path = self.txt_img_src.text().strip()
            if img_path and os.path.isfile(img_path):
                wm_img = QImage(img_path)
                if not wm_img.isNull():
                    lock_aspect = self.chk_lock_aspect.isChecked()
                    real_w = getattr(self, 'preview_page_w', 595.0)
                    real_h = getattr(self, 'preview_page_h', 842.0)
                    scale_pct = float(self.sld_img_scale.value()) / 100.0
                    opacity = float(self.sld_img_op.value()) / 100.0
                    img_angle = self._get_current_img_angle()
                    self.lbl_preview.current_angle = img_angle
                    self.lbl_preview.current_scale_value = self.sld_img_scale.value()

                    if lock_aspect:
                        target_img_w = max(15, int(pw * scale_pct))
                        aspect = float(wm_img.height()) / max(1.0, float(wm_img.width()))
                        target_img_h = max(15, int(round(target_img_w * aspect)))
                        scaled_img = wm_img.scaled(target_img_w, target_img_h, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
                    else:
                        target_img_w = max(15, int(pw * (self.spn_img_w.value() / max(1.0, real_w))))
                        target_img_h = max(15, int(ph * (self.spn_img_h.value() / max(1.0, real_h))))
                        scaled_img = wm_img.scaled(target_img_w, target_img_h, Qt.AspectRatioMode.IgnoreAspectRatio, Qt.TransformationMode.SmoothTransformation)

                    bw = scaled_img.width()
                    bh = scaled_img.height()

                    p.save()
                    p.translate(cx, cy)
                    p.rotate(-img_angle)
                    p.setOpacity(opacity)
                    p.drawImage(int(-bw / 2.0), int(-bh / 2.0), scaled_img)

                    # Bounding box & grip handles in rotated space
                    p.setOpacity(1.0)
                    dash_pen = QPen(QColor("#388bfd"), 1.2, Qt.PenStyle.DashLine)
                    p.setPen(dash_pen)
                    p.setBrush(Qt.BrushStyle.NoBrush)
                    p.drawRoundedRect(QRectF(-bw / 2.0 - 2, -bh / 2.0 - 2, bw + 4, bh + 4), 3, 3)

                    # Corner handles (4 köşe)
                    handle_pen = QPen(QColor("#1a73e8"), 1.5)
                    handle_brush = QBrush(QColor("#ffffff"))
                    p.setPen(handle_pen)
                    p.setBrush(handle_brush)
                    corners = [
                        QPointF(-bw / 2.0 - 2, -bh / 2.0 - 2),
                        QPointF(bw / 2.0 + 2, -bh / 2.0 - 2),
                        QPointF(-bw / 2.0 - 2, bh / 2.0 + 2),
                        QPointF(bw / 2.0 + 2, bh / 2.0 + 2),
                    ]
                    for cpt in corners:
                        p.drawEllipse(cpt, 3.5, 3.5)

                    # En-boy oranı korunduğunda sağ-sol çizgileri/kolları kalkar (kullanıcı talebi)
                    # En-boy oranı kilitli OLMADIĞINDA ise serbest esnetme için orta tutamaçlar gösterilir
                    if not lock_aspect:
                        side_pen = QPen(QColor("#60a5fa"), 1.2)
                        p.setPen(side_pen)
                        p.setBrush(QBrush(QColor("#ffffff")))
                        side_handles = [
                            QPointF(-bw / 2.0 - 2, 0.0), # Sol orta
                            QPointF(bw / 2.0 + 2, 0.0),  # Sağ orta
                            QPointF(0.0, -bh / 2.0 - 2), # Üst orta
                            QPointF(0.0, bh / 2.0 + 2),  # Alt orta
                        ]
                        for spt in side_handles:
                            p.drawRoundedRect(QRectF(spt.x() - 3.0, spt.y() - 3.0, 6.0, 6.0), 1, 1)

                    # Rotation stem & handle (22px above top center)
                    stem_pen = QPen(QColor("#388bfd"), 1.5)
                    p.setPen(stem_pen)
                    p.drawLine(QPointF(0.0, -bh / 2.0 - 2), QPointF(0.0, -bh / 2.0 - 22.0))

                    rot_handle_pen = QPen(QColor("#1a73e8"), 1.5)
                    rot_handle_brush = QBrush(QColor("#ffffff"))
                    p.setPen(rot_handle_pen)
                    p.setBrush(rot_handle_brush)
                    p.drawEllipse(QPointF(0.0, -bh / 2.0 - 22.0), 5.0, 5.0)

                    p.setPen(Qt.PenStyle.NoPen)
                    p.setBrush(QBrush(QColor("#1a73e8")))
                    p.drawEllipse(QPointF(0.0, -bh / 2.0 - 22.0), 2.0, 2.0)

                    p.restore()

                    # Update InteractiveWatermarkPreviewLabel state for unrotated hit-testing & resizing
                    ox = (self.lbl_preview.width() - pw) / 2.0
                    oy = (self.lbl_preview.height() - ph) / 2.0
                    self.lbl_preview.watermark_center = QPointF(ox + cx, oy + cy)
                    self.lbl_preview.watermark_rect = QRectF(ox + cx - bw / 2.0, oy + cy - bh / 2.0, bw, bh)
        
        p.end()
        self.lbl_preview.setPixmap(preview_pix)

    def execute_apply_watermark(self, mode="auto"):
        src_path = self.txt_src.text().strip()
        is_active_tab = self._is_active_tab_doc(src_path)

        if not is_active_tab and (not src_path or not os.path.isfile(src_path)):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return

        is_text_tab = (self.tab_widget.currentIndex() == 0)
        if is_text_tab:
            if not self.txt_watermark.text().strip():
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_text"))
                return
        else:
            img_path = self.txt_img_src.text().strip()
            if not img_path or not os.path.isfile(img_path):
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_image"))
                return

        settings = self._get_current_watermark_settings()
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()

        try:
            if is_active_tab and tab_doc and tab_doc.doc:
                # Direct apply to active document
                if mode == "auto" and getattr(tab_doc, 'has_watermark', False):
                    if getattr(tab_doc, 'pre_watermark_bytes', None):
                        try:
                            t_chk = fitz.open("pdf", tab_doc.pre_watermark_bytes)
                            if len(t_chk) == len(tab_doc.doc):
                                base_bytes = tab_doc.pre_watermark_bytes
                            else:
                                temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                                remove_watermark_from_doc(temp_doc)
                                base_bytes = temp_doc.tobytes()
                                tab_doc.pre_watermark_bytes = base_bytes
                                temp_doc.close()
                            t_chk.close()
                        except Exception:
                            base_bytes = tab_doc.doc.tobytes()
                    else:
                        temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                        remove_watermark_from_doc(temp_doc)
                        base_bytes = temp_doc.tobytes()
                        tab_doc.pre_watermark_bytes = base_bytes
                        temp_doc.close()
                else:
                    if not getattr(tab_doc, 'pre_watermark_bytes', None) or mode == "add_new":
                        tab_doc.pre_watermark_bytes = tab_doc.doc.tobytes()
                    base_bytes = tab_doc.pre_watermark_bytes if mode != "add_new" else tab_doc.doc.tobytes()

                doc = fitz.open("pdf", base_bytes)
                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations(doc)

                apply_watermark_to_doc(doc, settings)
                new_bytes = doc.tobytes()

                old_bytes = tab_doc.doc.tobytes()
                old_drawings = [dict(d) for d in getattr(tab_doc, 'drawings', [])]
                old_page = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))
                old_has_watermark = getattr(tab_doc, 'has_watermark', False)
                old_settings = copy.deepcopy(getattr(tab_doc, 'watermark_settings', None))
                old_pre_bytes = getattr(tab_doc, 'pre_watermark_bytes', None)

                tab_doc.doc.close()
                tab_doc.doc = doc
                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations()
                tab_doc.has_watermark = True
                tab_doc.watermark_settings = settings
                tab_doc.set_modified(True)
                tab_doc.render_pages_init()
                tab_doc.rerender_visible_pages()

                curr_p = max(0, min(old_page, len(doc) - 1))
                tab_doc.go_to_page(curr_p)
                if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                    if self.main_window.thumbnails_dock.isVisible():
                        self.main_window.thumbnails_dock.load_thumbnails_lazy()
                        self.main_window.thumbnails_dock.select_page(curr_p)

                tab_doc.push_undo_action({
                    "type": "watermark",
                    "old_bytes": old_bytes,
                    "old_drawings": old_drawings,
                    "old_page": curr_p,
                    "old_has_watermark": old_has_watermark,
                    "old_watermark_settings": old_settings,
                    "old_pre_watermark_bytes": old_pre_bytes,
                    "new_bytes": new_bytes,
                    "new_drawings": [dict(d) for d in getattr(tab_doc, 'drawings', [])],
                    "new_page": curr_p,
                    "new_has_watermark": True,
                    "new_watermark_settings": settings,
                    "new_pre_watermark_bytes": tab_doc.pre_watermark_bytes,
                })

                msg = LanguageManager.tr("watermark_success_update") if (self.has_existing_watermark and mode == "auto") else LanguageManager.tr("watermark_success_apply")
                QMessageBox.information(self, LanguageManager.tr("success"), msg)
                self.accept()
            else:
                # Standalone file apply -> save to target
                dst_path = self.txt_dst.text().strip()
                if not dst_path:
                    dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), src_path.replace(".pdf", "_filigranli.pdf"), "PDF (*.pdf)")
                    if not dst_path:
                        return
                    self.txt_dst.setText(dst_path)

                doc = fitz.open(src_path)
                if doc.is_encrypted:
                    from PyQt6.QtWidgets import QInputDialog
                    pw, ok = QInputDialog.getText(self, LanguageManager.tr("password_required_title"), LanguageManager.tr("encrypt_current_pw"), QLineEdit.EchoMode.Password)
                    if not ok or doc.authenticate(pw) <= 0:
                        doc.close()
                        QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_wrong_pw"))
                        return

                if mode != "add_new":
                    remove_watermark_from_doc(doc)
                apply_watermark_to_doc(doc, settings)
                doc.save(dst_path, garbage=3, deflate=True)
                doc.close()

                QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("watermark_success", path=dst_path))
                if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                    self.main_window.add_pdf_tab(dst_path)
                self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def execute_remove_watermark(self):
        reply = QMessageBox.question(
            self,
            LanguageManager.tr("confirm_title"),
            LanguageManager.tr("watermark_confirm_remove"),
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        src_path = self.txt_src.text().strip()
        is_active_tab = self._is_active_tab_doc(src_path)
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()

        try:
            if is_active_tab and tab_doc and tab_doc.doc:
                old_bytes = tab_doc.doc.tobytes()
                old_drawings = [dict(d) for d in getattr(tab_doc, 'drawings', [])]
                old_page = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))
                old_has_watermark = getattr(tab_doc, 'has_watermark', False)
                old_settings = copy.deepcopy(getattr(tab_doc, 'watermark_settings', None))
                old_pre_bytes = getattr(tab_doc, 'pre_watermark_bytes', None)

                if old_pre_bytes:
                    clean_doc = fitz.open("pdf", old_pre_bytes)
                    remove_watermark_from_doc(clean_doc)
                else:
                    clean_doc = fitz.open("pdf", old_bytes)
                    remove_watermark_from_doc(clean_doc)

                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations(clean_doc)

                new_bytes = clean_doc.tobytes()
                tab_doc.doc.close()
                tab_doc.doc = clean_doc
                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations()
                tab_doc.has_watermark = False
                tab_doc.watermark_settings = None
                tab_doc.pre_watermark_bytes = None
                tab_doc.set_modified(True)
                tab_doc.render_pages_init()
                tab_doc.rerender_visible_pages()

                curr_p = max(0, min(old_page, len(clean_doc) - 1))
                tab_doc.go_to_page(curr_p)
                if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                    if self.main_window.thumbnails_dock.isVisible():
                        self.main_window.thumbnails_dock.load_thumbnails_lazy()
                        self.main_window.thumbnails_dock.select_page(curr_p)

                tab_doc.push_undo_action({
                    "type": "watermark",
                    "old_bytes": old_bytes,
                    "old_drawings": old_drawings,
                    "old_page": curr_p,
                    "old_has_watermark": old_has_watermark,
                    "old_watermark_settings": old_settings,
                    "old_pre_watermark_bytes": old_pre_bytes,
                    "new_bytes": new_bytes,
                    "new_drawings": [dict(d) for d in getattr(tab_doc, 'drawings', [])],
                    "new_page": curr_p,
                    "new_has_watermark": False,
                    "new_watermark_settings": None,
                    "new_pre_watermark_bytes": None,
                })

                QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("watermark_success_remove"))
                self.accept()
            else:
                if not src_path or not os.path.isfile(src_path):
                    QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
                    return
                doc = fitz.open(src_path)
                remove_watermark_from_doc(doc)
                dst_path = self.txt_dst.text().strip() or src_path
                if os.path.abspath(dst_path) == os.path.abspath(src_path):
                    clean_bytes = doc.tobytes(garbage=3, deflate=True)
                    doc.close()
                    with open(dst_path, "wb") as f:
                        f.write(clean_bytes)
                else:
                    doc.save(dst_path, garbage=3, deflate=True)
                    doc.close()
                QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("watermark_success_remove"))
                self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def execute_save_as_watermark(self):
        src_path = self.txt_src.text().strip()
        is_active_tab = self._is_active_tab_doc(src_path)
        if not is_active_tab and (not src_path or not os.path.isfile(src_path)):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return

        cur = self.txt_dst.text().strip()
        default_name = cur if cur else (src_path.replace(".pdf", "_filigranli.pdf") if os.path.isfile(src_path) else "belge_filigranli.pdf")
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), default_name, "PDF (*.pdf)")
        if not path:
            return
        self.txt_dst.setText(path)

        is_text_tab = (self.tab_widget.currentIndex() == 0)
        if is_text_tab:
            if not self.txt_watermark.text().strip():
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_text"))
                return
        else:
            img_path = self.txt_img_src.text().strip()
            if not img_path or not os.path.isfile(img_path):
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_image"))
                return

        settings = self._get_current_watermark_settings()
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()

        try:
            if is_active_tab and tab_doc and tab_doc.doc:
                if getattr(tab_doc, 'has_watermark', False):
                    if getattr(tab_doc, 'pre_watermark_bytes', None):
                        base_bytes = tab_doc.pre_watermark_bytes
                    else:
                        temp_doc = fitz.open("pdf", tab_doc.doc.tobytes())
                        remove_watermark_from_doc(temp_doc)
                        base_bytes = temp_doc.tobytes()
                        temp_doc.close()
                else:
                    base_bytes = tab_doc.doc.tobytes()
                doc = fitz.open("pdf", base_bytes)
                if hasattr(tab_doc, 'export_drawings_to_pdf'):
                    tab_doc.export_drawings_to_pdf(doc)
            else:
                doc = fitz.open(src_path)
                if doc.is_encrypted:
                    from PyQt6.QtWidgets import QInputDialog
                    pw, ok = QInputDialog.getText(self, LanguageManager.tr("password_required_title"), LanguageManager.tr("encrypt_current_pw"), QLineEdit.EchoMode.Password)
                    if not ok or doc.authenticate(pw) <= 0:
                        doc.close()
                        QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_wrong_pw"))
                        return
                remove_watermark_from_doc(doc)

            apply_watermark_to_doc(doc, settings)
            doc.save(path, garbage=3, deflate=True)
            doc.close()

            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("watermark_success", path=path))
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                self.main_window.add_pdf_tab(path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def execute_watermark(self):
        self.execute_apply_watermark(mode="auto")


class PDFBackgroundDialog(QDialog):
    @staticmethod
    def _is_background_stream(stream_bytes: bytes, page_rect) -> bool:
        if not stream_bytes:
            return False
        if b'AuraPDF-Background' in stream_bytes:
            return True
        s = stream_bytes.strip()
        # Eski sürümde eklenen tam sayfa düz renk veya görsel arka planı tespit et
        if len(s) < 220:
            w_int = int(page_rect.width)
            h_int = int(page_rect.height)
            if b're' in s and (b'f' in s or b'F' in s) and (str(w_int).encode() in s or str(h_int).encode() in s):
                return True
            if b'Do' in s and b'cm' in s and len(s) < 140:
                return True
        return False

    @classmethod
    def _remove_page_background(cls, doc, page) -> bool:
        """Sayfadaki mevcut arka plan akışlarını tespit edip temizler."""
        try:
            contents = page.get_contents()
            if not contents:
                return False
            removed = False
            remaining = []
            for xref in contents:
                try:
                    stream = doc.xref_stream(xref)
                    if cls._is_background_stream(stream, page.rect):
                        removed = True
                        continue
                except Exception:
                    pass
                remaining.append(xref)

            if removed:
                if len(remaining) == 0:
                    doc.xref_set_key(page.xref, 'Contents', 'null')
                elif len(remaining) == 1:
                    doc.xref_set_key(page.xref, 'Contents', f'{remaining[0]} 0 R')
                else:
                    refs = ' '.join(f'{x} 0 R' for x in remaining)
                    doc.xref_set_key(page.xref, 'Contents', f'[ {refs} ]')
            return removed
        except Exception:
            return False

    def __init__(self, main_window):
        super().__init__(main_window)
        self.main_window = main_window
        self.current_bg_color = QColor("#fbf0d9")  # Varsayılan Sepya

        self.setWindowTitle(LanguageManager.tr("background_dialog_title"))
        self.resize(1040, 680)
        self.setMinimumSize(850, 560)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        main_layout = QHBoxLayout(self)
        main_layout.setContentsMargins(16, 16, 16, 16)
        main_layout.setSpacing(16)

        left_widget = QWidget()
        layout = QVBoxLayout(left_widget)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(10)

        # Kaynak PDF
        layout.addWidget(QLabel(LanguageManager.tr("encrypt_source_file")))
        h_src = QHBoxLayout()
        self.txt_src = QLineEdit()
        self.txt_src.setPlaceholderText("PDF Dosyası Seçin...")
        btn_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_src.clicked.connect(self._browse_source)
        h_src.addWidget(self.txt_src)
        h_src.addWidget(btn_src)
        layout.addLayout(h_src)

        self.tab_widget = QTabWidget(self)

        # --- TAB 1: Düz Renk Arka Plan ---
        tab_col = QWidget()
        v_col = QVBoxLayout(tab_col)
        v_col.setContentsMargins(14, 14, 14, 14)
        v_col.setSpacing(10)

        v_col.addWidget(QLabel(LanguageManager.tr("background_color_presets")))

        h_c_btns = QHBoxLayout()
        presets = [
            (LanguageManager.tr("background_color_sepia"), "#fbf0d9", "#222222"),
            (LanguageManager.tr("background_color_dark"), "#1e1e1e", "#ffffff"),
            (LanguageManager.tr("background_color_light_gray"), "#f0f2f5", "#222222"),
            (LanguageManager.tr("background_color_ivory"), "#fdfbf7", "#222222"),
        ]
        for name, hex_code, text_col in presets:
            btn = QPushButton(name)
            btn.setProperty("class", "ColorBtn")
            btn.setStyleSheet(f"background-color: {hex_code}; color: {text_col}; font-size: 11px; padding: 6px 10px;")
            btn.clicked.connect(lambda checked, col=hex_code: self._select_preset_color(col))
            h_c_btns.addWidget(btn)
        v_col.addLayout(h_c_btns)

        h_custom_c = QHBoxLayout()
        self.btn_custom_color = QPushButton(LanguageManager.tr("background_color_custom"))
        self.btn_custom_color.clicked.connect(self._pick_custom_color)
        self.lbl_color_preview = QLabel("         ")
        self._update_color_preview()
        h_custom_c.addWidget(self.btn_custom_color)
        h_custom_c.addWidget(self.lbl_color_preview)
        h_custom_c.addStretch()
        v_col.addLayout(h_custom_c)

        h_op = QHBoxLayout()
        self.lbl_opacity = QLabel(LanguageManager.tr("background_opacity", val=100))
        self.sld_opacity = QSlider(Qt.Orientation.Horizontal)
        self.sld_opacity.setRange(10, 100)
        self.sld_opacity.setValue(100)
        self.sld_opacity.valueChanged.connect(lambda val: self.lbl_opacity.setText(LanguageManager.tr("background_opacity", val=val)))
        h_op.addWidget(self.lbl_opacity)
        h_op.addWidget(self.sld_opacity)
        v_col.addLayout(h_op)

        v_col.addStretch()
        self.tab_widget.addTab(tab_col, LanguageManager.tr("background_tab_color"))

        # --- TAB 2: Görsel / Antet Şablonu ---
        tab_img = QWidget()
        v_img = QVBoxLayout(tab_img)
        v_img.setContentsMargins(14, 14, 14, 14)
        v_img.setSpacing(10)

        v_img.addWidget(QLabel(LanguageManager.tr("watermark_image_file")))
        h_img_src = QHBoxLayout()
        self.txt_img_src = QLineEdit()
        self.txt_img_src.setPlaceholderText("Görsel Dosyası Seçin (PNG, JPG, BMP)...")
        btn_img_src = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_img_src.clicked.connect(self._browse_image)
        h_img_src.addWidget(self.txt_img_src)
        h_img_src.addWidget(btn_img_src)
        v_img.addLayout(h_img_src)

        h_fit = QHBoxLayout()
        h_fit.addWidget(QLabel(LanguageManager.tr("background_image_fit")))
        self.cmb_fit = QComboBox()
        self.cmb_fit.addItem(LanguageManager.tr("background_fit_fill"), "fill")
        self.cmb_fit.addItem(LanguageManager.tr("background_fit_proportional"), "proportional")
        h_fit.addWidget(self.cmb_fit)
        h_fit.addStretch()
        v_img.addLayout(h_fit)

        h_img_op = QHBoxLayout()
        self.lbl_img_op = QLabel(LanguageManager.tr("background_opacity", val=100))
        self.sld_img_op = QSlider(Qt.Orientation.Horizontal)
        self.sld_img_op.setRange(10, 100)
        self.sld_img_op.setValue(100)
        self.sld_img_op.valueChanged.connect(lambda val: self.lbl_img_op.setText(LanguageManager.tr("background_opacity", val=val)))
        h_img_op.addWidget(self.lbl_img_op)
        h_img_op.addWidget(self.sld_img_op)
        v_img.addLayout(h_img_op)

        v_img.addStretch()
        self.tab_widget.addTab(tab_img, LanguageManager.tr("background_tab_image"))

        layout.addWidget(self.tab_widget)

        # Sayfa Aralığı
        grp_pages = QGroupBox(LanguageManager.tr("watermark_page_range_group"))
        v_pages = QVBoxLayout(grp_pages)
        self.rad_all_pages = QRadioButton(LanguageManager.tr("background_pages_all"))
        self.rad_all_pages.setChecked(True)

        tab_doc_bg = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        cur_p_idx = getattr(tab_doc_bg, 'current_page_idx', getattr(tab_doc_bg, 'current_page', 0)) if tab_doc_bg else 0
        cur_p_num = (cur_p_idx + 1) if (tab_doc_bg and tab_doc_bg.doc and len(tab_doc_bg.doc) > 0) else 1
        self.rad_current_page = QRadioButton(f"{LanguageManager.tr('background_pages_current', default='Yalnızca Mevcut Sayfa')} (Sayfa {cur_p_num})")

        self.rad_skip_first = QRadioButton(LanguageManager.tr("background_pages_skip_first"))
        self.rad_first_page = QRadioButton(LanguageManager.tr("background_pages_first"))
        self.rad_custom_pages = QRadioButton(LanguageManager.tr("background_pages_custom"))
        self.txt_custom_pages = QLineEdit()
        self.txt_custom_pages.setPlaceholderText("1-5, 8, 12")

        v_pages.addWidget(self.rad_all_pages)
        v_pages.addWidget(self.rad_current_page)
        v_pages.addWidget(self.rad_skip_first)
        v_pages.addWidget(self.rad_first_page)
        h_cp = QHBoxLayout()
        h_cp.addWidget(self.rad_custom_pages)
        h_cp.addWidget(self.txt_custom_pages)
        v_pages.addLayout(h_cp)
        layout.addWidget(grp_pages)

        # Hedef PDF
        layout.addWidget(QLabel(LanguageManager.tr("encrypt_target_file")))
        h_dst = QHBoxLayout()
        self.txt_dst = QLineEdit()
        btn_dst = QPushButton(LanguageManager.tr("split_btn_browse"))
        btn_dst.clicked.connect(self._browse_target)
        h_dst.addWidget(self.txt_dst)
        h_dst.addWidget(btn_dst)
        layout.addLayout(h_dst)

        # Alt Butonlar
        h_bot = QHBoxLayout()
        self.chk_open_tab = QCheckBox(LanguageManager.tr("encrypt_chk_open"))
        self.chk_open_tab.setChecked(True)
        h_bot.addWidget(self.chk_open_tab)
        h_bot.addStretch()

        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.setMinimumHeight(30)
        btn_cancel.clicked.connect(self.reject)
        h_bot.addWidget(btn_cancel)

        self.btn_remove = QPushButton(LanguageManager.tr("background_btn_remove"))
        self.btn_remove.setMinimumHeight(30)
        self.btn_remove.setStyleSheet("""
            QPushButton {
                background-color: rgba(239, 68, 68, 0.12);
                color: #ef4444;
                border: 1px solid rgba(239, 68, 68, 0.3);
                border-radius: 6px;
                padding: 4px 10px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: rgba(239, 68, 68, 0.22);
                border-color: #ef4444;
            }
        """)
        self.btn_remove.clicked.connect(self.remove_background)
        h_bot.addWidget(self.btn_remove)

        self.btn_save_as = QPushButton(LanguageManager.tr("watermark_btn_save_as", default="Farklı Kaydet..."))
        self.btn_save_as.setMinimumHeight(30)
        self.btn_save_as.clicked.connect(lambda: self.execute_background(to_file=True))
        h_bot.addWidget(self.btn_save_as)

        self.btn_apply = QPushButton(LanguageManager.tr("watermark_btn_apply_doc", default="Açık Belgeye Uygula"))
        self.btn_apply.setObjectName("PrimaryBtn")
        self.btn_apply.setMinimumHeight(30)
        self.btn_apply.clicked.connect(lambda: self.execute_background(to_file=False))
        h_bot.addWidget(self.btn_apply)

        # Sol Kolon Konteyneri (Kaydırma Alanı + Sabit Alt Butonlar)
        left_container = QWidget()
        v_left = QVBoxLayout(left_container)
        v_left.setContentsMargins(0, 0, 0, 0)
        v_left.setSpacing(8)

        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        left_scroll.setStyleSheet("""
            QScrollArea { border: none; background: transparent; }
            QScrollBar:vertical {
                border: none;
                background: #1a1a1a;
                width: 8px;
                margin: 0px;
                border-radius: 4px;
            }
            QScrollBar::handle:vertical {
                background: #444444;
                min-height: 24px;
                border-radius: 4px;
            }
            QScrollBar::handle:vertical:hover {
                background: #555555;
            }
        """)
        left_scroll.setWidget(left_widget)
        v_left.addWidget(left_scroll, 1)
        v_left.addLayout(h_bot)

        # Canlı Önizleme Kartı
        grp_preview = QGroupBox(LanguageManager.tr('preview_title'))
        grp_preview.setFixedWidth(325)
        grp_preview.setStyleSheet("""
            QGroupBox {
                border: 1px solid rgba(255, 255, 255, 0.1);
                border-radius: 6px;
                margin-top: 10px;
                padding-top: 14px;
            }
        """)
        v_prev = QVBoxLayout(grp_preview)
        v_prev.setContentsMargins(10, 12, 10, 12)
        v_prev.setSpacing(8)

        self.lbl_preview_info = QLabel(LanguageManager.tr("preview_page_1"))
        self.lbl_preview_info.setStyleSheet(" font-size: 11px;")
        v_prev.addWidget(self.lbl_preview_info)

        frame_box = QFrame()
        frame_box.setObjectName("PreviewFrame")
        frame_box.setProperty("class", "preview-card")
        v_box = QVBoxLayout(frame_box)
        v_box.setContentsMargins(8, 8, 8, 8)
        v_box.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.lbl_preview = QLabel()
        self.lbl_preview.setFixedSize(260, 360)
        self.lbl_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_preview.setStyleSheet("border: 1px solid rgba(255, 255, 255, 0.15); border-radius: 2px;")
        v_box.addWidget(self.lbl_preview)
        v_prev.addWidget(frame_box)

        self.preview_current_page = getattr(tab_doc_bg, 'current_page_idx', getattr(tab_doc_bg, 'current_page', 0)) if tab_doc_bg else 0
        self.preview_total_pages = 1

        # Sayfa Gezinti Çubuğu (Tüm sayfalara bakabilme)
        h_nav = QHBoxLayout()
        h_nav.setSpacing(6)
        self.btn_prev_page = QPushButton(LanguageManager.tr("preview_prev"))
        self.btn_prev_page.setProperty("class", "NavBtn")
        self.btn_prev_page.setIcon(ModernIconFactory.icon_chevron_left())
        self.btn_prev_page.setMinimumWidth(84)
        self.btn_prev_page.clicked.connect(self._preview_prev_page)

        self.spn_preview_page = QSpinBox()
        self.spn_preview_page.setObjectName("PreviewPageSpin")
        self.spn_preview_page.setButtonSymbols(QAbstractSpinBox.ButtonSymbols.NoButtons)
        self.spn_preview_page.setRange(1, 1)
        self.spn_preview_page.setValue(1)
        self.spn_preview_page.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.spn_preview_page.setFixedWidth(56)
        self.spn_preview_page.setStyleSheet("""
            QSpinBox#PreviewPageSpin {
                background-color: #1a1a1e;
                color: #f1f5f9;
                border: 1px solid rgba(255, 255, 255, 0.12);
                border-radius: 4px;
                padding: 2px 2px;
                font-size: 12px;
                font-weight: bold;
                text-align: center;
            }
            QSpinBox#PreviewPageSpin:focus {
                border: 1px solid #38bdf8;
            }
        """)
        self.spn_preview_page.valueChanged.connect(self._preview_page_spin_changed)

        self.lbl_total_preview_pages = QLabel("/ 1")
        self.lbl_total_preview_pages.setStyleSheet(" font-size: 11px; font-weight: 500;")

        self.btn_next_page = QPushButton(LanguageManager.tr("preview_next"))
        self.btn_next_page.setProperty("class", "NavBtn")
        self.btn_next_page.setIcon(ModernIconFactory.icon_chevron_right())
        self.btn_next_page.setMinimumWidth(84)
        self.btn_next_page.clicked.connect(self._preview_next_page)

        h_nav.addWidget(self.btn_prev_page)
        h_nav.addWidget(self.spn_preview_page)
        h_nav.addWidget(self.lbl_total_preview_pages)
        h_nav.addWidget(self.btn_next_page)
        v_prev.addLayout(h_nav)

        btn_refresh = QPushButton(LanguageManager.tr('preview_refresh'))
        btn_refresh.clicked.connect(self.update_preview)
        v_prev.addWidget(btn_refresh)
        v_prev.addStretch()

        main_layout.addWidget(left_container, 1)
        main_layout.addWidget(grp_preview, 0)

        # Aktif sekme varsa hazırla (kayıtlı veya boş belge)
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        if tab_doc and tab_doc.doc and len(tab_doc.doc) > 0:
            if getattr(tab_doc, 'file_path', None):
                p = tab_doc.file_path
                self.txt_src.setText(p)
                base_dir = os.path.dirname(p)
                base_name, _ = os.path.splitext(os.path.basename(p))
                self.txt_dst.setText(os.path.join(base_dir, f"{base_name}_arkaplanli.pdf"))
            else:
                doc_title = getattr(tab_doc, 'custom_title', None) or LanguageManager.tr("new_blank_doc")
                self.txt_src.setText(f"[{doc_title}]")
                desktop = os.path.join(os.path.expanduser("~"), "Desktop")
                self.txt_dst.setText(os.path.join(desktop, f"{doc_title}_arkaplanli.pdf"))

        # Önizleme sinyalleri bağlantısı
        self.txt_src.textChanged.connect(self._on_source_changed)
        self.sld_opacity.valueChanged.connect(self.update_preview)
        self.txt_img_src.textChanged.connect(self.update_preview)
        self.cmb_fit.currentIndexChanged.connect(self.update_preview)
        self.sld_img_op.valueChanged.connect(self.update_preview)
        self.tab_widget.currentChanged.connect(self.update_preview)

        self.update_preview()

    def _is_active_tab_doc(self, src_path=None) -> bool:
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        if not tab_doc or tab_doc.doc is None or len(tab_doc.doc) == 0:
            return False
        if src_path is None:
            src_path = self.txt_src.text().strip()
        if getattr(tab_doc, 'file_path', None) and os.path.isfile(src_path):
            return os.path.abspath(tab_doc.file_path) == os.path.abspath(src_path)
        if not getattr(tab_doc, 'file_path', None) and (src_path.startswith("[") or not src_path):
            return True
        return False

    def _on_source_changed(self):
        self.preview_current_page = 0
        self.update_preview()

    def _select_preset_color(self, hex_code: str):
        self.current_bg_color = QColor(hex_code)
        self._update_color_preview()
        self.update_preview()

    def _update_color_preview(self):
        hex_c = self.current_bg_color.name()
        self.lbl_color_preview.setStyleSheet(f"background-color: {hex_c}; border: 2px solid rgba(255, 255, 255, 0.3); border-radius: 4px; min-width: 48px; min-height: 24px;")

    def _pick_custom_color(self):
        color = QColorDialog.getColor(self.current_bg_color, self, "Arka Plan Rengi Seç")
        if color.isValid():
            self.current_bg_color = color
            self._update_color_preview()
            self.update_preview()

    def _browse_source(self):
        path, _ = QFileDialog.getOpenFileName(self, LanguageManager.tr("open_dialog_title"), "", "PDF (*.pdf)")
        if path:
            self.txt_src.setText(path)
            base_dir = os.path.dirname(path)
            base_name, _ = os.path.splitext(os.path.basename(path))
            self.txt_dst.setText(os.path.join(base_dir, f"{base_name}_arkaplanli.pdf"))
            self.preview_current_page = 0
            self.update_preview()

    def _browse_target(self):
        cur = self.txt_dst.text().strip()
        path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), cur or "belge_arkaplanli.pdf", "PDF (*.pdf)")
        if path:
            self.txt_dst.setText(path)

    def _browse_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "Görsel Seç", "", "Görseller (*.png *.jpg *.jpeg *.bmp *.webp)")
        if path:
            self.txt_img_src.setText(path)
            self.update_preview()

    def _preview_prev_page(self):
        if self.preview_current_page > 0:
            self.preview_current_page -= 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_next_page(self):
        if self.preview_current_page < self.preview_total_pages - 1:
            self.preview_current_page += 1
            self.spn_preview_page.blockSignals(True)
            self.spn_preview_page.setValue(self.preview_current_page + 1)
            self.spn_preview_page.blockSignals(False)
            self.update_preview()

    def _preview_page_spin_changed(self, val):
        target = val - 1
        if 0 <= target < self.preview_total_pages and target != self.preview_current_page:
            self.preview_current_page = target
            self.update_preview()

    def update_preview(self):
        src_path = self.txt_src.text().strip()
        is_color = self.tab_widget.currentIndex() == 0
        target_w = 260
        target_h = 360
        
        canvas = QPixmap(target_w, target_h)
        canvas.fill(QColor("#ffffff"))
        
        p = QPainter(canvas)
        p.setRenderHint(QPainter.RenderHint.Antialiasing)
        p.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
        
        if is_color:
            bg_col = QColor(self.current_bg_color)
            opacity = float(self.sld_opacity.value()) / 100.0
            bg_col.setAlphaF(opacity)
            p.fillRect(0, 0, target_w, target_h, bg_col)
        else:
            img_path = self.txt_img_src.text().strip()
            if img_path and os.path.isfile(img_path):
                bg_img = QImage(img_path)
                if not bg_img.isNull():
                    opacity = float(self.sld_img_op.value()) / 100.0
                    fit_mode = self.cmb_fit.currentData()
                    p.setOpacity(opacity)
                    if fit_mode == "fill":
                        p.drawImage(QRectF(0, 0, target_w, target_h), bg_img)
                    else:
                        scaled = bg_img.scaled(target_w, target_h, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
                        x = (target_w - scaled.width()) / 2.0
                        y = (target_h - scaled.height()) / 2.0
                        p.drawImage(int(x), int(y), scaled)
                    p.setOpacity(1.0)
                    
        has_rendered_doc = False
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        is_active_tab = self._is_active_tab_doc(src_path)
        doc = None
        should_close_doc = False

        if is_active_tab and tab_doc and tab_doc.doc:
            try:
                doc = fitz.open("pdf", tab_doc.doc.tobytes())
                should_close_doc = True
            except Exception:
                doc = tab_doc.doc
                should_close_doc = False
        elif src_path and os.path.isfile(src_path):
            try:
                doc = fitz.open(src_path)
                should_close_doc = True
            except Exception:
                doc = None

        if doc:
            try:
                if not doc.is_encrypted or doc.authenticate("") > 0:
                    total = len(doc)
                    self.preview_total_pages = max(1, total)
                    if self.preview_current_page >= self.preview_total_pages:
                        self.preview_current_page = self.preview_total_pages - 1
                    self.spn_preview_page.blockSignals(True)
                    self.spn_preview_page.setRange(1, self.preview_total_pages)
                    self.spn_preview_page.setValue(self.preview_current_page + 1)
                    self.spn_preview_page.blockSignals(False)
                    self.lbl_total_preview_pages.setText(f"/ {self.preview_total_pages}")
                    self.btn_prev_page.setEnabled(self.preview_current_page > 0)
                    self.btn_next_page.setEnabled(self.preview_current_page < self.preview_total_pages - 1)
                    self.spn_preview_page.setEnabled(self.preview_total_pages > 1)

                    page = doc[self.preview_current_page]
                    if should_close_doc:
                        self._remove_page_background(doc, page)
                    rect = page.rect
                    scale = min(target_w / max(1.0, rect.width), target_h / max(1.0, rect.height))
                    mat = fitz.Matrix(scale, scale)
                    pix = page.get_pixmap(matrix=mat, alpha=True)
                    img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGBA8888)
                    p.drawImage(0, 0, img)
                    has_rendered_doc = True
                    if hasattr(self, 'lbl_preview_info'):
                        self.lbl_preview_info.setText(LanguageManager.tr("preview_page_counter", cur=self.preview_current_page + 1, total=self.preview_total_pages))
            except Exception:
                pass
            finally:
                if should_close_doc:
                    doc.close()

        if not has_rendered_doc:
            if hasattr(self, 'lbl_preview_info'):
                self.lbl_preview_info.setText(LanguageManager.tr("preview_sample"))
            if hasattr(self, 'btn_prev_page'):
                self.btn_prev_page.setEnabled(False)
                self.btn_next_page.setEnabled(False)
                self.spn_preview_page.setEnabled(False)
                self.lbl_total_preview_pages.setText("/ 1")
            p.setPen(Qt.PenStyle.NoPen)
            p.setBrush(QBrush(QColor("#2d3748")))
            p.drawRoundedRect(22, 28, target_w - 90, 11, 2, 2)
            
            p.setBrush(QBrush(QColor("#4a5568")))
            p.drawRoundedRect(22, 48, target_w - 140, 6, 2, 2)
            
            p.setBrush(QBrush(QColor("#718096")))
            y = 75
            while y < target_h - 35:
                w_line = target_w - 44 if (y // 14) % 4 != 0 else target_w - 100
                p.drawRoundedRect(22, y, w_line, 5, 2, 2)
                y += 13

        # Kenarlık
        p.setPen(QPen(QColor("#4b5563"), 1))
        p.setBrush(Qt.BrushStyle.NoBrush)
        p.drawRect(0, 0, target_w - 1, target_h - 1)

        p.end()
        self.lbl_preview.setPixmap(canvas)

    def execute_background(self, to_file=False):
        src_path = self.txt_src.text().strip()
        is_active_tab = self._is_active_tab_doc(src_path)

        if not is_active_tab and (not src_path or not os.path.isfile(src_path)):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return

        is_color_tab = self.tab_widget.currentIndex() == 0

        # Görsel arka plan kontrolü
        img_stream = None
        aspect = 1.0
        fit_mode = "fill"
        if not is_color_tab:
            img_path = self.txt_img_src.text().strip()
            if not img_path or not os.path.isfile(img_path):
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_image"))
                return

            base_img = QImage(img_path)
            if base_img.isNull():
                QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("watermark_err_no_image"))
                return

            opacity = float(self.sld_img_op.value()) / 100.0
            fit_mode = self.cmb_fit.currentData()

            res_img = QImage(base_img.size(), QImage.Format.Format_ARGB32)
            res_img.fill(Qt.GlobalColor.transparent)
            p = QPainter(res_img)
            p.setOpacity(opacity)
            p.drawImage(0, 0, base_img)
            p.end()

            ba = QByteArray()
            buf = QBuffer(ba)
            buf.open(QIODevice.OpenModeFlag.WriteOnly)
            res_img.save(buf, "PNG")
            img_stream = bytes(ba)
            aspect = base_img.height() / max(1, base_img.width())

        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()

        # 1. Açık Belgeye Doğrudan Uygula (Hafızadaki Boş veya Kayıtlı Belge)
        if not to_file and is_active_tab and tab_doc and tab_doc.doc:
            try:
                doc = fitz.open("pdf", tab_doc.doc.tobytes())
                total_pages = len(doc)

                if hasattr(self, 'rad_current_page') and self.rad_current_page.isChecked():
                    curr_p = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))
                    pages_to_process = [max(0, min(curr_p, total_pages - 1))]
                elif self.rad_all_pages.isChecked():
                    pages_to_process = list(range(total_pages))
                elif self.rad_skip_first.isChecked():
                    pages_to_process = list(range(1, total_pages))
                elif self.rad_first_page.isChecked():
                    pages_to_process = [0]
                else:
                    pages_to_process = parse_page_range(self.txt_custom_pages.text().strip(), total_pages)

                if not pages_to_process:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("warning"), "Lütfen geçerli en az bir sayfa belirleyin.")
                    return

                if is_color_tab:
                    opacity = float(self.sld_opacity.value()) / 100.0
                    r_norm = self.current_bg_color.redF()
                    g_norm = self.current_bg_color.greenF()
                    b_norm = self.current_bg_color.blueF()

                for page_idx in pages_to_process:
                    page = doc[page_idx]
                    if page.rotation != 0:
                        page.remove_rotation()

                    # 1. Önceki tüm arka plan akışlarını temizle
                    self._remove_page_background(doc, page)

                    # 2. Yeni arka planı uygula
                    if is_color_tab:
                        page.draw_rect(
                            page.rect,
                            fill=(r_norm, g_norm, b_norm),
                            color=None,
                            fill_opacity=opacity,
                            overlay=False
                        )
                    else:
                        rect = page.rect
                        if fit_mode == "fill":
                            target_rect = rect
                        else:
                            page_aspect = rect.height / max(1.0, rect.width)
                            if aspect > page_aspect:
                                th = rect.height
                                tw = th / aspect
                            else:
                                tw = rect.width
                                th = tw * aspect
                            cx = (rect.x0 + rect.x1) / 2.0
                            cy = (rect.y0 + rect.y1) / 2.0
                            target_rect = fitz.Rect(cx - tw / 2.0, cy - th / 2.0, cx + tw / 2.0, cy + th / 2.0)
                        page.insert_image(target_rect, stream=img_stream, overlay=False)

                    # 3. Eklenen yeni arka plan akışını etiketle
                    if bg_contents := page.get_contents():
                        try:
                            doc.update_stream(bg_contents[0], b'% AuraPDF-Background\n' + doc.xref_stream(bg_contents[0]))
                        except Exception:
                            pass

                old_bytes = tab_doc.doc.tobytes()
                new_bytes = doc.tobytes()
                old_drawings = [dict(d) for d in getattr(tab_doc, 'drawings', [])]
                old_page = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))

                tab_doc.doc.close()
                tab_doc.doc = doc
                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations()
                tab_doc.set_modified(True)
                tab_doc.render_pages_init()
                tab_doc.rerender_visible_pages()

                curr_p = max(0, min(old_page, len(doc) - 1))
                tab_doc.go_to_page(curr_p)
                if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                    if self.main_window.thumbnails_dock.isVisible():
                        self.main_window.thumbnails_dock.load_thumbnails_lazy()
                        self.main_window.thumbnails_dock.select_page(curr_p)

                tab_doc.push_undo_action({
                    "type": "reorder_pages",
                    "old_bytes": old_bytes,
                    "old_drawings": old_drawings,
                    "old_page": curr_p,
                    "new_bytes": new_bytes,
                    "new_drawings": [dict(d) for d in getattr(tab_doc, 'drawings', [])],
                    "new_page": curr_p,
                    "new_order": list(range(len(tab_doc.doc)))
                })

                QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("background_success_apply", default="Arka plan açık belgeye başarıyla uygulandı!"))
                self.accept()
                return
            except Exception as e:
                QMessageBox.critical(self, LanguageManager.tr("error"), str(e))
                return

        # 2. Farklı Kaydet (Dosyaya Yazma)
        dst_path = self.txt_dst.text().strip()
        if not dst_path or is_active_tab:
            default_name = "belge_arkaplanli.pdf"
            if is_active_tab and tab_doc:
                doc_title = getattr(tab_doc, 'custom_title', None) or "Adsiz"
                default_name = f"{doc_title}_arkaplanli.pdf"
            dst_path, _ = QFileDialog.getSaveFileName(self, LanguageManager.tr("save_dialog_title"), dst_path or default_name, "PDF (*.pdf)")
            if not dst_path:
                return
            self.txt_dst.setText(dst_path)

        try:
            if is_active_tab and tab_doc and tab_doc.doc:
                doc = fitz.open("pdf", tab_doc.doc.tobytes())
            else:
                doc = fitz.open(src_path)
                if doc.is_encrypted:
                    from PyQt6.QtWidgets import QInputDialog
                    pw, ok = QInputDialog.getText(self, LanguageManager.tr("password_required_title"), LanguageManager.tr("encrypt_current_pw"), QLineEdit.EchoMode.Password)
                    if not ok or doc.authenticate(pw) <= 0:
                        doc.close()
                        QMessageBox.warning(self, LanguageManager.tr("error"), LanguageManager.tr("encrypt_err_wrong_pw"))
                        return

            total_pages = len(doc)
            if hasattr(self, 'rad_current_page') and self.rad_current_page.isChecked():
                curr_p = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0)) if tab_doc else 0
                pages_to_process = [max(0, min(curr_p, total_pages - 1))]
            elif self.rad_all_pages.isChecked():
                pages_to_process = list(range(total_pages))
            elif self.rad_skip_first.isChecked():
                pages_to_process = list(range(1, total_pages))
            elif self.rad_first_page.isChecked():
                pages_to_process = [0]
            else:
                pages_to_process = parse_page_range(self.txt_custom_pages.text().strip(), total_pages)

            if not pages_to_process:
                doc.close()
                QMessageBox.warning(self, LanguageManager.tr("warning"), "Lütfen geçerli en az bir sayfa belirleyin.")
                return

            if is_color_tab:
                opacity = float(self.sld_opacity.value()) / 100.0
                r_norm = self.current_bg_color.redF()
                g_norm = self.current_bg_color.greenF()
                b_norm = self.current_bg_color.blueF()

            for page_idx in pages_to_process:
                page = doc[page_idx]
                if page.rotation != 0:
                    page.remove_rotation()

                # 1. Önceki tüm arka plan akışlarını temizle
                self._remove_page_background(doc, page)

                # 2. Yeni arka planı uygula
                if is_color_tab:
                    page.draw_rect(
                        page.rect,
                        fill=(r_norm, g_norm, b_norm),
                        color=None,
                        fill_opacity=opacity,
                        overlay=False
                    )
                else:
                    rect = page.rect
                    if fit_mode == "fill":
                        target_rect = rect
                    else:
                        page_aspect = rect.height / max(1.0, rect.width)
                        if aspect > page_aspect:
                            th = rect.height
                            tw = th / aspect
                        else:
                            tw = rect.width
                            th = tw * aspect
                        cx = (rect.x0 + rect.x1) / 2.0
                        cy = (rect.y0 + rect.y1) / 2.0
                        target_rect = fitz.Rect(cx - tw / 2.0, cy - th / 2.0, cx + tw / 2.0, cy + th / 2.0)
                    page.insert_image(target_rect, stream=img_stream, overlay=False)

                # 3. Eklenen yeni arka plan akışını etiketle
                if bg_contents := page.get_contents():
                    try:
                        doc.update_stream(bg_contents[0], b'% AuraPDF-Background\n' + doc.xref_stream(bg_contents[0]))
                    except Exception:
                        pass

            doc.save(dst_path, garbage=3, deflate=True)
            doc.close()

            QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("background_success", path=dst_path))
            if self.chk_open_tab.isChecked() and hasattr(self.main_window, 'add_pdf_tab'):
                self.main_window.add_pdf_tab(dst_path)
            self.accept()
        except Exception as e:
            QMessageBox.critical(self, LanguageManager.tr("error"), str(e))

    def remove_background(self):
        """Seçili sayfalardaki veya tüm belgedeki mevcut arka planı kaldırır."""
        src_path = self.txt_src.text().strip()
        is_active_tab = self._is_active_tab_doc(src_path)
        tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()

        if not is_active_tab and (not src_path or not os.path.isfile(src_path)):
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("encrypt_err_no_file"))
            return

        if is_active_tab and tab_doc and tab_doc.doc:
            try:
                doc = fitz.open("pdf", tab_doc.doc.tobytes())
                total_pages = len(doc)
                if hasattr(self, 'rad_current_page') and self.rad_current_page.isChecked():
                    curr_p = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))
                    pages_to_process = [max(0, min(curr_p, total_pages - 1))]
                elif self.rad_all_pages.isChecked():
                    pages_to_process = list(range(total_pages))
                elif self.rad_skip_first.isChecked():
                    pages_to_process = list(range(1, total_pages))
                elif self.rad_first_page.isChecked():
                    pages_to_process = [0]
                else:
                    pages_to_process = parse_page_range(self.txt_custom_pages.text().strip(), total_pages)

                if not pages_to_process:
                    doc.close()
                    QMessageBox.warning(self, LanguageManager.tr("warning"), "Lütfen geçerli en az bir sayfa belirleyin.")
                    return

                for p_idx in pages_to_process:
                    self._remove_page_background(doc, doc[p_idx])

                old_bytes = tab_doc.doc.tobytes()
                new_bytes = doc.tobytes()
                old_drawings = [dict(d) for d in getattr(tab_doc, 'drawings', [])]
                old_page = getattr(tab_doc, 'current_page_idx', getattr(tab_doc, 'current_page', 0))

                tab_doc.doc.close()
                tab_doc.doc = doc
                if hasattr(tab_doc, 'hide_managed_annotations'):
                    tab_doc.hide_managed_annotations()
                tab_doc.set_modified(True)
                tab_doc.render_pages_init()
                tab_doc.rerender_visible_pages()

                curr_p = max(0, min(old_page, len(doc) - 1))
                tab_doc.go_to_page(curr_p)
                if hasattr(self.main_window, 'thumbnails_dock') and self.main_window.thumbnails_dock:
                    if self.main_window.thumbnails_dock.isVisible():
                        self.main_window.thumbnails_dock.load_thumbnails_lazy()
                        self.main_window.thumbnails_dock.select_page(curr_p)

                tab_doc.push_undo_action({
                    "type": "reorder_pages",
                    "old_bytes": old_bytes,
                    "old_drawings": old_drawings,
                    "old_page": curr_p,
                    "new_bytes": new_bytes,
                    "new_drawings": [dict(d) for d in getattr(tab_doc, 'drawings', [])],
                    "new_page": curr_p,
                    "new_order": list(range(len(tab_doc.doc)))
                })

                QMessageBox.information(self, LanguageManager.tr("success"), LanguageManager.tr("background_success_remove"))
                self.accept()
                return
            except Exception as e:
                QMessageBox.critical(self, LanguageManager.tr("error"), str(e))
                return


class PDFPasswordDialog(QDialog):
    def __init__(self, parent=None, file_path: str = "", hint: Optional[str] = None):
        if isinstance(parent, str) and not file_path:
            file_path = parent
            parent = None
        super().__init__(parent)
        self.file_path = file_path or ""
        self.hint = hint
        self.password = ""
        self.recovered_file = None
        self.found_pw = None

        self.setWindowTitle(LanguageManager.tr("password_required_title"))
        self.resize(520, 360 if not hint else 420)
        self.setMinimumSize(420, 280)
        self.setSizeGripEnabled(True)
        apply_modern_dialog_style(self)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(14)

        # Başlık ve dosya bilgisi
        fname = os.path.basename(self.file_path) if self.file_path else ""
        lbl_prompt = QLabel(LanguageManager.tr("password_required_prompt", filename=fname))
        lbl_prompt.setWordWrap(True)
        lbl_prompt.setStyleSheet("font-size: 13px; font-weight: 500;")
        layout.addWidget(lbl_prompt)

        # Şifre İpucu (varsa)
        if self.hint:
            card_hint = ModernInfoCard(
                text=LanguageManager.tr("password_hint_display", hint=self.hint),
                parent=self
            )
            layout.addWidget(card_hint)

        # Parola Giriş Alanı
        self.txt_pw = QLineEdit()
        self.txt_pw.setEchoMode(QLineEdit.EchoMode.Password)
        self.txt_pw.setPlaceholderText(LanguageManager.tr("password_placeholder"))
        self.txt_pw.returnPressed.connect(self._on_unlock)
        layout.addWidget(self.txt_pw)

        # Parolayı Göster Checkbox
        self.chk_show = QCheckBox(LanguageManager.tr("encrypt_show_pw"))
        self.chk_show.toggled.connect(self._toggle_show_pw)
        layout.addWidget(self.chk_show)

        layout.addStretch()

        # Alt Butonlar
        h_btn = QHBoxLayout()
        btn_recover = QPushButton(LanguageManager.tr("password_btn_recover"))
        btn_recover.setObjectName("RecoverBtn")
        btn_recover.clicked.connect(self._open_recovery)
        h_btn.addWidget(btn_recover)

        h_btn.addStretch()

        btn_cancel = QPushButton(LanguageManager.tr("print_btn_cancel"))
        btn_cancel.clicked.connect(self.reject)
        h_btn.addWidget(btn_cancel)

        btn_unlock = QPushButton(LanguageManager.tr("password_btn_unlock"))
        btn_unlock.setObjectName("PrimaryBtn")
        btn_unlock.clicked.connect(self._on_unlock)
        h_btn.addWidget(btn_unlock)

        layout.addLayout(h_btn)

    def _toggle_show_pw(self, checked: bool):
        mode = QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
        self.txt_pw.setEchoMode(mode)

    def _open_recovery(self):
        dlg = PDFPasswordRecoveryDialog(self, initial_file=self.file_path)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            if dlg.found_pw:
                self.found_pw = dlg.found_pw
                self.password = dlg.found_pw
                self.accept()
            elif hasattr(dlg, 'txt_brute_dst') and os.path.isfile(dlg.txt_brute_dst.text().strip()):
                self.recovered_file = dlg.txt_brute_dst.text().strip()
                self.accept()

    def _on_unlock(self):
        self.password = self.txt_pw.text()
        self.accept()

    def get_password(self) -> str:
        return self.password
