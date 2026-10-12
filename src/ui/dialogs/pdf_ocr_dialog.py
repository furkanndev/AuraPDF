"""
AuraPDF - Taranmış Belgeleri Düzenlenebilir Yapma (Gelişmiş Yerel OCR)
=====================================================================
Taranmış PDF veya resim sayfalarını çevrimdışı RapidOCR motoruyla işleyerek:
1. Görünmez metin katmanı ile Aranabilir & Kopyalanabilir PDF (Sandviç PDF)
2. Düzenlenebilir Microsoft Word Belgesi (.docx)
3. Yapılandırılmış Metin Dosyası (.txt)
olarak dönüştüren profesyonel araç.
"""

import os
import io

from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QFileDialog, QWidget, QFrame, QRadioButton, QButtonGroup,
    QProgressBar, QTextEdit, QLineEdit, QMessageBox, QSplitter
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QFont

import pymupdf as fitz
from src.core.i18n import LanguageManager
from src.ui.dialog_theme import apply_dialog_theme, get_dialog_theme_stylesheet, ModernInfoCard
from src.ui.icons import ModernIconFactory


class OCRWorker(QThread):
    progress = pyqtSignal(int, int, str)  # current, total, status_text
    page_done = pyqtSignal(int, str, list)  # page_idx, text, boxes
    finished_all = pyqtSignal(bool, str)   # success, message

    def __init__(self, doc, page_indices, mode, output_path=None):
        super().__init__()
        self.doc = doc
        self.page_indices = page_indices
        self.mode = mode  # 'sandwich', 'docx', 'txt'
        self.output_path = output_path
        self._is_cancelled = False

    def cancel(self):
        self._is_cancelled = True

    def run(self):
        try:
            import numpy as np
            from rapidocr_onnxruntime import RapidOCR
            ocr_engine = RapidOCR()
        except Exception as e:
            self.finished_all.emit(False, f"RapidOCR motoru başlatılamadı:\n{e}")
            return

        total = len(self.page_indices)
        all_extracted_text = []

        # Eğer sandwich modunda ise ve yeni bir dosya oluşturulacaksa
        target_doc = None
        if self.mode == "sandwich":
            target_doc = fitz.open()
            target_doc.insert_pdf(self.doc)

        for step, p_idx in enumerate(self.page_indices):
            if self._is_cancelled:
                self.finished_all.emit(False, "İşlem kullanıcı tarafından iptal edildi.")
                return

            self.progress.emit(step + 1, total, f"Sayfa {p_idx + 1} taranıyor ve metin tanınıyor...")

            try:
                page = self.doc[p_idx]
                zoom = 2.0  # 144 DPI yüksek doğruluk
                mat = fitz.Matrix(zoom, zoom)
                pix = page.get_pixmap(matrix=mat, alpha=False)

                img_arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape((pix.height, pix.width, 3))
                result, _ = ocr_engine(img_arr)

                page_text_lines = []
                page_boxes = []

                if result:
                    for line_info in result:
                        box = line_info[0]  # [[x0, y0], [x1, y1], [x2, y2], [x3, y3]]
                        text = line_info[1]
                        score = line_info[2] if len(line_info) > 2 else 1.0

                        if text and text.strip():
                            page_text_lines.append(text.strip())
                            # Sayfa koordinatlarına geri dönüştür
                            x0 = min(pt[0] for pt in box) / zoom
                            y0 = min(pt[1] for pt in box) / zoom
                            x1 = max(pt[0] for pt in box) / zoom
                            y1 = max(pt[1] for pt in box) / zoom
                            page_boxes.append((fitz.Rect(x0, y0, x1, y1), text.strip()))

                page_full_text = "\n".join(page_text_lines)
                all_extracted_text.append(f"--- Sayfa {p_idx + 1} ---\n" + page_full_text)
                self.page_done.emit(p_idx, page_full_text, page_boxes)

                # Sandviç modunda görünmez metin katmanı ekle
                if self.mode == "sandwich" and target_doc:
                    t_page = target_doc[p_idx]
                    for rect, text in page_boxes:
                        try:
                            # render_mode=3 PDF standardında görünmez metindir (seçilebilir, aranabilir)
                            fs = max(7, min(rect.height * 0.75, 24))
                            t_page.insert_textbox(rect, text, fontsize=fs, render_mode=3)
                        except Exception:
                            pass

            except Exception as e:
                print(f"Sayfa {p_idx+1} OCR hatası:", e)

        # Sonuçları kaydet
        try:
            if self.mode == "sandwich":
                if self.output_path and target_doc:
                    target_doc.save(self.output_path, garbage=4, deflate=True)
                    target_doc.close()
                self.finished_all.emit(True, f"Aranabilir ve kopyalanabilir PDF oluşturuldu:\n{self.output_path}")

            elif self.mode == "docx":
                import docx
                doc_out = docx.Document()
                for t in all_extracted_text:
                    for p_line in t.splitlines():
                        if p_line.startswith("--- Sayfa"):
                            doc_out.add_heading(p_line, level=2)
                        elif p_line.strip():
                            doc_out.add_paragraph(p_line)
                doc_out.save(self.output_path)
                self.finished_all.emit(True, f"Düzenlenebilir Word belgesi oluşturuldu:\n{self.output_path}")

            elif self.mode == "txt":
                with open(self.output_path, "w", encoding="utf-8") as f:
                    f.write("\n\n".join(all_extracted_text))
                self.finished_all.emit(True, f"Metin dosyası kaydedildi:\n{self.output_path}")

        except Exception as e:
            self.finished_all.emit(False, f"Dosya kaydedilirken hata oluştu:\n{e}")


class PDFScanOCRDialog(QDialog):
    def __init__(self, main_window, parent=None):
        super().__init__(parent or main_window)
        self.main_window = main_window
        self.setWindowTitle(LanguageManager.tr("ocr_dialog_title"))
        self.setWindowIcon(ModernIconFactory.icon_ocr())
        self.resize(950, 660)
        self.setMinimumSize(780, 520)
        self.setSizeGripEnabled(True)

        self.current_theme = getattr(self.main_window, 'current_theme', 'dark')

        self.tab_doc = getattr(self.main_window, 'get_current_tab_doc', lambda: None)()
        self.doc = None
        self.file_path = ""
        self.current_page = 0
        self.worker = None

        if self.tab_doc and getattr(self.tab_doc, 'doc', None):
            self.doc = self.tab_doc.doc
            self.file_path = getattr(self.tab_doc, 'file_path', "")
            self.current_page = getattr(self.tab_doc, 'current_page_idx', 0)

        self._init_ui()
        self._apply_style()

    def _apply_style(self):
        apply_dialog_theme(self, self.main_window, extra_css="""
            QFrame.card {
                border-radius: 8px;
                padding: 10px;
            }
            QPushButton#PrimaryBtn {
                font-weight: bold;
                border: none;
            }
            QProgressBar {
                border-radius: 5px;
                text-align: center;
            }
            QTextEdit {
                border-radius: 6px;
                padding: 8px;
            }
        """)

    def _init_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(10)

        # Üst Bilgi Kartı
        header = ModernInfoCard(
            text=LanguageManager.tr("ocr_desc"),
            title=LanguageManager.tr("ocr_dialog_header_title"),
            icon_type="ocr",
            parent=self
        )
        root.addWidget(header)

        # Bölücü (Sol Ayarlar, Sağ Canlı Metin Önizlemesi)
        splitter = QSplitter(Qt.Orientation.Horizontal)

        # Sol Ayarlar
        left_box = QWidget()
        l_lay = QVBoxLayout(left_box)
        l_lay.setContentsMargins(0, 0, 0, 0)
        l_lay.setSpacing(10)

        # Dönüştürme Modu
        mod_card = QFrame()
        mod_card.setProperty("class", "card")
        mc_lay = QVBoxLayout(mod_card)
        mc_lay.setSpacing(8)

        mc_lay.addWidget(QLabel(f"<b>{LanguageManager.tr('ocr_conversion_format_lbl')}</b>"))
        self.rb_sandwich = QRadioButton(LanguageManager.tr("ocr_mode_sandwich"))
        self.rb_sandwich.setChecked(True)
        mc_lay.addWidget(self.rb_sandwich)

        sub1 = QLabel(LanguageManager.tr("ocr_sandwich_sub_desc"))
        sub1.setStyleSheet(" font-size: 11px; margin-left: 20px;")
        sub1.setWordWrap(True)
        mc_lay.addWidget(sub1)

        self.rb_docx = QRadioButton(LanguageManager.tr("ocr_mode_docx"))
        mc_lay.addWidget(self.rb_docx)

        self.rb_text = QRadioButton(LanguageManager.tr("ocr_mode_text"))
        mc_lay.addWidget(self.rb_text)

        group_mode = QButtonGroup(self)
        group_mode.addButton(self.rb_sandwich)
        group_mode.addButton(self.rb_docx)
        group_mode.addButton(self.rb_text)

        l_lay.addWidget(mod_card)

        # Sayfa Aralığı
        page_card = QFrame()
        page_card.setProperty("class", "card")
        pc_lay = QVBoxLayout(page_card)
        pc_lay.setSpacing(8)

        pc_lay.addWidget(QLabel(f"<b>{LanguageManager.tr('ocr_scope_pages')}</b>"))
        self.rb_scope_curr = QRadioButton(LanguageManager.tr("ocr_scope_curr_page", page=self.current_page + 1))
        self.rb_scope_all = QRadioButton(LanguageManager.tr("ocr_scope_all_pages", count=len(self.doc) if self.doc else 0))
        self.rb_scope_all.setChecked(True)
        self.rb_scope_range = QRadioButton(LanguageManager.tr("ocr_scope_range"))
        self.txt_page_range = QLineEdit()
        self.txt_page_range.setPlaceholderText(LanguageManager.tr("ocr_range_placeholder"))

        range_lay = QHBoxLayout()
        range_lay.addWidget(self.rb_scope_range)
        range_lay.addWidget(self.txt_page_range)

        group_scope = QButtonGroup(self)
        group_scope.addButton(self.rb_scope_curr)
        group_scope.addButton(self.rb_scope_all)
        group_scope.addButton(self.rb_scope_range)

        pc_lay.addWidget(self.rb_scope_curr)
        pc_lay.addWidget(self.rb_scope_all)
        pc_lay.addLayout(range_lay)

        l_lay.addWidget(page_card)
        l_lay.addStretch()

        splitter.addWidget(left_box)

        # Sağ: Tanınan Metin Önizlemesi
        right_box = QWidget()
        r_lay = QVBoxLayout(right_box)
        r_lay.setContentsMargins(0, 0, 0, 0)
        r_lay.setSpacing(6)

        r_lay.addWidget(QLabel(f"<b>{LanguageManager.tr('ocr_recognized_stream_title')}</b>"))
        self.txt_preview = QTextEdit()
        self.txt_preview.setReadOnly(True)
        self.txt_preview.setFont(QFont("Consolas", 10))
        self.txt_preview.setPlaceholderText(LanguageManager.tr("ocr_preview_placeholder"))
        r_lay.addWidget(self.txt_preview)

        splitter.addWidget(right_box)
        splitter.setSizes([420, 500])

        root.addWidget(splitter, 1)

        # İlerleme Çubuğu & Durum
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.progress_bar.setFixedHeight(18)
        root.addWidget(self.progress_bar)

        self.lbl_status = QLabel(LanguageManager.tr("ocr_status_start_hint"))
        self.lbl_status.setStyleSheet(" font-size: 11px;")
        root.addWidget(self.lbl_status)

        # Alt Butonlar
        bottom_bar = QHBoxLayout()

        self.btn_start = QPushButton(LanguageManager.tr("ocr_btn_start"))
        self.btn_start.setObjectName("PrimaryBtn")
        self.btn_start.clicked.connect(self._start_ocr_process)
        bottom_bar.addWidget(self.btn_start)

        self.btn_cancel_process = QPushButton(LanguageManager.tr("ocr_btn_stop"))
        self.btn_cancel_process.setVisible(False)
        self.btn_cancel_process.clicked.connect(self._cancel_ocr_process)
        bottom_bar.addWidget(self.btn_cancel_process)

        bottom_bar.addStretch()

        btn_close = QPushButton(LanguageManager.tr("btn_cancel_settings"))
        btn_close.clicked.connect(self.reject)
        bottom_bar.addWidget(btn_close)

        root.addLayout(bottom_bar)

    def _start_ocr_process(self):
        if not self.doc or len(self.doc) == 0:
            QMessageBox.warning(self, LanguageManager.tr("warning"), LanguageManager.tr("ocr_no_doc"))
            return

        # Sayfa aralığını hesapla
        total_pages = len(self.doc)
        pages_to_process = []

        if self.rb_scope_curr.isChecked():
            pages_to_process = [self.current_page]
        elif self.rb_scope_all.isChecked():
            pages_to_process = list(range(total_pages))
        elif self.rb_scope_range.isChecked():
            raw = self.txt_page_range.text().strip()
            if not raw:
                QMessageBox.warning(self, "Uyarı", "Lütfen geçerli bir sayfa aralığı girin (örn: 1-5).")
                return
            try:
                for part in raw.split(","):
                    part = part.strip()
                    if "-" in part:
                        s, e = part.split("-")
                        pages_to_process.extend(range(int(s) - 1, int(e)))
                    elif part.isdigit():
                        pages_to_process.append(int(part) - 1)
                pages_to_process = [p for p in sorted(set(pages_to_process)) if 0 <= p < total_pages]
            except Exception:
                QMessageBox.warning(self, "Uyarı", "Sayfa aralığı formatı geçersiz.")
                return

        if not pages_to_process:
            QMessageBox.warning(self, "Uyarı", "İşlenecek sayfa bulunamadı.")
            return

        # Mod ve Hedef dosya seçimi
        mode = "sandwich"
        filter_str = "PDF Dosyaları (*.pdf)"
        default_ext = "_aranabilir.pdf"

        if self.rb_docx.isChecked():
            mode = "docx"
            filter_str = "Word Belgeleri (*.docx)"
            default_ext = "_duzenlenebilir.docx"
        elif self.rb_text.isChecked():
            mode = "txt"
            filter_str = "Metin Belgeleri (*.txt)"
            default_ext = "_metin.txt"

        base_name = os.path.splitext(self.file_path)[0] if self.file_path else "belge"
        out_name = f"{base_name}{default_ext}"

        save_path, _ = QFileDialog.getSaveFileName(
            self, "Dönüştürülen Dosyayı Kaydet", out_name, filter_str
        )
        if not save_path:
            return

        # UI güncelle
        self.btn_start.setEnabled(False)
        self.btn_cancel_process.setVisible(True)
        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(len(pages_to_process))
        self.progress_bar.setValue(0)
        self.txt_preview.clear()

        # Worker başlat
        self.worker = OCRWorker(self.doc, pages_to_process, mode, save_path)
        self.worker.progress.connect(self._on_progress)
        self.worker.page_done.connect(self._on_page_done)
        self.worker.finished_all.connect(self._on_finished)
        self.worker.start()

    def _cancel_ocr_process(self):
        if self.worker and self.worker.isRunning():
            self.worker.cancel()
            self.lbl_status.setText(LanguageManager.tr("ocr_cancelling"))

    def _on_progress(self, current, total, text):
        self.progress_bar.setValue(current)
        self.lbl_status.setText(f"[{current}/{total}] {text}")

    def _on_page_done(self, p_idx, text, boxes):
        if text.strip():
            self.txt_preview.append(LanguageManager.tr("ocr_page_box_header", page=p_idx + 1, count=len(boxes), text=text))
        else:
            self.txt_preview.append(LanguageManager.tr("ocr_page_no_text", page=p_idx + 1))

    def _on_finished(self, success, message):
        self.btn_start.setEnabled(True)
        self.btn_cancel_process.setVisible(False)
        self.progress_bar.setVisible(False)

        if success:
            self.lbl_status.setText(LanguageManager.tr("ocr_completed"))
            QMessageBox.information(self, LanguageManager.tr("success"), message)
            self.accept()
        else:
            self.lbl_status.setText(LanguageManager.tr("ocr_failed"))
            QMessageBox.warning(self, LanguageManager.tr("info"), message)


PDFOCRDialog = PDFScanOCRDialog
